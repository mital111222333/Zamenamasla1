"""
Веб-панель системы учёта замены масла — мультитенантная версия.

- /login — вход по логину/паролю (для точки замены масла ИЛИ платформенного
  администратора). Никакого доступа без входа.
- Точка (role='shop') видит и редактирует ТОЛЬКО свои данные — во всех
  запросах ниже используется g.shop_id из сессии, и все обращения к БД идут
  с этим shop_id. Это и есть изоляция данных между точками.
- Админ (role='admin') управляет списком точек на /admin — создаёт новые,
  включает/выключает — но не видит клиентских данных ни одной из них.
"""

import os
import re
import time
import secrets
import logging
import urllib.parse
import threading
import requests
from functools import wraps
from urllib.parse import quote
from flask import Flask, request, jsonify, render_template_string, Response, session, redirect, url_for, g

import database as db
import i18n

logger = logging.getLogger(__name__)

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")

PUBLIC_URL = os.environ.get("PUBLIC_URL", "")
BOT_USERNAME = os.environ.get("BOT_USERNAME", "")
BOT_TOKEN = os.environ.get("BOT_TOKEN", "")
DISPLAY_SHOW_SECONDS = int(os.environ.get("DISPLAY_SHOW_SECONDS", "45"))


def _send_telegram_message(chat_id, text) -> bool:
    """Отправка сообщения напрямую через HTTP API Telegram — синхронно, без
    участия основного бот-процесса (веб-панель работает в отдельном потоке
    того же процесса, но без доступа к его асинхронному event loop).
    Частая причина отказа: получатель ни разу не писал этому боту — Telegram
    не разрешает боту писать первым, даже если chat_id указан верно."""
    if not BOT_TOKEN:
        logger.warning("Telegram-сообщение не отправлено: BOT_TOKEN не задан")
        return False
    if not chat_id:
        logger.warning("Telegram-сообщение не отправлено: chat_id не указан")
        return False
    try:
        resp = requests.post(
            f"https://api.telegram.org/bot{BOT_TOKEN}/sendMessage",
            json={"chat_id": chat_id, "text": text},
            timeout=10,
        )
        if not resp.ok:
            logger.warning(f"Telegram отклонил сообщение для chat_id={chat_id}: "
                            f"HTTP {resp.status_code} — {resp.text[:300]}")
        return resp.ok
    except Exception as e:
        logger.error(f"Не удалось отправить Telegram-сообщение для chat_id={chat_id}: {e}")
        return False

# Состояние табло — отдельно для каждой точки (по shop_id), чтобы камера
# одной точки не могла подсветить экран другой.
_display_lock = threading.Lock()
_display_states = {}

CAR_BRANDS = [
    "Chevrolet", "Daewoo", "Ravon", "Kia", "Hyundai", "Toyota", "Lexus",
    "Nissan", "Isuzu", "BMW", "Mercedes-Benz", "Audi", "Volkswagen",
    "Lada (ВАЗ)", "Datsun", "Honda", "Mazda", "Ford", "Mitsubishi", "Другое",
]
SERVICE_TYPES = ["Замена масла", "Замена масла + фильтр", "Полное ТО", "Другое"]


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if session.get("role") not in ("shop", "branch") or not session.get("shop_id"):
            return redirect(url_for("login_page"))
        shop = db.get_shop(session["shop_id"])
        if not shop or not shop["is_active"]:
            session.clear()
            return redirect(url_for("login_page"))
        g.shop_id = session["shop_id"]
        g.lang = shop.get("language") or "ru"
        g.T = i18n.get_texts(g.lang)
        g.is_employee = bool(session.get("is_employee"))
        g.is_branch = shop.get("role") == "branch"
        g.parent_shop_id = shop.get("parent_shop_id")
        return view(*args, **kwargs)
    return wrapped


def employee_blocked(view):
    """Закрывает API-эндпоинт для логина сотрудника (прибыль, цены закупки,
    экспорт, управление складом) — даже если кто-то обратится к нему напрямую,
    в обход интерфейса. Ставить ПОСЛЕ @login_required."""
    @wraps(view)
    def wrapped(*args, **kwargs):
        if getattr(g, "is_employee", False):
            return jsonify({"ok": False, "error": "недоступно для этого аккаунта"}), 403
        return view(*args, **kwargs)
    return wrapped


def profit_blocked(view):
    """Закрывает прибыль и от сотрудника, и от филиала — прибыль по филиалам
    видит только их главный аккаунт (в сложенном виде, см. агрегированную
    статистику). Остальное (статистика по выручке, экспорт, рассылка, SMS)
    филиалу по-прежнему доступно, поэтому это отдельный декоратор, а не
    расширение employee_blocked."""
    @wraps(view)
    def wrapped(*args, **kwargs):
        if getattr(g, "is_employee", False) or getattr(g, "is_branch", False):
            return jsonify({"ok": False, "error": "недоступно для этого аккаунта"}), 403
        return view(*args, **kwargs)
    return wrapped



def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if session.get("role") != "admin":
            return redirect(url_for("login_page"))
        return view(*args, **kwargs)
    return wrapped


def _client_link(token):
    if not token or not BOT_USERNAME:
        return None
    return f"https://t.me/{BOT_USERNAME}?start={token}"


LOGIN_PAGE = """
<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{{ T.app_title }}</title>
<style>
  * { box-sizing: border-box; }
  body {
    margin:0; background:#0f1115; color:#f2f2f2; font-family: -apple-system, Segoe UI, Roboto, sans-serif;
    height:100vh; display:flex; align-items:center; justify-content:center;
  }
  .box { background:#1a1d24; border:1px solid #2a2e37; border-radius:14px; padding:28px; width:90%; max-width:340px; }
  h1 { font-size:20px; margin:0 0 20px; text-align:center; }
  label { display:block; font-size:13px; color:#9a9a9a; margin-bottom:4px; }
  input { width:100%; padding:11px; border-radius:8px; border:1px solid #2a2e37; background:#11141a; color:#fff; font-size:15px; margin-bottom:14px; }
  button { width:100%; padding:12px; border:none; border-radius:10px; background:#3a86ff; color:#fff; font-size:16px; font-weight:600; cursor:pointer; }
  .error { background:#3a1e1e; color:#dc6f6f; padding:10px; border-radius:8px; margin-bottom:14px; font-size:14px; }
  .ok-msg { background:#1e3a24; color:#6fdc86; padding:10px; border-radius:8px; margin-bottom:14px; font-size:14px; }
  .lang-link { display:block; text-align:center; margin-top:14px; color:#9a9a9a; font-size:12px; text-decoration:none; }
  .forgot-link { display:block; text-align:center; margin-top:12px; color:#3a86ff; font-size:13px; background:none; border:none; cursor:pointer; padding:0; }
  .hint { font-size:12px; color:#7a7a7a; margin:-8px 0 14px; }
</style>
</head>
<body>
  <form class="box" method="POST" id="loginForm">
    <h1>🔧 {{ T.login_title }}</h1>
    {% if error %}<div class="error">{{ error }}</div>{% endif %}
    <input type="hidden" name="_lang" value="{{ lang }}">
    <label>{{ T.login_username }}</label>
    <input name="username" autofocus required>
    <label>{{ T.login_password }}</label>
    <input name="password" type="password" required>
    <button type="submit">{{ T.login_button }}</button>
    <button type="button" class="forgot-link" onclick="showForgot()">{{ T.forgot_password_link }}</button>
    <a class="lang-link" href="/login?lang={{ other_lang }}">{{ T.lang_switch }}</a>
  </form>

  <div class="box" id="forgotBox" style="display:none;">
    <h1>🔑 {{ T.forgot_title }}</h1>
    <div id="forgotMsg"></div>
    <div id="forgotStep1">
      <label>{{ T.login_username }}</label>
      <input id="forgot_username" autofocus>
      <p class="hint">{{ T.forgot_hint }}</p>
      <button type="button" onclick="requestResetCode()">{{ T.forgot_send_code }}</button>
    </div>
    <div id="forgotStep2" style="display:none;">
      <label>{{ T.forgot_code_label }}</label>
      <input id="forgot_code" inputmode="numeric" maxlength="6">
      <label>{{ T.forgot_new_password }}</label>
      <input id="forgot_new_password" type="password">
      <button type="button" onclick="confirmResetCode()">{{ T.forgot_save_btn }}</button>
    </div>
    <button type="button" class="forgot-link" onclick="hideForgot()">{{ T.forgot_back }}</button>
  </div>

  <script>
    function showForgot() {
      document.getElementById('loginForm').style.display = 'none';
      document.getElementById('forgotBox').style.display = 'block';
    }
    function hideForgot() {
      document.getElementById('forgotBox').style.display = 'none';
      document.getElementById('loginForm').style.display = 'block';
      document.getElementById('forgotStep1').style.display = 'block';
      document.getElementById('forgotStep2').style.display = 'none';
      document.getElementById('forgotMsg').innerHTML = '';
    }
    async function requestResetCode() {
      const username = document.getElementById('forgot_username').value.trim();
      const msg = document.getElementById('forgotMsg');
      if (!username) return;
      const res = await fetch('/api/forgot_password/request', {
        method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({username})
      });
      const data = await res.json();
      if (data.ok) {
        msg.innerHTML = `<div class="ok-msg">{{ T.forgot_code_sent }}</div>`;
        document.getElementById('forgotStep1').style.display = 'none';
        document.getElementById('forgotStep2').style.display = 'block';
      } else {
        msg.innerHTML = `<div class="error">${data.error}</div>`;
      }
    }
    async function confirmResetCode() {
      const username = document.getElementById('forgot_username').value.trim();
      const code = document.getElementById('forgot_code').value.trim();
      const new_password = document.getElementById('forgot_new_password').value;
      const msg = document.getElementById('forgotMsg');
      const res = await fetch('/api/forgot_password/confirm', {
        method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({username, code, new_password})
      });
      const data = await res.json();
      if (data.ok) {
        msg.innerHTML = `<div class="ok-msg">{{ T.forgot_success }}</div>`;
        setTimeout(() => { window.location.href = '/login'; }, 1800);
      } else {
        msg.innerHTML = `<div class="error">${data.error}</div>`;
      }
    }
  </script>
</body>
</html>
"""


@app.route("/login", methods=["GET", "POST"])
def login_page():
    error = None
    lang = request.form.get("_lang") or request.args.get("lang")
    if lang not in ("ru", "uz"):
        lang = "ru"
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        shop = db.authenticate_shop(username, password)
        if shop:
            session.clear()
            session["shop_id"] = shop["id"]
            session["role"] = shop["role"]
            session["username"] = shop["username"]
            session["shop_name"] = shop.get("shop_name") or shop["username"]
            session["is_employee"] = False
            session.permanent = True
            return redirect(url_for("admin_page") if shop["role"] == "admin" else url_for("index"))
        employee = db.authenticate_shop_employee(username, password)
        if employee:
            parent_shop = db.get_shop(employee["shop_id"])
            session.clear()
            session["shop_id"] = employee["shop_id"]
            session["role"] = "shop"
            session["username"] = employee["username"]
            session["shop_name"] = (parent_shop.get("shop_name") or "") if parent_shop else ""
            session["is_employee"] = True
            session.permanent = True
            return redirect(url_for("index"))
        error = i18n.t("login_error", lang)
    T = i18n.get_texts(lang)
    return render_template_string(LOGIN_PAGE, error=error, T=T, lang=lang, other_lang="uz" if lang == "ru" else "ru")


@app.route("/api/forgot_password/request", methods=["POST"])
def api_forgot_password_request():
    """Первый шаг восстановления пароля — отправляет 6-значный код в
    привязанный Telegram точки. Если Telegram не привязан, честно говорит
    об этом (это не секретная информация — точка и так это знает),
    предлагая обратиться к администратору."""
    data = request.get_json(force=True)
    username = (data.get("username") or "").strip()
    if not username:
        return jsonify({"ok": False, "error": "укажите логин"}), 400
    shop = db.find_shop_by_username(username)
    if not shop:
        return jsonify({"ok": False, "error": "точка с таким логином не найдена"}), 404
    if not shop.get("notify_telegram_id"):
        return jsonify({"ok": False, "error": "к этой точке не привязан Telegram — обратитесь к администратору платформы для сброса пароля", "no_telegram": True}), 400
    code = db.create_password_reset_code(shop["id"])
    lang = shop.get("language") or "ru"
    text = i18n.t("password_reset_code_message", lang, code=code)
    sent = _send_telegram_message(shop["notify_telegram_id"], text)
    if not sent:
        return jsonify({"ok": False, "error": "не удалось отправить код в Telegram, попробуйте позже или обратитесь к администратору"}), 500
    return jsonify({"ok": True})


@app.route("/api/forgot_password/confirm", methods=["POST"])
def api_forgot_password_confirm():
    """Второй шаг — проверяет код и, если верный, задаёт новый пароль."""
    data = request.get_json(force=True)
    username = (data.get("username") or "").strip()
    code = (data.get("code") or "").strip()
    new_password = data.get("new_password") or ""
    if not username or not code or len(new_password) < 6:
        return jsonify({"ok": False, "error": "заполните все поля (пароль — минимум 6 символов)"}), 400
    ok = db.reset_password_with_code(username, code, new_password)
    if not ok:
        return jsonify({"ok": False, "error": "неверный или просроченный код"}), 400
    return jsonify({"ok": True})


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login_page"))


PAGE = """
<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{{ shop_name }} — панель</title>
<script src="https://telegram.org/js/telegram-web-app.js"></script>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:ital,wght@0,400;0,500;0,600;0,700;0,800;1,700&family=Space+Grotesk:wght@600;700&family=IBM+Plex+Mono:wght@500;600&display=swap" rel="stylesheet">
<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.0/chart.umd.min.js"></script>
<style>
  :root {
    --bg: var(--tg-theme-bg-color, #F1F5F9);
    --text: var(--tg-theme-text-color, #1E293B);
    --hint: var(--tg-theme-hint-color, #94A3B8);
    --btn: var(--tg-theme-button-color, #E63946);
    --btn-text: var(--tg-theme-button-text-color, #FFFFFF);
    --blue: #0F52BA;
    --darkblue: #0A2540;
    --cyan: #00A8E8;
    --darkred: #9B111E;
    --carbon: #181E29;
    --card: #FFFFFF;
    --border: #E2E8F0;
    --field-bg: #F4F7FA;
    --ok: #059669;
    --ok-bg: #ECFDF5;
    --danger: #B3241C;
    --danger-bg: #FDECEA;
    --accent2: #0F52BA;
    --font-display: 'Space Grotesk', sans-serif;
    --font-body: 'Plus Jakarta Sans', -apple-system, sans-serif;
    --font-mono: 'IBM Plex Mono', monospace;
  }
  * { box-sizing: border-box; }
  body { margin:0; background:var(--bg); color:var(--text); font-family: var(--font-body); }
  .speedline {
    height:6px; width:100%;
    background: linear-gradient(90deg, var(--blue) 0%, var(--blue) 33%, #fff 33%, #fff 38%, var(--btn) 38%, var(--btn) 70%, #fff 70%, #fff 75%, var(--cyan) 75%, var(--cyan) 100%);
  }
  .container { padding: 12px; max-width: 960px; margin: 0 auto; }
  .topbar { display:flex; justify-content:space-between; align-items:center; margin: 14px 0 16px; }
  .logo-badge {
    display:inline-flex; align-items:center; justify-content:center; width:40px; height:40px;
    background:var(--darkblue); color:var(--cyan); border-radius:12px; transform:rotate(-8deg);
    box-shadow:0 4px 10px rgba(10,37,64,.25); font-size:16px; flex:none;
  }
  h1 { font-family: var(--font-display); font-weight:800; font-style:italic; letter-spacing:-0.3px; font-size: 22px; margin: 0; line-height:1.1; color:var(--darkblue); }
  h1 .accent { color:var(--btn); }
  .logo-sub { font-size:10px; font-weight:700; letter-spacing:1.5px; text-transform:uppercase; color:var(--hint); }
  .logout { color: var(--btn); font-size: 12px; font-weight:700; text-decoration:none; background:var(--danger-bg); padding:6px 10px; border-radius:10px; }
  .lang-btn { background: #EFF6FF; border: 1px solid #BFDBFE; color: var(--blue); font-size: 12px; font-weight:700; padding: 6px 10px; border-radius: 10px; cursor: pointer; }
  .tabs { display:grid; grid-template-columns: repeat(3, 1fr); gap:8px; margin-bottom: 10px; }
  .subtabs { display:flex; gap:8px; margin-bottom:14px; }
  .subtab {
    flex:1; text-align:center; padding:10px; border-radius:10px; background:var(--field-bg);
    border:1.5px solid var(--border); cursor:pointer; font-weight:700; font-size:13px; color:var(--hint);
  }
  .subtab.active { background:var(--btn); color:#fff; border-color:var(--btn); }
  .tab {
    text-align:center; padding: 10px 4px; border-radius: 14px; background: var(--card);
    border:2px solid var(--border); cursor:pointer; font-weight:700; font-family: var(--font-display);
    font-size:11px; text-transform:uppercase; letter-spacing:0.2px; transition: all .15s ease;
    display:flex; flex-direction:column; align-items:center; gap:5px;
    box-shadow: 0 1px 3px rgba(20,20,25,.04);
  }
  .tab .tab-icon {
    width:30px; height:30px; border-radius:50%; background:var(--field-bg); color:var(--blue);
    display:flex; align-items:center; justify-content:center; font-size:13px; flex:none; transition: all .15s ease;
  }
  .tab#tab-add { border-color:var(--btn); }
  .tab#tab-add .tab-icon { background:var(--danger-bg); color:var(--btn); }
  .tab#tab-add span:last-child { color:var(--btn); }
  .tab.active { color: var(--btn); }
  .tab.active .tab-icon { background:var(--btn); color:#fff; transform:scale(1.08); }
  .wh-banner {
    width:100%; background:var(--carbon); border:2px solid #0d1117; border-radius:16px; padding:12px 16px;
    display:flex; align-items:center; justify-content:center; gap:12px; cursor:pointer; margin-bottom:12px;
    box-shadow:0 8px 20px rgba(0,0,0,.25);
  }
  .wh-banner .stripe-pair { display:flex; gap:4px; }
  .wh-banner .stripe-pair span { width:8px; height:26px; border-radius:2px; transform:skewX(-12deg); }
  .wh-banner .wh-label { font-family:var(--font-display); font-weight:800; font-style:italic; font-size:18px; color:#fff; text-transform:uppercase; letter-spacing:0.5px; }
  .card {
    background: rgba(255,255,255,.94); backdrop-filter: blur(10px); border:2px solid #DBEAFE; border-radius: 22px; padding: 16px; margin-bottom: 12px;
    box-shadow: 0 10px 25px -5px rgba(15,82,186,.08); overflow:hidden;
  }
  .field { margin-bottom: 10px; }
  .row2 { display:flex; gap:10px; }
  .row2 .field { flex:1; }
  label { display:block; font-size: 12px; color: var(--hint); margin-bottom: 4px; text-transform:uppercase; letter-spacing:0.4px; }
  label i { color:var(--btn); margin-right:4px; width:12px; text-align:center; }
  input, select, textarea {
    width: 100%; padding: 10px; border-radius: 10px; border: 1.5px solid var(--border);
    background: var(--field-bg); color: var(--text); font-size: 15px; font-family: var(--font-body);
  }
  input:focus, select:focus, textarea:focus { outline:none; border-color: var(--btn); box-shadow:0 0 0 3px rgba(225,6,0,.12); }
  .card-header {
    background:linear-gradient(90deg, var(--blue), #1a6fd4); color:#fff; padding:14px 16px; border-radius:20px 20px 0 0; margin:-16px -16px 14px;
    display:flex; align-items:center; gap:10px;
  }
  .card-header .dot { width:9px; height:9px; border-radius:50%; background:var(--btn); flex:none; box-shadow:0 0 0 4px rgba(225,6,0,.25); animation: pulse 1.6s infinite; }
  @keyframes pulse { 0%,100% { opacity:1; } 50% { opacity:.4; } }
  .card-header h3 { margin:0; font-family:var(--font-display); font-weight:700; font-style:italic; font-size:17px; letter-spacing:0.3px; text-transform:uppercase; }
  .plate-wrap { position:relative; }
  .plate-chip {
    position:absolute; left:10px; top:50%; transform:translateY(-50%);
    font-size:11px; font-weight:700; color:var(--hint); background:var(--border); padding:2px 6px; border-radius:5px;
  }
  #plate { padding-left:44px; font-family:var(--font-mono); font-weight:600; letter-spacing:1px; text-transform:uppercase; }
  .plate-suggest {
    position:absolute; top:100%; left:0; right:0; margin-top:4px; background:var(--card);
    border:1.5px solid var(--border); border-radius:12px; box-shadow:0 8px 20px rgba(20,20,25,.12);
    z-index:20; overflow:hidden; max-height:240px; overflow-y:auto;
  }
  .plate-suggest .ps-item { padding:10px 14px; cursor:pointer; border-bottom:1px solid var(--border); }
  .plate-suggest .ps-item:last-child { border-bottom:none; }
  .plate-suggest .ps-item:active, .plate-suggest .ps-item:hover { background:var(--field-bg); }
  .plate-suggest .ps-plate { font-family:var(--font-mono); font-weight:700; font-size:14px; letter-spacing:0.5px; color:var(--text); }
  .plate-suggest .ps-owner { font-size:12px; color:var(--hint); margin-top:2px; }
  .checkbox-row { display:flex; align-items:center; gap:8px; }
  .checkbox-row input { width:auto; }
  .item-row { display:flex; gap:8px; align-items:center; margin-bottom:8px; }
  .item-row .item-name { flex:1.3; font-size:13px; color:var(--hint); min-width:0; }
  .item-row input { flex:1; padding:8px; font-size:13px; min-width:0; }
  .item-row select { flex:1.6; padding:8px; font-size:13px; min-width:0; }
  .other-stock-row { display:flex; gap:6px; align-items:center; margin-bottom:8px; }
  .other-stock-row select { flex:2; padding:8px; font-size:13px; min-width:0; }
  .other-stock-row input { flex:1; padding:8px; font-size:13px; min-width:0; }
  .other-stock-row button { flex:none; width:32px; height:32px; border-radius:8px; border:none; background:var(--danger-bg); color:var(--danger); font-size:14px; cursor:pointer; }
  .item-row .item-total { flex:0.9; font-size:12px; color:var(--hint); text-align:right; }
  button.submit {
    width: 100%; padding: 14px; border: none; border-radius: 14px;
    background: linear-gradient(135deg, #1D4ED8 0%, #1E40AF 55%, #312E81 100%); color:#fff; font-size: 15px; font-weight: 700; font-family: var(--font-display);
    letter-spacing:0.5px; text-transform:uppercase; cursor: pointer; margin-top: 6px;
    box-shadow: 0 8px 18px rgba(29,78,216,.30);
  }
  button.submit:active { transform:scale(.98); }
  table { width: 100%; border-collapse: collapse; font-size: 13px; }
  th, td { text-align: left; padding: 8px 6px; border-bottom: 1px solid var(--border); white-space: nowrap; }
  th { color:#fff; font-weight: 700; position: sticky; top: 0; background: #1E293B; font-family: var(--font-body); letter-spacing:0.4px; font-size:10px; text-transform:uppercase; }
  .table-wrap { overflow-x: auto; border:1px solid var(--border); border-radius: 12px; }
  .badge { display:inline-block; padding: 3px 9px; border-radius: 20px; font-size: 11px; font-weight:600; cursor:pointer; border:none; text-transform:uppercase; letter-spacing:0.3px; }
  .badge.linked { background: var(--ok-bg); color: var(--ok); }
  .badge.unlinked { background: var(--danger-bg); color: var(--danger); }
  .search { margin-bottom: 10px; }
  .hint-text { color: var(--hint); font-size: 12px; margin-top: 6px; }
  .msg { padding: 12px; border-radius: 10px; margin-bottom: 10px; font-size: 14px; }
  .msg.ok { background:var(--ok-bg); color:var(--ok); }
  .msg.err { background:var(--danger-bg); color:var(--danger); }
  .modal-overlay { display:none; position:fixed; inset:0; background:rgba(20,20,25,.55); align-items:center; justify-content:center; z-index:50; }
  .modal-overlay.open { display:flex; }
  .modal { background:var(--card); border:1px solid var(--border); border-radius:16px; padding:18px; max-width:320px; width:90%; text-align:center; box-shadow: 0 12px 34px rgba(20,20,25,.25); }
  .modal-wide { max-width:480px; max-height:85vh; overflow-y:auto; }
  .modal img { width:180px; height:180px; margin: 10px auto; display:block; border-radius:10px; background:#fff; }
  .modal .link-text { font-size:12px; word-break:break-all; color:var(--hint); background:var(--field-bg); padding:8px; border-radius:10px; margin-bottom:10px; }
  .modal button { margin-top:8px; }
  .modal a.wa-btn { display:block; text-decoration:none; }
  .close-btn { background:transparent; border:none; color:var(--hint); font-size:14px; cursor:pointer; margin-top:6px; width:100%; padding:8px; }
  .history-toggle { background:transparent; border:none; color:var(--btn); font-size:12px; cursor:pointer; text-decoration:underline; padding:0; }
  .history-entry { padding:6px 0; border-bottom:1px dashed var(--border); font-size:12px; }
  .stats-grid { display:grid; grid-template-columns: repeat(auto-fit, minmax(150px, 1fr)); gap:10px; }
  .stats-card {
    background:linear-gradient(135deg, #EFF6FF, #EEF2FF); border:1px solid #DBEAFE; border-radius:16px; padding:14px;
    text-align:center;
  }
  .stats-card .label { font-size:10px; color:#64748B; margin-bottom:4px; font-weight:700; letter-spacing:0.5px; text-transform:uppercase; }
  .stats-card .amount { font-size:22px; font-weight:800; color:var(--blue); font-family: var(--font-display); }
  .stats-card .count { font-size:11px; color:var(--hint); margin-top:4px; }
  .known-client {
    margin-top:12px; background:var(--card); border:2px solid #BFDBFE; border-radius:18px;
    overflow:hidden; box-shadow: 0 6px 16px rgba(15,82,186,.08);
  }
  .known-client .kc-header {
    background:linear-gradient(90deg, var(--blue), #1a6fd4); color:#fff; padding:12px 14px;
    display:flex; align-items:center; gap:8px;
  }
  .known-client .kc-header i { font-size:15px; }
  .known-client .kc-header span { font-family:var(--font-display); font-weight:700; font-style:italic; font-size:14px; text-transform:uppercase; letter-spacing:0.3px; }
  .known-client .kc-body { padding:16px; }
  .known-client .kc-person { display:flex; align-items:center; gap:14px; margin-bottom:16px; }
  .known-client .kc-avatar {
    width:58px; height:58px; border-radius:50%; background:var(--field-bg); color:var(--blue);
    display:flex; align-items:center; justify-content:center; font-family:var(--font-display);
    font-weight:700; font-size:22px; flex:none; border:2px solid #DBEAFE;
  }
  .known-client .kc-name { font-size:18px; font-weight:700; color:var(--text); line-height:1.25; }
  .known-client .kc-meta { font-size:13px; color:var(--hint); margin-top:3px; }
  .known-client .kc-history-label { font-size:11px; font-weight:700; color:var(--hint); text-transform:uppercase; letter-spacing:0.4px; margin-bottom:8px; }
  .known-client .kc-last-visit {
    background:var(--field-bg); border-radius:12px; padding:11px 13px; margin-bottom:14px;
  }
  .known-client .kc-last-visit .kc-lv-top { display:flex; justify-content:space-between; align-items:center; gap:8px; }
  .known-client .kc-last-visit .kc-lv-service { font-size:14px; font-weight:700; color:var(--text); }
  .known-client .kc-last-visit .kc-lv-date { font-size:12px; color:var(--hint); margin-top:2px; }
  .known-client .kc-lv-mileage { font-size:12px; color:var(--blue); font-weight:600; margin-top:8px; }
  .mileage-compare { font-size:12.5px; margin-top:6px; padding:8px 10px; border-radius:8px; }
  .mileage-compare.over { background:#FEF2F2; color:#B3241C; font-weight:600; }
  .mileage-compare.ok { background:#F0FDF4; color:#15803D; }
  .known-client .kc-last-visit .kc-lv-cost { font-size:16px; font-weight:700; color:var(--blue); flex:none; }
  .known-client .kc-lv-items { margin-top:9px; padding-top:9px; border-top:1px dashed var(--border); }
  .known-client .kc-lv-item {
    display:flex; justify-content:space-between; gap:8px; font-size:12.5px; color:var(--text);
    padding:3px 0;
  }
  .known-client .kc-lv-item span:last-child { color:var(--hint); flex:none; }
  .known-client .kc-action-btn {
    width:100%; padding:13px; border:none; border-radius:12px; background:linear-gradient(135deg, #1D4ED8 0%, #1E40AF 55%, #312E81 100%);
    color:#fff; font-size:15px; font-weight:700; font-family:var(--font-body); cursor:pointer;
    display:flex; align-items:center; justify-content:center; gap:8px; box-shadow:0 6px 14px rgba(29,78,216,.28);
  }
  .known-client .kc-action-hint { font-size:11.5px; color:var(--hint); text-align:center; margin-top:8px; }
  .car-card {
    background:var(--card); border:1.5px solid var(--border); border-radius:14px; padding:14px;
    margin-bottom:10px; cursor:pointer; transition:border-color .12s, box-shadow .12s;
  }
  .car-card:active { border-color:var(--blue); box-shadow:0 2px 10px rgba(15,82,186,.1); }
  .car-card .cc-top { display:flex; justify-content:space-between; align-items:center; margin-bottom:9px; gap:8px; }
  .car-card .cc-plate {
    font-family:var(--font-mono); font-weight:700; font-size:14.5px; letter-spacing:0.5px; color:var(--blue);
    background:var(--field-bg); padding:3px 9px; border-radius:7px; flex:none;
  }
  .car-card .cc-linkbtn { flex:none; }
  .car-card .cc-owner { font-size:15px; font-weight:700; color:var(--text); }
  .car-card .cc-meta { font-size:12.5px; color:var(--hint); margin-top:2px; }
  .car-card .cc-bottom {
    display:flex; justify-content:space-between; align-items:center; margin-top:10px;
    padding-top:10px; border-top:1px dashed var(--border);
  }
  .car-card .cc-cost { font-size:15px; font-weight:700; color:var(--blue); }
  .car-card .cc-cost-date { font-size:11px; color:var(--hint); font-weight:400; }
  .car-card .cc-next { font-size:11.5px; color:var(--hint); text-align:right; }
  .car-card .cc-chevron { color:var(--hint); font-size:13px; flex:none; }
  .debt-card { background:var(--card); border:1.5px solid var(--border); border-radius:14px; padding:14px; margin-bottom:10px; }
  .debt-card.overdue { border-color:#F87171; background:linear-gradient(135deg, #FFFFFF, #FEF2F2); }
  .debt-card .dc-top { display:flex; justify-content:space-between; align-items:flex-start; gap:8px; }
  .debt-card .dc-owner { font-size:15px; font-weight:700; color:var(--text); }
  .debt-card .dc-meta { font-size:12.5px; color:var(--hint); margin-top:2px; }
  .debt-card .dc-remaining { font-size:18px; font-weight:700; color:var(--btn); font-family:var(--font-mono); text-align:right; }
  .debt-card .dc-due { font-size:11.5px; color:var(--hint); margin-top:2px; text-align:right; }
  .debt-card .dc-due.overdue-text { color:#B3241C; font-weight:700; }
  .debt-card .dc-pay-row { display:flex; gap:6px; margin-top:10px; padding-top:10px; border-top:1px dashed var(--border); }
  .dash-row {
    display:flex; justify-content:space-between; align-items:center; padding:8px 0;
    border-bottom:1px dashed var(--border); font-size:13.5px;
  }
  .dash-row:last-child { border-bottom:none; }
  .dash-row .dash-row-rank {
    display:inline-flex; align-items:center; justify-content:center; width:20px; height:20px;
    background:var(--field-bg); border-radius:50%; font-size:11px; font-weight:700; color:var(--hint); margin-right:8px; flex:none;
  }
  .dash-row .dash-row-name { color:var(--text); font-weight:600; }
  .dash-row .dash-row-value { color:var(--blue); font-weight:700; flex:none; }
  .dash-row .dash-row-value.warn { color:#B3241C; }
  .dash-summary-grid { display:grid; grid-template-columns:repeat(2, 1fr); gap:10px; }
  .dash-summary-box { background:var(--field-bg); border-radius:12px; padding:12px; text-align:center; }
  .dash-summary-box .dsb-num { font-size:20px; font-weight:700; color:var(--text); font-family:var(--font-mono); }
  .dash-summary-box .dsb-num.warn { color:#B3241C; }
  .dash-summary-box .dsb-label { font-size:11px; color:var(--hint); margin-top:2px; }
  .dash-brands-panel { background:var(--field-bg); border-radius:10px; margin:4px 0 8px; padding:2px 8px; }
  .recur-exp-card {
    background:var(--field-bg); border-radius:12px; padding:12px; margin-bottom:8px;
  }
  .recur-exp-card.overdue { background:#FEF2F2; }
  .recur-exp-card .rec-top { display:flex; justify-content:space-between; align-items:flex-start; gap:8px; }
  .recur-exp-card .rec-name { font-size:14px; font-weight:700; color:var(--text); }
  .recur-exp-card .rec-meta { font-size:11.5px; color:var(--hint); margin-top:2px; }
  .recur-exp-card .rec-amount { font-size:15px; font-weight:700; color:var(--blue); flex:none; }
  .recur-exp-card .rec-actions { display:flex; gap:6px; margin-top:10px; padding-top:10px; border-top:1px dashed var(--border); }
  .journal-row { display:flex; justify-content:space-between; align-items:center; padding:8px 0; border-bottom:1px dashed var(--border); font-size:13px; }
  .journal-row:last-child { border-bottom:none; }
  .journal-row .jr-cat { font-weight:700; color:var(--text); }
  .journal-row .jr-meta { font-size:11.5px; color:var(--hint); margin-top:2px; }
  .journal-row .jr-amount { color:#B3241C; font-weight:700; flex:none; }
  .jr-icon-btn {
    background:none; border:none; cursor:pointer; font-size:13px; padding:4px; color:var(--hint); flex:none;
  }
  .jr-icon-btn:active { color:var(--blue); }
  .jr-edit-form { background:var(--field-bg); border-radius:10px; padding:12px; margin:2px 0 10px; }
  .kc-hist-list { margin-top:14px; }
  .kc-hist-entry { background:var(--field-bg); border-radius:12px; padding:12px 13px; margin-bottom:8px; }
  .kc-hist-entry .kc-he-top { display:flex; justify-content:space-between; align-items:flex-start; gap:8px; }
  .kc-hist-entry .kc-he-service { font-size:14px; font-weight:700; color:var(--text); }
  .kc-hist-entry .kc-he-date { font-size:11.5px; color:var(--hint); margin-top:2px; }
  .kc-hist-entry .kc-he-cost { font-size:15px; font-weight:700; color:var(--blue); flex:none; white-space:nowrap; }
  .kc-hist-entry .kc-he-meta { font-size:12px; color:var(--hint); margin-top:6px; }
  .kc-hist-entry .kc-he-items { font-size:12px; color:var(--text); margin-top:6px; line-height:1.5; }
  .kc-hist-entry .kc-he-notes { font-size:12px; color:var(--hint); margin-top:6px; }
  .kc-hist-entry .kc-he-actions { margin-top:8px; padding-top:8px; border-top:1px dashed var(--border); }
</style>
</head>
<body>
<div class="speedline"></div>
<div class="container">
  <div class="topbar">
    <div style="display:flex; align-items:center; gap:10px;">
      <div class="logo-badge"><i class="fa-solid fa-wrench"></i></div>
      <div>
        <h1>{{ shop_name }}</h1>
        <div class="logo-sub">MITAL AUTO SERVICE</div>
      </div>
    </div>
    <div style="display:flex; align-items:center; gap:8px;">
      <button class="lang-btn" onclick="switchLanguage()">{{ T.lang_switch }}</button>
      <a class="logout" href="/logout">{{ T.logout }}</a>
    </div>
  </div>

  <div class="tabs">
    <div class="tab active" id="tab-add" onclick="showTab('add')"><span class="tab-icon"><i class="fa-solid fa-oil-can"></i></span><span>{{ T.tab_add }}</span></div>
    <div class="tab" id="tab-table" onclick="showTab('table')"><span class="tab-icon"><i class="fa-solid fa-car"></i></span><span>{{ T.tab_table }}</span></div>
    <div class="tab" id="tab-debts" onclick="showTab('debts')"><span class="tab-icon"><i class="fa-solid fa-hand-holding-dollar"></i></span><span>{{ T.tab_debts }}</span></div>
    {% if not is_employee %}
    <div class="tab" id="tab-expenses" onclick="showTab('expenses')"><span class="tab-icon"><i class="fa-solid fa-receipt"></i></span><span>{{ T.tab_expenses }}</span></div>
    {% endif %}
    <div class="tab" id="tab-broadcast" onclick="showTab('broadcast')"><span class="tab-icon"><i class="fa-solid fa-bullhorn"></i></span><span>{{ T.tab_broadcast }}</span></div>
    {% if not is_employee %}
    <div class="tab" id="tab-export" onclick="showTab('export')"><span class="tab-icon"><i class="fa-solid fa-file-arrow-down"></i></span><span>{{ T.tab_export }}</span></div>
    <div class="tab" id="tab-stats" onclick="showTab('stats')"><span class="tab-icon"><i class="fa-solid fa-chart-column"></i></span><span>{{ T.tab_stats }}</span></div>
    {% endif %}
    {% if sms_enabled %}<div class="tab" id="tab-sms" onclick="showTab('sms')"><span class="tab-icon"><i class="fa-solid fa-comment-sms"></i></span><span>{{ T.tab_sms }}</span></div>{% endif %}
  </div>

  {% if warehouse_enabled and not is_employee %}
  <div class="wh-banner" id="tab-warehouse" onclick="showTab('warehouse')">
    <div class="stripe-pair"><span style="background:var(--blue);"></span><span style="background:var(--btn);"></span></div>
    <div class="wh-label"><i class="fa-solid fa-boxes-stacked"></i> {{ T.tab_warehouse }}</div>
    <div class="stripe-pair"><span style="background:var(--btn);"></span><span style="background:var(--cyan);"></span></div>
  </div>
  {% endif %}

  <div id="msg"></div>

  <div id="view-add" class="card">
    <div class="card-header"><span class="dot"></span><h3>{{ T.tab_add }}</h3></div>
    <div class="field">
      <label><i class="fa-solid fa-id-card"></i>{{ T.field_plate }}</label>
      <div class="plate-wrap">
        <span class="plate-chip">UZ</span>
        <input id="plate" placeholder="01A123BC" oninput="onPlateInput()" onblur="onPlateBlur()" autocomplete="off">
        <div id="plateSuggest" class="plate-suggest" style="display:none;"></div>
      </div>
      <div id="knownClientPanel"></div>
    </div>
    <div class="field">
      <label><i class="fa-solid fa-user"></i>{{ T.field_owner_name }}</label>
      <input id="owner_name" placeholder="Имя Фамилия">
    </div>
    <div class="field">
      <label><i class="fa-solid fa-phone"></i>{{ T.field_owner_phone }}</label>
      <input id="owner_phone" placeholder="+998 90 123 45 67">
      <div class="hint-text">{{ T.hint_owner_phone }}</div>
    </div>
    <div class="row2">
      <div class="field">
        <label><i class="fa-solid fa-car"></i>{{ T.field_car_brand }}</label>
        <select id="car_brand">
          {% for b in brands %}<option value="{{b}}">{{b}}</option>{% endfor %}
        </select>
      </div>
      <div class="field">
        <label><i class="fa-solid fa-car-side"></i>{{ T.field_car_model }}</label>
        <input id="car_model" placeholder="Cobalt, Nexia, Malibu...">
      </div>
    </div>
    <div class="row2">
      <div class="field">
        <label>{{ T.field_mileage }}</label>
        <input id="mileage" type="number" placeholder="45000" oninput="checkMileageVsDue()">
        <div id="mileageCompare"></div>
      </div>
      <div class="field">
        <label>{{ T.field_next_mileage }}</label>
        <input id="next_mileage" type="number" placeholder="55000">
      </div>
    </div>
    <div class="field">
      <label style="font-size:15px; color:var(--text); font-weight:600;">{{ T.section_fluids }}</label>
    </div>
    <div id="fluidsList"></div>

    <div class="field" style="margin-top:14px;">
      <label style="font-size:15px; color:var(--text); font-weight:600;">{{ T.section_filters }}</label>
    </div>
    <div id="filtersList"></div>

    <div class="field" style="margin-top:14px;">
      <label style="font-size:15px; color:var(--text); font-weight:600;">{{ T.section_other }}</label>
    </div>
    <div class="row2">
      <div class="field">
        <label>{{ T.field_other_name }}</label>
        <input id="other_name" placeholder="{{ T.field_other_name_ph }}">
      </div>
      <div class="field">
        <label>{{ T.field_price }}</label>
        <input id="other_price" type="number" placeholder="0" oninput="updateTotal()">
      </div>
    </div>

    {% if warehouse_enabled %}
    <div class="field" style="margin-top:14px;">
      <label style="font-size:15px; color:var(--text); font-weight:600;">{{ T.wh_other_stock_title }}</label>
    </div>
    <div id="otherStockRows"></div>
    <button type="button" class="submit" style="padding:8px; background:var(--border);" onclick="addOtherStockRow('other')">{{ T.wh_add_row }}</button>
    {% endif %}

    <div class="field" style="margin-top:14px; padding:12px; background:var(--field-bg); border-radius:10px;">
      <label style="font-size:15px;">{{ T.field_total }}</label>
      <div id="totalCost" style="font-size:24px; font-weight:600; color:var(--btn); font-family:var(--font-mono);">0</div>
    </div>

    <div class="field" style="margin-top:14px;">
      <label>{{ T.payment_split_label }}</label>
      <div class="row2">
        <div class="field">
          <label style="font-size:11px;"><i class="fa-solid fa-money-bill"></i> {{ T.payment_cash }}</label>
          <input id="pay_cash" type="number" placeholder="0" oninput="onPayCashInput()">
        </div>
        <div class="field">
          <label style="font-size:11px;"><i class="fa-solid fa-credit-card"></i> {{ T.payment_card }}</label>
          <input id="pay_card" type="number" placeholder="0" oninput="onPayCardInput()">
        </div>
      </div>
      <div class="checkbox-row" style="margin-top:10px; cursor:pointer;" onclick="document.getElementById('debt_enabled').click()">
        <input type="checkbox" id="debt_enabled" onchange="toggleDebtSection()" onclick="event.stopPropagation()">
        <label style="margin:0; cursor:pointer;">{{ T.debt_enable_label }}</label>
      </div>
      <div id="debtFields" style="display:none; margin-top:10px; padding:12px; background:var(--card); border:1.5px dashed var(--border); border-radius:10px;">
        <div class="field">
          <label style="font-size:11px;">{{ T.debt_remaining_label }}</label>
          <div id="debtRemaining" style="font-size:18px; font-weight:700; color:var(--btn); font-family:var(--font-mono);">0</div>
        </div>
        <div class="row2">
          <div class="field">
            <label style="font-size:11px;">{{ T.debt_installment_amount }}</label>
            <input id="debt_installment_amount" type="number" placeholder="100000">
          </div>
          <div class="field">
            <label style="font-size:11px;">{{ T.debt_interval_days }}</label>
            <input id="debt_interval_days" type="number" placeholder="14">
          </div>
        </div>
      </div>
    </div>

    <div class="field" style="margin-top:14px;">
      <label>{{ T.field_interval }}</label>
      <div style="display:flex; gap:8px;">
        <input id="interval_value" type="number" placeholder="3" value="3" style="flex:1;">
        <select id="interval_unit" style="flex:1;">
          <option value="months">{{ T.unit_months }}</option>
          <option value="days">{{ T.unit_days }}</option>
        </select>
      </div>
    </div>
    <div class="field">
      <label>{{ T.field_notes }}</label>
      <textarea id="notes" rows="2" placeholder="{{ T.notes_ph }}"></textarea>
    </div>
    <button class="submit" onclick="submitCar()">{{ T.btn_save }}</button>
  </div>

  <div id="view-table" style="display:none;">
    <input class="search" id="search" placeholder="{{ T.search_ph }}" oninput="renderTable()">
    <div id="clientCardPanel" style="display:none; margin-bottom:12px;"></div>
    <div id="table-body"></div>
  </div>

  <div id="view-debts" style="display:none;">
    <div id="debtsList"></div>
  </div>

  {% if not is_employee %}
  <div id="view-expenses" style="display:none;">
    {% if not is_branch %}
    <div class="card" id="usdRateCardExp" style="margin-bottom:14px;">
      <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.usd_rate_title }}</label>
      <div class="row2">
        <div class="field">
          <label>{{ T.usd_rate_label }}</label>
          <input id="usd_rate_input_exp" type="number" step="0.01" placeholder="12700" value="{{ usd_rate or '' }}">
        </div>
      </div>
      <button class="submit" onclick="saveUsdRate('usd_rate_input_exp', 'usdRateSaved_exp')">{{ T.usd_rate_save }}</button>
      <div id="usdRateSaved_exp" style="display:none; color:#1B8A5A; font-size:13px; margin-top:8px;">✓ {{ T.usd_rate_saved }}</div>
    </div>
    {% endif %}

    <div class="card">
      <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.expense_add_title }}</label>
      <div class="field">
        <label>{{ T.expense_category_label }}</label>
        <select id="exp_category" onchange="onExpenseCategoryChange('exp_category', 'exp_category_custom')"></select>
        <input id="exp_category_custom" placeholder="{{ T.expense_custom_category_ph }}" style="display:none; margin-top:6px;">
      </div>
      <div class="field">
        <label>{{ T.expense_name_label }}</label>
        <input id="exp_name" placeholder="{{ T.expense_name_ph }}">
      </div>
      <div class="row2">
        <div class="field">
          <label>{{ T.expense_amount_label }}</label>
          <input id="exp_amount" type="number" placeholder="100000" oninput="onSumFieldEdited('exp_amount', 'exp_amount_usd')">
        </div>
        <div class="field">
          <label>{{ T.usd_price_label }}</label>
          <input id="exp_amount_usd" type="number" step="0.01" placeholder="$" oninput="onUsdFieldEdited('exp_amount_usd', 'exp_amount')">
        </div>
      </div>
      <button class="submit" onclick="submitExpense()">{{ T.expense_add_btn }}</button>
    </div>

    <div class="card" style="margin-top:14px;">
      <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.expense_recurring_title }}</label>
      <div id="recurringExpensesList"></div>
      <div style="margin-top:10px; padding-top:10px; border-top:1px dashed var(--border);">
        <div class="field">
          <label>{{ T.expense_category_label }}</label>
          <select id="rec_category" onchange="onExpenseCategoryChange('rec_category', 'rec_category_custom')"></select>
          <input id="rec_category_custom" placeholder="{{ T.expense_custom_category_ph }}" style="display:none; margin-top:6px;">
        </div>
        <div class="field">
          <label>{{ T.expense_name_label }}</label>
          <input id="rec_name" placeholder="{{ T.expense_recurring_name_ph }}">
        </div>
        <div class="row2">
          <div class="field">
            <label>{{ T.expense_amount_label }}</label>
            <input id="rec_amount" type="number" placeholder="2000000" oninput="onSumFieldEdited('rec_amount', 'rec_amount_usd')">
          </div>
          <div class="field">
            <label>{{ T.expense_day_of_month_label }}</label>
            <input id="rec_day" type="number" min="1" max="28" placeholder="5">
          </div>
        </div>
        <div class="field">
          <label>{{ T.usd_price_label }}</label>
          <input id="rec_amount_usd" type="number" step="0.01" placeholder="$" oninput="onUsdFieldEdited('rec_amount_usd', 'rec_amount')">
        </div>
        <button class="submit" onclick="createRecurringExpense()">{{ T.expense_recurring_add_btn }}</button>
      </div>
    </div>

    <div class="card" style="margin-top:14px;">
      <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.expense_journal_title }}</label>
      <div id="expensesJournal">{{ T.stats_loading }}</div>
    </div>
  </div>
  {% endif %}

  <div id="view-broadcast" class="card" style="display:none;">
    <div class="field">
      <label>{{ T.broadcast_msg_label }}</label>
      <textarea id="broadcast_message" rows="5" placeholder="{{ T.broadcast_msg_ph }}"></textarea>
      <div class="hint-text" id="broadcastRecipients">{{ T.broadcast_loading_recipients }}</div>
    </div>
    <button class="submit" onclick="sendBroadcast()">{{ T.broadcast_send_btn }}</button>
    <div style="margin-top:18px;">
      <div class="hint-text" style="margin-bottom:8px;">{{ T.broadcast_history_title }}</div>
      <div id="broadcastHistory"></div>
    </div>
  </div>

  {% if not is_employee %}
  <div id="view-export" class="card" style="display:none;">
    <p style="margin-top:0;">{{ T.export_p1 }}</p>
    <p class="hint-text">{{ T.export_p2 }}</p>
    <button class="submit" onclick="window.location.href='/api/export'">{{ T.export_btn }}</button>
    <button class="submit" style="background:#1a7a3d; margin-top:8px;" onclick="window.location.href='/api/export_excel'">{{ T.export_excel_btn }}</button>
    <p class="hint-text">{{ T.export_excel_hint }}</p>
  </div>

  <div id="view-stats" style="display:none;">
    <div class="subtabs">
      <div class="subtab active" id="substat-own" onclick="showStatsSubTab('own')">{{ T.stats_sub_own }}</div>
      <div class="subtab" id="substat-branches" onclick="showStatsSubTab('branches')" style="display:none;">{{ T.stats_sub_branches }}</div>
    </div>

    <div id="statsOwnView">
      <div id="statsGrid" class="stats-grid">{{ T.stats_loading }}</div>

      {% if not is_employee %}
      <div class="card" style="margin-top:16px; background:linear-gradient(135deg, #F0FDF4, #ECFDF5); border-color:#86EFAC;">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.dash_net_profit_title }}</label>
        <div id="dashNetProfit">{{ T.stats_loading }}</div>
      </div>
      {% endif %}

      <div class="card" style="margin-top:16px;">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.dash_revenue_chart_title }}</label>
        <canvas id="revenueChart" height="180"></canvas>
      </div>

      <div class="card" style="margin-top:16px;">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.dash_top_products_title }}</label>
        <div id="dashTopProducts">{{ T.stats_loading }}</div>
      </div>

      <div class="card" style="margin-top:16px;">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.dash_debt_summary_title }}</label>
        <div id="dashDebtSummary">{{ T.stats_loading }}</div>
      </div>

      {% if warehouse_enabled and not is_employee %}
      <div class="card" style="margin-top:16px;">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.dash_low_stock_title }}</label>
        <div id="dashLowStock">{{ T.stats_loading }}</div>
      </div>
      {% endif %}

      <div class="card" style="margin-top:16px;">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.stats_custom_title }}</label>
        <div id="statsPresets" style="display:flex; flex-wrap:wrap; gap:8px; margin-bottom:14px;"></div>
        <div class="row2">
          <div class="field">
            <label>{{ T.stats_from }}</label>
            <input id="stats_from" type="date">
          </div>
          <div class="field">
            <label>{{ T.stats_to }}</label>
            <input id="stats_to" type="date">
          </div>
        </div>
        <button class="submit" onclick="applyStatsRange()">{{ T.stats_apply }}</button>
        <div id="statsRangeResult" style="margin-top:14px;"></div>
      </div>
    </div>

    <div id="statsBranchesView" style="display:none;">
      <div id="statsAggregated"></div>

      <div class="card" style="margin-top:16px;">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.stats_custom_title }}</label>
        <div id="branchStatsPresets" style="display:flex; flex-wrap:wrap; gap:8px; margin-bottom:14px;"></div>
        <div class="row2">
          <div class="field">
            <label>{{ T.stats_from }}</label>
            <input id="branch_stats_from" type="date">
          </div>
          <div class="field">
            <label>{{ T.stats_to }}</label>
            <input id="branch_stats_to" type="date">
          </div>
        </div>
        <button class="submit" onclick="applyBranchStatsRange()">{{ T.stats_apply }}</button>
        <div id="branchStatsRangeResult" style="margin-top:14px;"></div>
        <div id="branchChartWrap" style="margin-top:16px; display:none;">
          <canvas id="branchChart" height="220"></canvas>
        </div>
      </div>
    </div>
  </div>
  {% endif %}

  {% if sms_enabled %}
  <div id="view-sms" class="card" style="display:none;">
    <p style="margin-top:0;">{{ T.sms_intro }}</p>
    <div class="field">
      <label>{{ T.sms_email }}</label>
      <input id="eskiz_email" value="{{ eskiz_email }}" placeholder="you@example.com">
    </div>
    <div class="field">
      <label>{{ T.sms_password }}</label>
      <input id="eskiz_password" type="password" placeholder="{{ T.sms_password_ph }}">
      <div class="hint-text">{{ T.sms_password_hint }}</div>
    </div>
    <button class="submit" onclick="saveSmsSettings()">{{ T.btn_save }}</button>
  </div>
  {% endif %}

  {% if warehouse_enabled and not is_employee %}
  <div id="view-warehouse" style="display:none;">
    <div class="subtabs">
      <div class="subtab active" id="subwh-own" onclick="showWhSubTab('own')">{{ T.wh_sub_own }}</div>
      <div class="subtab" id="subwh-branches" onclick="showWhSubTab('branches')" style="display:none;">{{ T.wh_sub_branches }}</div>
    </div>

    <div id="whOwnView">
      {% if not is_branch %}
      <div class="card" id="usdRateCard">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.usd_rate_title }}</label>
        <div class="row2">
          <div class="field">
            <label>{{ T.usd_rate_label }}</label>
            <input id="usd_rate_input" type="number" step="0.01" placeholder="12700" value="{{ usd_rate or '' }}">
          </div>
        </div>
        <button class="submit" onclick="saveUsdRate()">{{ T.usd_rate_save }}</button>
        <div id="usdRateSaved" style="display:none; color:#1B8A5A; font-size:13px; margin-top:8px;">✓ {{ T.usd_rate_saved }}</div>
        <p class="hint-text" style="margin-top:10px; margin-bottom:0;">{{ T.usd_rate_hint }}</p>
      </div>
      {% endif %}

      <div class="card">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.wh_add_product }}</label>
        <div class="row2">
          <div class="field">
            <label>{{ T.wh_category }}</label>
            <select id="wh_new_category" onchange="onWhCategoryChanged()"></select>
          </div>
          <div class="field">
            <label>{{ T.wh_product_name }}</label>
            <input id="wh_new_name" placeholder="MITANOL 5W-30">
          </div>
        </div>
        <div class="field" id="wh_new_unit_row" style="display:none;">
          <label>{{ T.wh_unit }}</label>
          <select id="wh_new_unit">
            <option value="pc">{{ T.unit_pc }}</option>
            <option value="l">{{ T.unit_l }}</option>
          </select>
        </div>
        <div class="row2">
          <div class="field">
            <label>{{ T.wh_sell_price }}</label>
            <input id="wh_new_sell_price" type="number" placeholder="45000">
          </div>
          <div class="field" {% if is_branch %}style="display:none;"{% endif %}>
            <label>{{ T.wh_purchase_price }}</label>
            <input id="wh_new_purchase_price" type="number" placeholder="30000" oninput="onSumFieldEdited('wh_new_purchase_price', 'wh_new_purchase_usd')">
          </div>
        </div>
        <div class="field" {% if is_branch %}style="display:none;"{% endif %}>
          <label>{{ T.usd_price_label }}</label>
          <input id="wh_new_purchase_usd" type="number" step="0.01" placeholder="$" oninput="onUsdFieldEdited('wh_new_purchase_usd', 'wh_new_purchase_price')">
        </div>
        <div class="field">
          <label>{{ T.wh_initial_stock }}</label>
          <input id="wh_new_stock" type="number" placeholder="0">
        </div>
        <button class="submit" onclick="createProduct()">{{ T.wh_add_btn }}</button>
      </div>

      <div class="card" style="margin-top:14px;">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.wh_products_title }}</label>
        {% if not is_branch %}
        <div id="whStockValue" class="hint-text" style="margin-bottom:10px;"></div>
        {% endif %}
        <div class="table-wrap" style="overflow-x:auto;">
          <table>
            <thead><tr>
              <th>{{ T.wh_category }}</th><th>{{ T.wh_product_name }}</th><th>{{ T.wh_stock }}</th>
              <th>{{ T.wh_sell_price }}</th>{% if not is_branch %}<th>{{ T.wh_purchase_price }}</th>{% endif %}<th></th>
            </tr></thead>
            <tbody id="products-body"></tbody>
          </table>
        </div>
      </div>

      <div class="card" style="margin-top:14px;">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.wh_restock_history }}</label>
        <div id="restockHistory"></div>
      </div>
    </div>

    <div id="whBranchesView" style="display:none;">
      <div class="card">
        <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">{{ T.branch_prices_title }}</label>
        <div id="branchWarehouseSummary" style="margin-bottom:12px;"></div>
        <select id="branchPriceSelect" onchange="loadBranchProducts(this.value)">
          <option value="">{{ T.branch_prices_pick }}</option>
        </select>
        <div id="branchProductsPanel" style="margin-top:10px;"></div>
      </div>
    </div>
  </div>
  {% endif %}
</div>
"""

MODAL_AND_SCRIPT = """
<div class="modal-overlay" id="restockModal">
  <div class="modal" style="text-align:left;">
    <h3 style="text-align:center;" id="restockModalTitle">{{ T.wh_restock_title }}</h3>
    <div class="field">
      <label>{{ T.wh_restock_qty }}</label>
      <input id="restock_qty" type="number">
    </div>
    <div class="field" {% if is_branch %}style="display:none;"{% endif %}>
      <label>{{ T.wh_purchase_price }}</label>
      <input id="restock_price" type="number" oninput="onSumFieldEdited('restock_price', 'restock_price_usd')">
    </div>
    <div class="field" {% if is_branch %}style="display:none;"{% endif %}>
      <label>{{ T.usd_price_label }}</label>
      <input id="restock_price_usd" type="number" step="0.01" placeholder="$" oninput="onUsdFieldEdited('restock_price_usd', 'restock_price')">
    </div>
    <div class="field">
      <label>{{ T.wh_restock_date }}</label>
      <input id="restock_date" type="date">
    </div>
    <button class="submit" onclick="submitRestock()">{{ T.wh_restock_btn }}</button>
    <button class="close-btn" onclick="closeRestockModal()">{{ T.modal_close }}</button>
  </div>
</div>

<div class="modal-overlay" id="linkModal">
  <div class="modal">
    <h3 id="modalPlate"></h3>
    <div class="hint-text">{{ T.modal_hint }}</div>
    <img id="modalQr" alt="QR">
    <div class="link-text" id="modalLink"></div>
    <a class="wa-btn" id="modalTg" href="#" target="_blank"><button class="submit" type="button" style="background:#2AABEE;">{{ T.modal_send_tg }}</button></a>
    <a class="wa-btn" id="modalWa" href="#" target="_blank"><button class="submit" type="button">{{ T.modal_send_wa }}</button></a>
    <button class="submit" onclick="copyLink()" style="background:var(--border);">{{ T.modal_copy }}</button>
    <button class="close-btn" onclick="closeModal()">{{ T.modal_close }}</button>
  </div>
</div>

<div class="modal-overlay" id="editModal">
  <div class="modal modal-wide" style="text-align:left;">
    <h3 style="text-align:center;" id="svcModalTitle">{{ T.entry_edit_title }}</h3>
    <div class="field">
      <label>{{ T.field_mileage }}</label>
      <input id="edit_mileage" type="number">
    </div>
    <div class="field">
      <label>{{ T.field_next_mileage }}</label>
      <input id="edit_next_mileage" type="number">
    </div>

    <div class="field">
      <label style="font-size:15px; color:var(--text); font-weight:600;">{{ T.section_fluids }}</label>
    </div>
    <div id="svcFluidsList"></div>

    <div class="field" style="margin-top:10px;">
      <label style="font-size:15px; color:var(--text); font-weight:600;">{{ T.section_filters }}</label>
    </div>
    <div id="svcFiltersList"></div>

    <div class="field" style="margin-top:10px;">
      <label style="font-size:15px; color:var(--text); font-weight:600;">{{ T.section_other }}</label>
    </div>
    <div class="row2">
      <div class="field">
        <label>{{ T.field_other_name }}</label>
        <input id="svc_other_name" placeholder="{{ T.field_other_name_ph }}">
      </div>
      <div class="field">
        <label>{{ T.field_price }}</label>
        <input id="svc_other_price" type="number" placeholder="0" oninput="updateSvcTotal()">
      </div>
    </div>

    {% if warehouse_enabled %}
    <div class="field" style="margin-top:10px;">
      <label style="font-size:15px; color:var(--text); font-weight:600;">{{ T.wh_other_stock_title }}</label>
    </div>
    <div id="svcOtherStockRows"></div>
    <button type="button" class="submit" style="padding:8px; background:var(--border);" onclick="addOtherStockRow('svcOther')">{{ T.wh_add_row }}</button>
    {% endif %}

    <div class="field" style="margin-top:10px; padding:12px; background:var(--field-bg); border-radius:10px;">
      <label style="font-size:15px;">{{ T.field_total }}</label>
      <div id="svcTotalCost" style="font-size:22px; font-weight:600; color:var(--btn); font-family:var(--font-mono);">0</div>
    </div>

    <div class="field" style="margin-top:10px;">
      <label>{{ T.payment_split_label }}</label>
      <div class="row2">
        <div class="field">
          <label style="font-size:11px;"><i class="fa-solid fa-money-bill"></i> {{ T.payment_cash }}</label>
          <input id="svc_pay_cash" type="number" placeholder="0" oninput="onSvcPayCashInput()">
        </div>
        <div class="field">
          <label style="font-size:11px;"><i class="fa-solid fa-credit-card"></i> {{ T.payment_card }}</label>
          <input id="svc_pay_card" type="number" placeholder="0" oninput="onSvcPayCardInput()">
        </div>
      </div>
    </div>

    <div class="field" style="margin-top:10px;">
      <label>{{ T.field_interval }}</label>
      <div style="display:flex; gap:8px;">
        <input id="edit_interval_value" type="number" style="flex:1;">
        <select id="edit_interval_unit" style="flex:1;">
          <option value="months">{{ T.unit_months }}</option>
          <option value="days">{{ T.unit_days }}</option>
        </select>
      </div>
    </div>
    <div class="field">
      <label>{{ T.field_notes }}</label>
      <textarea id="edit_notes" rows="2"></textarea>
    </div>
    <button class="submit" onclick="saveEdit()">{{ T.btn_save }}</button>
    <button class="close-btn" onclick="closeEditModal()">{{ T.modal_close }}</button>
  </div>
</div>

<script>
const T = {{ t_json|safe }};
const LANG = {{ lang|tojson }};
const WAREHOUSE_ENABLED = {{ warehouse_enabled|tojson }};
const IS_BRANCH = {{ is_branch|tojson }};
let USD_RATE = {{ usd_rate|tojson }};
const tg = window.Telegram ? window.Telegram.WebApp : null;
if (tg) { tg.ready(); tg.expand(); }

let carsCache = [];
let lastKnownNextMileage = null;

function checkMileageVsDue() {
  const el = document.getElementById('mileageCompare');
  if (!el) return;
  const entered = parseFloat(document.getElementById('mileage').value);
  if (!lastKnownNextMileage || isNaN(entered)) { el.innerHTML = ''; return; }
  const diff = entered - lastKnownNextMileage;
  if (diff > 0) {
    el.innerHTML = `<div class="mileage-compare over">⚠️ ${T.mileage_over_due} ${lastKnownNextMileage.toLocaleString('ru-RU')} ${T.km_short} — ${T.mileage_over_by} ${diff.toLocaleString('ru-RU')} ${T.km_short}</div>`;
  } else {
    el.innerHTML = `<div class="mileage-compare ok">${T.mileage_within_due} ${lastKnownNextMileage.toLocaleString('ru-RU')} ${T.km_short}</div>`;
  }
}

let openHistoryRow = null;
let historyDataCache = {};  // { plate: [entry, entry, ...] } — чтобы кнопки не тащили сырые данные записи (с заметками, апострофами и т.п.) прямо в HTML-атрибут onclick, а брали их отсюда по id

function showTab(t) {
  document.getElementById('view-add').style.display = t === 'add' ? 'block' : 'none';
  document.getElementById('view-table').style.display = t === 'table' ? 'block' : 'none';
  document.getElementById('view-debts').style.display = t === 'debts' ? 'block' : 'none';
  document.getElementById('view-broadcast').style.display = t === 'broadcast' ? 'block' : 'none';
  document.getElementById('tab-add').classList.toggle('active', t === 'add');
  document.getElementById('tab-table').classList.toggle('active', t === 'table');
  document.getElementById('tab-debts').classList.toggle('active', t === 'debts');
  document.getElementById('tab-broadcast').classList.toggle('active', t === 'broadcast');
  const exportView = document.getElementById('view-export');
  const exportTab = document.getElementById('tab-export');
  if (exportView) exportView.style.display = t === 'export' ? 'block' : 'none';
  if (exportTab) exportTab.classList.toggle('active', t === 'export');
  const statsView = document.getElementById('view-stats');
  const statsTab = document.getElementById('tab-stats');
  if (statsView) statsView.style.display = t === 'stats' ? 'block' : 'none';
  if (statsTab) statsTab.classList.toggle('active', t === 'stats');
  const smsView = document.getElementById('view-sms');
  const smsTab = document.getElementById('tab-sms');
  if (smsView) smsView.style.display = t === 'sms' ? 'block' : 'none';
  if (smsTab) smsTab.classList.toggle('active', t === 'sms');
  const whView = document.getElementById('view-warehouse');
  const whTab = document.getElementById('tab-warehouse');
  if (whView) whView.style.display = t === 'warehouse' ? 'block' : 'none';
  if (whTab) whTab.classList.toggle('active', t === 'warehouse');
  const expView = document.getElementById('view-expenses');
  const expTab = document.getElementById('tab-expenses');
  if (expView) expView.style.display = t === 'expenses' ? 'block' : 'none';
  if (expTab) expTab.classList.toggle('active', t === 'expenses');
  if (t === 'table') loadCars();
  if (t === 'debts') loadDebts();
  if (t === 'expenses') loadExpensesTab();
  if (t === 'broadcast') loadBroadcastInfo();
  if (t === 'stats') loadStats();
  if (t === 'warehouse') loadWarehouse();
}

async function saveSmsSettings() {
  const emailEl = document.getElementById('eskiz_email');
  const passwordEl = document.getElementById('eskiz_password');
  if (!emailEl || !passwordEl) return;  // вкладка SMS не отрисована (выключена для этой точки)
  const email = emailEl.value.trim();
  const password = passwordEl.value;
  const res = await fetch('/api/sms_settings', {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({email, password})
  });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.sms_saved, true);
    document.getElementById('eskiz_password').value = '';
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

// ---- Склад ----
let productsCache = [];
let restockingProductId = null;

function renderWarehouseCategoryOptions() {
  const sel = document.getElementById('wh_new_category');
  if (!sel) return;
  const allKeys = FLUID_KEYS.concat(FILTER_KEYS);
  sel.innerHTML = allKeys.map(key => `<option value="${key}">${T[key]}</option>`).join('')
    + `<option value="other">${T.wh_category_other}</option>`;
}

function onWhCategoryChanged() {
  const sel = document.getElementById('wh_new_category');
  const unitRow = document.getElementById('wh_new_unit_row');
  if (unitRow) unitRow.style.display = sel.value === 'other' ? 'block' : 'none';
}

async function loadWarehouse() {
  renderWarehouseCategoryOptions();
  const res = await fetch('/api/products');
  productsCache = await res.json();
  renderProductsTable();
  loadRestockHistory();
  // склад мог поменяться (добавили/удалили/пополнили товар) — обновляем поля
  // "марка" в форме "Внести замену" и в форме редактирования/добавления из
  // истории, чтобы новый товар сразу стал доступен в списке, без
  // перезагрузки всей страницы
  renderItemLists();
  renderSvcItemLists();

  if (!IS_BRANCH) {
    try {
      const branches = await (await fetch('/api/my_branches')).json();
      const branchBtn = document.getElementById('subwh-branches');
      if (branches.length) {
        branchBtn.style.display = '';
        document.getElementById('branchPriceSelect').innerHTML =
          `<option value="">${T.branch_prices_pick}</option>` +
          branches.map(b => `<option value="${b.id}">${escapeHtml(b.shop_name || b.username)}</option>`).join('');
        document.getElementById('branchWarehouseSummary').innerHTML = branches.map(b => `
          <div style="padding:6px 0; font-size:12px; border-bottom:1px dashed var(--border);">
            <div style="display:flex; justify-content:space-between;">
              <span>${escapeHtml(b.shop_name || b.username)}</span>
              <span>${T.branch_products_count} ${b.product_count}${b.missing_price_count > 0 ? ` · <span style="color:#B3241C;">⚠️ ${T.branch_missing_price} ${b.missing_price_count}</span>` : ''}</span>
            </div>
            <div style="color:#9A3412; font-weight:600; margin-top:2px;">${T.branch_stock_value} ${b.stock_value.toLocaleString('ru-RU')} ${T.currency}</div>
          </div>
        `).join('');
      } else {
        branchBtn.style.display = 'none';
      }
    } catch (e) { /* не главный аккаунт или ошибка - просто не показываем подвкладку */ }
  }
}

function showWhSubTab(t) {
  document.getElementById('whOwnView').style.display = t === 'own' ? 'block' : 'none';
  document.getElementById('whBranchesView').style.display = t === 'branches' ? 'block' : 'none';
  document.getElementById('subwh-own').classList.toggle('active', t === 'own');
  document.getElementById('subwh-branches').classList.toggle('active', t === 'branches');
}


function renderProductsTable() {
  const body = document.getElementById('products-body');
  if (!body) return;
  const colCount = IS_BRANCH ? 5 : 6;

  const valueEl = document.getElementById('whStockValue');
  if (valueEl) {
    const totalValue = productsCache.reduce((sum, p) =>
      sum + (p.purchase_price != null ? p.stock_qty * p.purchase_price : 0), 0);
    const missingCount = productsCache.filter(p => p.purchase_price == null).length;
    valueEl.innerHTML = `${T.wh_stock_value_label} <b style="color:#9A3412;">${totalValue.toLocaleString('ru-RU')} ${T.currency}</b>` +
      (missingCount > 0 ? ` <span style="color:#B3241C;">(⚠️ ${T.branch_missing_price} ${missingCount})</span>` : '');
  }

  if (!productsCache.length) {
    body.innerHTML = `<tr><td colspan="${colCount}">${T.wh_no_products}</td></tr>`;
    return;
  }
  body.innerHTML = productsCache.map(p => {
    const isLow = p.stock_qty < 0;
    const unitLabel = p.unit === 'pc' ? T.unit_pc : T.unit_l;
    return `
    <tr>
      <td>${T[p.category] || p.category}</td>
      <td>${escapeHtml(p.name)}</td>
      <td style="${isLow ? 'color:#B3241C; font-weight:700;' : ''}">${isLow ? '⚠️ ' : ''}${p.stock_qty} ${unitLabel}</td>
      <td>${p.sell_price ? p.sell_price.toLocaleString('ru-RU') + ' ' + T.currency : '—'}</td>
      ${IS_BRANCH ? '' : `<td>${p.purchase_price ? p.purchase_price.toLocaleString('ru-RU') + ' ' + T.currency : '—'}</td>`}
      <td>
        <button class="history-toggle" onclick="openRestockModal(${p.id}, ${escapeHtml(JSON.stringify(p.name))})">${T.wh_restock_action}</button>
        &nbsp;·&nbsp;
        <button class="history-toggle" style="color:#B3241C;" onclick="deleteProduct(${p.id}, ${escapeHtml(JSON.stringify(p.name))})">${T.wh_delete_action}</button>
      </td>
    </tr>
  `;
  }).join('');
}

function onUsdFieldEdited(usdFieldId, sumFieldId) {
  const usdVal = parseFloat(document.getElementById(usdFieldId).value);
  const sumField = document.getElementById(sumFieldId);
  if (!USD_RATE || isNaN(usdVal)) return;
  sumField.value = Math.round(usdVal * USD_RATE);
}

function onSumFieldEdited(sumFieldId, usdFieldId) {
  // если сумма меняется вручную (не через пересчёт из $), поле $ больше не
  // соответствует новой сумме однозначно - просто очищаем его, чтобы не
  // вводить в заблуждение несоответствующим числом
  const usdField = document.getElementById(usdFieldId);
  if (usdField) usdField.value = '';
}

async function saveUsdRate(inputId, savedId) {
  inputId = inputId || 'usd_rate_input';
  savedId = savedId || 'usdRateSaved';
  const rateInput = document.getElementById(inputId);
  const rate = rateInput.value;
  const res = await fetch('/api/usd_rate', {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({rate})
  });
  const data = await res.json();
  if (data.ok) {
    USD_RATE = data.rate;
    // синхронизируем оба виджета курса, если на странице есть второй (Расходы)
    ['usd_rate_input', 'usd_rate_input_exp'].forEach(id => {
      const el = document.getElementById(id);
      if (el && id !== inputId) el.value = data.rate ?? '';
    });
    const saved = document.getElementById(savedId);
    saved.style.display = 'block';
    setTimeout(() => { saved.style.display = 'none'; }, 1500);
  } else {
    showMsg(T.usd_rate_error || data.error, false);
  }
}

async function createProduct() {
  const catEl = document.getElementById('wh_new_category');
  const nameEl = document.getElementById('wh_new_name');
  if (!catEl || !nameEl) return;  // вкладка "Склад" не отрисована (выключена для этой точки)
  const category = catEl.value;
  const name = nameEl.value.trim();
  if (!category || !name) { showMsg(T.wh_fill_required, false); return; }
  const payload = {
    category, name,
    unit: document.getElementById('wh_new_unit').value,
    sell_price: document.getElementById('wh_new_sell_price').value || null,
    purchase_price: document.getElementById('wh_new_purchase_price').value || null,
    initial_stock: document.getElementById('wh_new_stock').value || 0,
  };
  const res = await fetch('/api/products', {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify(payload)
  });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.wh_product_added, true);
    ['wh_new_name','wh_new_sell_price','wh_new_purchase_price','wh_new_purchase_usd','wh_new_stock'].forEach(id => document.getElementById(id).value = '');
    document.getElementById('wh_new_unit_row').style.display = 'none';
    loadWarehouse();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

async function deleteProduct(id, name) {
  if (!confirm(T.wh_delete_confirm)) return;
  const res = await fetch('/api/products/' + id, { method: 'DELETE' });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.wh_product_deleted, true);
    loadWarehouse();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

function openRestockModal(id, name) {
  restockingProductId = id;
  document.getElementById('restockModalTitle').textContent = T.wh_restock_title + ': ' + name;
  document.getElementById('restock_qty').value = '';
  document.getElementById('restock_price').value = '';
  const restockUsd = document.getElementById('restock_price_usd');
  if (restockUsd) restockUsd.value = '';
  document.getElementById('restock_date').value = fmtDate(new Date());
  document.getElementById('restockModal').classList.add('open');
}

function closeRestockModal() {
  document.getElementById('restockModal').classList.remove('open');
}

async function submitRestock() {
  const payload = {
    quantity: document.getElementById('restock_qty').value,
    purchase_price: document.getElementById('restock_price').value || null,
    restock_date: document.getElementById('restock_date').value || null,
  };
  const res = await fetch(`/api/products/${restockingProductId}/restock`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify(payload)
  });
  const data = await res.json();
  if (data.ok) {
    closeRestockModal();
    showMsg(T.wh_restocked, true);
    loadWarehouse();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

async function loadRestockHistory() {
  const el = document.getElementById('restockHistory');
  if (!el) return;
  const res = await fetch('/api/restock_history');
  const history = await res.json();
  el.innerHTML = history.length ? history.map(r => `
    <div style="padding:8px 0;border-bottom:1px dashed var(--border);font-size:13px;">
      ${r.restock_date} — ${escapeHtml(r.product_name)}: +${r.quantity} ${r.unit === 'pc' ? T.unit_pc : T.unit_l}
      ${r.purchase_price ? ' (' + r.purchase_price.toLocaleString('ru-RU') + ' ' + T.currency + '/ед.)' : ''}
    </div>
  `).join('') : `<div class="hint-text">${T.wh_no_restocks}</div>`;
}


let revenueChartInstance = null;

function renderRevenueChart(dailyData) {
  const canvas = document.getElementById('revenueChart');
  if (!canvas || typeof Chart === 'undefined') return;
  if (revenueChartInstance) { revenueChartInstance.destroy(); }
  revenueChartInstance = new Chart(canvas.getContext('2d'), {
    type: 'line',
    data: {
      labels: dailyData.map(d => d.date.slice(5)),
      datasets: [{
        label: T.dash_revenue_chart_title,
        data: dailyData.map(d => d.total),
        borderColor: '#E63946',
        backgroundColor: 'rgba(230,57,70,0.08)',
        fill: true, tension: 0.3, pointRadius: 0,
      }]
    },
    options: { responsive: true, plugins: { legend: { display: false } }, scales: { y: { beginAtZero: true } } }
  });
}

async function toggleCategoryBrands(categoryName, idx) {
  const panel = document.getElementById('catBrands_' + idx);
  if (!panel) return;
  const isOpen = panel.style.display !== 'none';
  document.querySelectorAll('.dash-brands-panel').forEach(el => { if (el !== panel) el.style.display = 'none'; });
  if (isOpen) { panel.style.display = 'none'; return; }
  panel.innerHTML = T.stats_loading;
  panel.style.display = 'block';
  const brands = await (await fetch('/api/dashboard/top_brands?category=' + encodeURIComponent(categoryName))).json();
  panel.innerHTML = brands.length ? brands.map((b, j) => `
    <div class="dash-row" style="padding-left:28px; font-size:12.5px;">
      <span class="dash-row-name">${j + 1}. ${escapeHtml(b.name)}</span>
      <span class="dash-row-value">${b.qty.toLocaleString('ru-RU')}</span>
    </div>
  `).join('') : `<div class="hint-text" style="padding-left:28px;">${T.dash_no_data}</div>`;
}

async function loadDashboard() {
  const res = await fetch('/api/dashboard');
  const data = await res.json();

  renderRevenueChart(data.daily_revenue);

  const netEl = document.getElementById('dashNetProfit');
  if (netEl && data.net_profit) {
    const np = data.net_profit;
    const isNegative = np.net_profit < 0;
    netEl.innerHTML = `
      <div class="dash-summary-grid">
        <div class="dash-summary-box">
          <div class="dsb-num">${np.oil_profit.toLocaleString('ru-RU')} ${T.currency}</div>
          <div class="dsb-label">${T.dash_oil_profit_label}</div>
        </div>
        <div class="dash-summary-box">
          <div class="dsb-num warn">${np.expenses_total.toLocaleString('ru-RU')} ${T.currency}</div>
          <div class="dsb-label">${T.dash_expenses_label}</div>
        </div>
      </div>
      <div style="text-align:center; margin-top:12px; padding-top:12px; border-top:1px dashed #86EFAC;">
        <div style="font-size:24px; font-weight:700; font-family:var(--font-mono); color:${isNegative ? '#B3241C' : '#15803D'};">${np.net_profit.toLocaleString('ru-RU')} ${T.currency}</div>
        <div style="font-size:11px; color:var(--hint); margin-top:2px;">${T.dash_net_profit_label}</div>
      </div>
    `;
  }

  const topEl = document.getElementById('dashTopProducts');
  if (topEl) {
    topEl.innerHTML = data.top_products.length ? data.top_products.map((p, i) => `
      <div>
        <div class="dash-row" style="cursor:pointer;" onclick="toggleCategoryBrands(${escapeHtml(JSON.stringify(p.name))}, ${i})">
          <span><span class="dash-row-rank">${i + 1}</span><span class="dash-row-name">${escapeHtml(p.name)}</span> <i class="fa-solid fa-chevron-down" style="font-size:10px; color:var(--hint); margin-left:4px;"></i></span>
          <span class="dash-row-value">${p.qty.toLocaleString('ru-RU')}</span>
        </div>
        <div id="catBrands_${i}" class="dash-brands-panel" style="display:none;"></div>
      </div>
    `).join('') : `<div class="hint-text">${T.dash_no_data}</div>`;
  }

  const debtEl = document.getElementById('dashDebtSummary');
  if (debtEl) {
    const ds = data.debt_summary;
    debtEl.innerHTML = `
      <div class="dash-summary-grid">
        <div class="dash-summary-box">
          <div class="dsb-num">${ds.total_remaining.toLocaleString('ru-RU')} ${T.currency}</div>
          <div class="dsb-label">${T.dash_total_owed}</div>
        </div>
        <div class="dash-summary-box">
          <div class="dsb-num ${ds.overdue_count > 0 ? 'warn' : ''}">${ds.overdue_count} / ${ds.count}</div>
          <div class="dsb-label">${T.dash_overdue_of_total}</div>
        </div>
      </div>
    `;
  }

  const lowEl = document.getElementById('dashLowStock');
  if (lowEl) {
    lowEl.innerHTML = data.low_stock.length ? data.low_stock.map(p => `
      <div class="dash-row">
        <span class="dash-row-name">${escapeHtml(p.name)}</span>
        <span class="dash-row-value warn">${p.stock_qty} ${p.unit === 'pc' ? T.unit_pc : T.unit_l}</span>
      </div>
    `).join('') : `<div class="hint-text">${T.dash_no_data}</div>`;
  }
}

async function loadStats() {
  loadDashboard();
  const res = await fetch('/api/stats');
  const s = await res.json();
  let profit = null;
  if (WAREHOUSE_ENABLED && !IS_BRANCH) {
    const pRes = await fetch('/api/profit_stats');
    profit = await pRes.json();
  }
  const periods = [
    ['today', T.stats_today], ['yesterday', T.stats_yesterday], ['week', T.stats_week],
    ['month', T.stats_month], ['year', T.stats_year],
  ];

  document.getElementById('statsGrid').innerHTML = periods.map(([key, label]) => `
    <div class="stats-card">
      <div class="label">${label}</div>
      <div class="amount">${s[key].total.toLocaleString('ru-RU')} ${T.currency}</div>
      <div class="count">${T.stats_services_count} ${s[key].count}</div>
      <div class="count">${T.payment_cash}: ${(s[key].cash || 0).toLocaleString('ru-RU')} ${T.currency} · ${T.payment_card}: ${(s[key].card || 0).toLocaleString('ru-RU')} ${T.currency}</div>
      ${profit ? `<div class="count" style="color:#1B8A5A;">${T.stats_profit_label} ${profit[key].toLocaleString('ru-RU')} ${T.currency}</div>` : ''}
    </div>
  `).join('');

  if (!IS_BRANCH) {
    try {
      const agg = await (await fetch('/api/aggregated_stats')).json();
      const branchBtn = document.getElementById('substat-branches');
      if (agg.has_branches) {
        branchBtn.style.display = '';
        const breakdownRows = (agg.breakdown || []).map(r => `
          <tr>
            <td>${escapeHtml(r.shop_name)}${r.is_head ? ` <span class="hint-text">(${T.branch_head_label})</span>` : ''}</td>
            <td>${r.revenue.today.total.toLocaleString('ru-RU')} ${T.currency}</td>
            <td style="color:#1B8A5A;">${r.profit.today.toLocaleString('ru-RU')} ${T.currency}</td>
          </tr>
        `).join('');
        document.getElementById('statsAggregated').innerHTML = `
          <div class="card" style="border-color:#FDBA74;">
            <label style="font-size:15px; color:var(--text); font-weight:600; display:block; margin-bottom:10px;">
              ${T.stats_all_branches_title} (${agg.branch_count})
            </label>
            <div class="stats-grid">
              ${periods.map(([key, label]) => `
                <div class="stats-card" style="background:linear-gradient(135deg, #FFF7ED, #FEF3C7); border-color:#FDBA74;">
                  <div class="label">${label}</div>
                  <div class="amount" style="color:#9A3412;">${agg.revenue[key].total.toLocaleString('ru-RU')} ${T.currency}</div>
                  <div class="count">${T.stats_services_count} ${agg.revenue[key].count}</div>
                  <div class="count">${T.payment_cash}: ${(agg.revenue[key].cash || 0).toLocaleString('ru-RU')} ${T.currency} · ${T.payment_card}: ${(agg.revenue[key].card || 0).toLocaleString('ru-RU')} ${T.currency}</div>
                  <div class="count" style="color:#1B8A5A;">${T.stats_profit_label} ${agg.profit[key].toLocaleString('ru-RU')} ${T.currency}</div>
                </div>
              `).join('')}
            </div>
            <div style="margin-top:14px; padding-top:14px; border-top:1px dashed #FDBA74;">
              <label style="font-size:13px; font-weight:600; display:block; margin-bottom:8px;">${T.branch_breakdown_title}</label>
              <div class="table-wrap"><table>
                <thead><tr><th>${T.branch_col_label}</th><th>${T.stats_today}</th><th>${T.stats_profit_label}</th></tr></thead>
                <tbody>${breakdownRows}</tbody>
              </table></div>
            </div>
          </div>
        `;
      } else {
        branchBtn.style.display = 'none';
      }
    } catch (e) { /* не главный аккаунт или ошибка - просто не показываем подвкладку */ }
  }
}

function showStatsSubTab(t) {
  document.getElementById('statsOwnView').style.display = t === 'own' ? 'block' : 'none';
  document.getElementById('statsBranchesView').style.display = t === 'branches' ? 'block' : 'none';
  document.getElementById('substat-own').classList.toggle('active', t === 'own');
  document.getElementById('substat-branches').classList.toggle('active', t === 'branches');
}

async function loadBranchProducts(branchId) {
  const panel = document.getElementById('branchProductsPanel');
  if (!branchId) { panel.innerHTML = ''; return; }
  const products = await (await fetch(`/api/branches/${branchId}/products`)).json();
  if (!products.length) { panel.innerHTML = `<div class="hint-text">${T.wh_no_products}</div>`; return; }
  panel.innerHTML = products.map(p => `
    <div style="padding:6px 0; border-bottom:1px dashed var(--border); font-size:13px;">
      <div>${escapeHtml(p.name)} <span class="hint-text">(${p.stock_qty} ${p.unit === 'pc' ? T.unit_pc : T.unit_l})</span></div>
      <div style="display:flex; align-items:center; gap:6px; margin-top:4px; flex-wrap:wrap;">
        <input id="branch_price_${branchId}_${p.id}" type="number" value="${p.purchase_price ?? ''}" placeholder="${T.wh_purchase_price}"
               style="width:100px; padding:5px 8px; font-size:12px; flex:none;"
               oninput="onSumFieldEdited('branch_price_${branchId}_${p.id}', 'branch_price_usd_${branchId}_${p.id}')">
        <input id="branch_price_usd_${branchId}_${p.id}" type="number" step="0.01" placeholder="$"
               style="width:80px; padding:5px 8px; font-size:12px; flex:none;"
               oninput="onUsdFieldEdited('branch_price_usd_${branchId}_${p.id}', 'branch_price_${branchId}_${p.id}')">
        <button class="badge active" style="flex:none;" onclick="setBranchPurchasePrice(${branchId}, ${p.id})">${T.branch_save_price}</button>
        <span id="branch_price_saved_${branchId}_${p.id}" style="display:none; color:#1B8A5A; font-size:15px; flex:none;">✓</span>
      </div>
    </div>
  `).join('');
}

async function setBranchPurchasePrice(branchId, productId) {
  const input = document.getElementById(`branch_price_${branchId}_${productId}`);
  await fetch(`/api/branches/${branchId}/products/${productId}/purchase_price`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({purchase_price: input.value})
  });
  const check = document.getElementById(`branch_price_saved_${branchId}_${productId}`);
  if (check) {
    check.style.display = 'inline';
    setTimeout(() => { check.style.display = 'none'; }, 1500);
  }
}

function fmtDate(d) {
  return d.getFullYear() + '-' + String(d.getMonth() + 1).padStart(2, '0') + '-' + String(d.getDate()).padStart(2, '0');
}

function setStatsRange(from, to) {
  document.getElementById('stats_from').value = fmtDate(from);
  document.getElementById('stats_to').value = fmtDate(to);
  applyStatsRange();
}

function renderStatsPresets() {
  const presetsEl = document.getElementById('statsPresets');
  if (!presetsEl) return;  // вкладка "Статистика" скрыта (например, для сотрудника)
  const today = new Date();
  const presets = [
    [T.stats_preset_today, () => setStatsRange(today, today)],
    [T.stats_preset_yesterday, () => {
      const y = new Date(today); y.setDate(y.getDate() - 1);
      setStatsRange(y, y);
    }],
    [T.stats_preset_7days, () => {
      const start = new Date(today); start.setDate(start.getDate() - 6);
      setStatsRange(start, today);
    }],
    [T.stats_preset_this_month, () => {
      const start = new Date(today.getFullYear(), today.getMonth(), 1);
      setStatsRange(start, today);
    }],
    [T.stats_preset_prev_month, () => {
      const start = new Date(today.getFullYear(), today.getMonth() - 1, 1);
      const end = new Date(today.getFullYear(), today.getMonth(), 0);
      setStatsRange(start, end);
    }],
  ];
  document.getElementById('statsPresets').innerHTML = presets.map(([label], i) =>
    `<button class="lang-btn" onclick="STATS_PRESETS[${i}]()">${label}</button>`
  ).join('');
  window.STATS_PRESETS = presets.map(p => p[1]);
  const from = document.getElementById('stats_from');
  const to = document.getElementById('stats_to');
  if (!from.value) from.value = fmtDate(today);
  if (!to.value) to.value = fmtDate(today);
}

async function applyStatsRange() {
  const from = document.getElementById('stats_from').value;
  const to = document.getElementById('stats_to').value;
  if (!from || !to) return;
  const res = await fetch(`/api/stats/range?from=${from}&to=${to}`);
  const data = await res.json();
  if (!data.ok) { document.getElementById('statsRangeResult').innerHTML = ''; return; }
  document.getElementById('statsRangeResult').innerHTML = `
    <div class="stats-card">
      <div class="label">${T.stats_range_result} ${from} — ${to}</div>
      <div class="amount">${data.total.toLocaleString('ru-RU')} ${T.currency}</div>
      <div class="count">${T.stats_services_count} ${data.count}</div>
      <div class="count">${T.payment_cash}: ${(data.cash || 0).toLocaleString('ru-RU')} ${T.currency} · ${T.payment_card}: ${(data.card || 0).toLocaleString('ru-RU')} ${T.currency}</div>
    </div>
  `;
}

renderStatsPresets();

function renderBranchStatsPresets() {
  const presetsEl = document.getElementById('branchStatsPresets');
  if (!presetsEl) return;  // подвкладка "Все филиалы" скрыта (нет филиалов) или недоступна
  const today = new Date();
  const presets = [
    [T.stats_preset_today, () => setBranchStatsRange(today, today)],
    [T.stats_preset_yesterday, () => {
      const y = new Date(today); y.setDate(y.getDate() - 1);
      setBranchStatsRange(y, y);
    }],
    [T.stats_preset_7days, () => {
      const start = new Date(today); start.setDate(start.getDate() - 6);
      setBranchStatsRange(start, today);
    }],
    [T.stats_preset_this_month, () => {
      const start = new Date(today.getFullYear(), today.getMonth(), 1);
      setBranchStatsRange(start, today);
    }],
    [T.stats_preset_prev_month, () => {
      const start = new Date(today.getFullYear(), today.getMonth() - 1, 1);
      const end = new Date(today.getFullYear(), today.getMonth(), 0);
      setBranchStatsRange(start, end);
    }],
  ];
  presetsEl.innerHTML = presets.map(([label], i) =>
    `<button class="lang-btn" onclick="BRANCH_STATS_PRESETS[${i}]()">${label}</button>`
  ).join('');
  window.BRANCH_STATS_PRESETS = presets.map(p => p[1]);
  const from = document.getElementById('branch_stats_from');
  const to = document.getElementById('branch_stats_to');
  if (from && !from.value) from.value = fmtDate(today);
  if (to && !to.value) to.value = fmtDate(today);
}

function setBranchStatsRange(from, to) {
  document.getElementById('branch_stats_from').value = fmtDate(from);
  document.getElementById('branch_stats_to').value = fmtDate(to);
  applyBranchStatsRange();
}

async function applyBranchStatsRange() {
  const from = document.getElementById('branch_stats_from').value;
  const to = document.getElementById('branch_stats_to').value;
  if (!from || !to) return;
  const res = await fetch(`/api/aggregated_stats/range?from=${from}&to=${to}`);
  const data = await res.json();
  if (!data.ok) { document.getElementById('branchStatsRangeResult').innerHTML = ''; return; }

  const breakdown = data.breakdown || [];
  const breakdownRows = breakdown.map(r => `
    <tr>
      <td>${escapeHtml(r.shop_name)}${r.is_head ? ` <span class="hint-text">(${T.branch_head_label})</span>` : ''}</td>
      <td>${r.revenue.total.toLocaleString('ru-RU')} ${T.currency}</td>
      <td style="color:#1B8A5A;">${r.profit.toLocaleString('ru-RU')} ${T.currency}</td>
    </tr>
  `).join('');

  document.getElementById('branchStatsRangeResult').innerHTML = `
    <div class="stats-card" style="background:linear-gradient(135deg, #FFF7ED, #FEF3C7); border-color:#FDBA74; margin-bottom:14px;">
      <div class="label">${T.stats_range_result} ${from} — ${to}</div>
      <div class="amount" style="color:#9A3412;">${data.revenue.total.toLocaleString('ru-RU')} ${T.currency}</div>
      <div class="count">${T.stats_services_count} ${data.revenue.count}</div>
      <div class="count">${T.payment_cash}: ${(data.revenue.cash || 0).toLocaleString('ru-RU')} ${T.currency} · ${T.payment_card}: ${(data.revenue.card || 0).toLocaleString('ru-RU')} ${T.currency}</div>
      <div class="count" style="color:#1B8A5A;">${T.stats_profit_label} ${data.profit.toLocaleString('ru-RU')} ${T.currency}</div>
    </div>
    <div class="table-wrap"><table>
      <thead><tr><th>${T.branch_col_label}</th><th>${T.stats_revenue_label}</th><th>${T.stats_profit_label}</th></tr></thead>
      <tbody>${breakdownRows}</tbody>
    </table></div>
  `;

  renderBranchChart(breakdown);
}

function renderBranchChart(breakdown) {
  const wrap = document.getElementById('branchChartWrap');
  if (!wrap) return;
  if (typeof Chart === 'undefined' || !breakdown.length) { wrap.style.display = 'none'; return; }
  wrap.style.display = 'block';
  const canvas = document.getElementById('branchChart');
  if (window.branchChartInstance) { window.branchChartInstance.destroy(); }
  window.branchChartInstance = new Chart(canvas.getContext('2d'), {
    type: 'bar',
    data: {
      labels: breakdown.map(r => r.shop_name),
      datasets: [
        { label: T.stats_revenue_label, data: breakdown.map(r => r.revenue.total), backgroundColor: '#E63946' },
        { label: T.branch_chart_profit_label, data: breakdown.map(r => r.profit), backgroundColor: '#1B8A5A' },
      ]
    },
    options: { responsive: true, plugins: { legend: { position: 'bottom' } }, scales: { y: { beginAtZero: true } } }
  });
}

renderBranchStatsPresets();

async function switchLanguage() {
  const newLang = LANG === 'ru' ? 'uz' : 'ru';
  await fetch('/api/set_language', {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({language: newLang})
  });
  window.location.reload();
}

async function loadBroadcastInfo() {
  const res = await fetch('/api/broadcast/recipients');
  const data = await res.json();
  document.getElementById('broadcastRecipients').textContent =
    `${T.broadcast_recipients_text} ${data.count} ${T.broadcast_recipients_suffix}`;

  const res2 = await fetch('/api/broadcast/history');
  const items = await res2.json();
  document.getElementById('broadcastHistory').innerHTML = items.length ? items.map(b => `
    <div style="padding:8px 0;border-bottom:1px dashed var(--border);font-size:13px;">
      <div>${b.message.length > 80 ? b.message.slice(0,80) + '…' : b.message}</div>
      <div class="hint-text">${b.created_at} — ${T.broadcast_status_label} ${b.status}${b.status === 'done' ? `, ${T.broadcast_delivered} ${b.total_sent}, ${T.broadcast_failed} ${b.total_failed}` : ''}</div>
    </div>
  `).join('') : `<div class="hint-text">${T.broadcast_none_yet}</div>`;
}

async function sendBroadcast() {
  const message = document.getElementById('broadcast_message').value.trim();
  if (!message) { showMsg(T.broadcast_empty_msg, false); return; }
  if (!confirm(T.broadcast_confirm)) return;
  const res = await fetch('/api/broadcast', {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({message})
  });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.broadcast_queued, true);
    document.getElementById('broadcast_message').value = '';
    setTimeout(loadBroadcastInfo, 4000);
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

function showMsg(text, ok) {
  const el = document.getElementById('msg');
  el.innerHTML = `<div class="msg ${ok ? 'ok' : 'err'}">${text}</div>`;
  setTimeout(() => { el.innerHTML = ''; }, 6000);
}

// ---- Жидкости и фильтры (детализация замены) ----
const FLUID_KEYS = ["fluid_0", "fluid_1", "fluid_2", "fluid_3", "fluid_4"];
const FILTER_KEYS = ["filter_0", "filter_1", "filter_2", "filter_3"];

// Общая логика для полей "марка" в форме внесения/редактирования замены —
// используется и в обычной форме "Внести замену", и в форме
// добавления/редактирования из истории (svcModal). Если для категории есть
// товары на складе — поле становится выпадающим списком с остатком и ценой,
// иначе (или если склад не включён точке) — обычное текстовое поле, как и
// раньше.
function productsForCategory(key) {
  return WAREHOUSE_ENABLED ? productsCache.filter(p => p.category === key) : [];
}

function brandFieldHtml(key, fieldId, onPickHandler) {
  const prods = productsForCategory(key);
  if (!prods.length) {
    return `<input id="${fieldId}" placeholder="${T.brand_ph}">`;
  }
  const options = ['<option value="">' + T.brand_ph + '</option>'].concat(
    prods.map(p => {
      const unitLabel = p.unit === 'pc' ? T.unit_pc : T.unit_l;
      const stockWarn = p.stock_qty < 0 ? ' ⚠️' : '';
      return `<option value="${p.id}" data-price="${p.sell_price || ''}" data-name="${escapeHtml(p.name)}">${escapeHtml(p.name)} (${p.stock_qty} ${unitLabel}${stockWarn})</option>`;
    })
  ).join('');
  return `<select id="${fieldId}" data-is-product="1" onchange="${onPickHandler}">${options}</select>`;
}

function readBrandField(fieldId) {
  const el = document.getElementById(fieldId);
  if (!el) return { brand: null, product_id: null };
  if (el.tagName === 'SELECT') {
    if (!el.value) return { brand: null, product_id: null };
    const opt = el.options[el.selectedIndex];
    return { brand: opt.dataset.name || null, product_id: parseInt(el.value) };
  }
  return { brand: el.value.trim() || null, product_id: null };
}

function selectOrPreserveBrand(brandEl, it) {
  // Подставляет сохранённую позицию в поле "марка" при открытии редактирования.
  // Если это обычное текстовое поле — просто вписываем текст, как раньше.
  // Если это выпадающий список (склад включён и в категории есть товары) —
  // пытаемся выбрать ТОТ ЖЕ товар. Но если товар с тех пор удалили со склада
  // (его больше нет среди опций) — список НЕ должен молча показывать пустое
  // поле: тогда добавляем в список одноразовую "историческую" опцию с
  // сохранённым названием, чтобы при обычном пересохранении (без изменения
  // этого поля) марка не терялась.
  if (brandEl.tagName === 'INPUT') {
    brandEl.value = it.brand || '';
    return;
  }
  if (brandEl.tagName !== 'SELECT') return;
  if (it.product_id && [...brandEl.options].some(o => o.value === String(it.product_id))) {
    brandEl.value = String(it.product_id);
  } else if (it.brand) {
    const phantom = document.createElement('option');
    phantom.value = it.product_id || ('legacy_' + Math.random().toString(36).slice(2));
    phantom.textContent = it.brand + ' (' + T.wh_product_unavailable + ')';
    phantom.dataset.name = it.brand;
    if (!it.product_id) phantom.dataset.legacy = '1';
    brandEl.appendChild(phantom);
    brandEl.value = phantom.value;
  }
}

// ---- "Прочие товары" — динамический список строк (как в чеке), для товаров
// со склада, которые не подходят ни под одну из фиксированных категорий
// (жидкости/фильтры). prefix: 'other' — основная форма "Внести замену",
// 'svcOther' — форма редактирования/добавления из истории. ----
let otherStockRows = [];
let otherStockSeq = 0;
let svcOtherStockRows = [];
let svcOtherStockSeq = 0;

function renderOneOtherStockRow(prefix, id) {
  const prods = productsForCategory('other');
  const updateFn = prefix === 'other' ? 'updateTotal' : 'updateSvcTotal';
  const options = ['<option value="">' + T.brand_ph + '</option>'].concat(
    prods.map(p => {
      const unitLabel = p.unit === 'pc' ? T.unit_pc : T.unit_l;
      const stockWarn = p.stock_qty < 0 ? ' ⚠️' : '';
      return `<option value="${p.id}" data-price="${p.sell_price || ''}" data-name="${escapeHtml(p.name)}">${escapeHtml(p.name)} (${p.stock_qty} ${unitLabel}${stockWarn})</option>`;
    })
  ).join('');
  return `
    <div class="other-stock-row" id="${prefix}_stockrow_${id}">
      <select id="${prefix}_stock_product_${id}" onchange="onOtherStockPicked('${prefix}', ${id})">${options}</select>
      <input id="${prefix}_stock_price_${id}" type="number" placeholder="${T.price_ph}" oninput="${updateFn}()">
      <input id="${prefix}_stock_qty_${id}" type="number" step="0.1" placeholder="${T.qty_ph}" oninput="${updateFn}()">
      <button type="button" onclick="removeOtherStockRow('${prefix}', ${id})">✕</button>
    </div>
  `;
}

function addOtherStockRow(prefix) {
  let id;
  if (prefix === 'other') { otherStockSeq++; id = otherStockSeq; otherStockRows.push(id); }
  else { svcOtherStockSeq++; id = svcOtherStockSeq; svcOtherStockRows.push(id); }
  const container = document.getElementById(prefix + 'StockRows');
  if (!container) return;
  // добавляем ТОЛЬКО новую строку в конец — не трогаем уже существующие,
  // чтобы не стереть то, что пользователь уже успел ввести в них
  container.insertAdjacentHTML('beforeend', renderOneOtherStockRow(prefix, id));
}

function removeOtherStockRow(prefix, id) {
  if (prefix === 'other') { otherStockRows = otherStockRows.filter(x => x !== id); }
  else { svcOtherStockRows = svcOtherStockRows.filter(x => x !== id); }
  // удаляем ТОЛЬКО конкретный элемент строки — соседние строки не трогаем
  const rowEl = document.getElementById(`${prefix}_stockrow_${id}`);
  if (rowEl) rowEl.remove();
  if (prefix === 'other') updateTotal(); else updateSvcTotal();
}

function renderOtherStockRows(prefix) {
  // Полная перерисовка С НУЛЯ — используется только при сбросе формы или
  // загрузке сохранённых позиций (когда в DOM ещё нет строк, которые можно
  // было бы случайно затереть). Для добавления/удаления ОДНОЙ строки к уже
  // заполненному списку используются addOtherStockRow/removeOtherStockRow —
  // они не перерисовывают соседние строки.
  const container = document.getElementById(prefix + 'StockRows');
  if (!container) return;
  const ids = prefix === 'other' ? otherStockRows : svcOtherStockRows;
  container.innerHTML = ids.map(id => renderOneOtherStockRow(prefix, id)).join('');
}

function onOtherStockPicked(prefix, id) {
  const el = document.getElementById(`${prefix}_stock_product_${id}`);
  const opt = el.options[el.selectedIndex];
  if (opt && opt.dataset.price) document.getElementById(`${prefix}_stock_price_${id}`).value = opt.dataset.price;
  if (prefix === 'other') updateTotal(); else updateSvcTotal();
}

function collectOtherStockItems(prefix) {
  const ids = prefix === 'other' ? otherStockRows : svcOtherStockRows;
  const items = [];
  ids.forEach(id => {
    const selectEl = document.getElementById(`${prefix}_stock_product_${id}`);
    const priceEl = document.getElementById(`${prefix}_stock_price_${id}`);
    const qtyEl = document.getElementById(`${prefix}_stock_qty_${id}`);
    if (!selectEl || !priceEl || !qtyEl) return;
    const price = parseFloat(priceEl.value) || 0;
    const qty = parseFloat(qtyEl.value) || 0;
    if (price > 0 && qty > 0 && selectEl.value) {
      const opt = selectEl.options[selectEl.selectedIndex];
      items.push({
        key: 'other_stock', name: opt.dataset.name || '', brand: null, product_id: parseInt(selectEl.value),
        unit_price: price, qty, total: Math.round(price * qty),
      });
    }
  });
  return items;
}

function fillOtherStockRowsFrom(prefix, items) {
  // Восстанавливает СРАЗУ ВСЕ позиции "Прочее" при открытии редактирования.
  // Важно: сначала регистрируем id всех строк и рисуем их ОДНИМ вызовом
  // renderOtherStockRows — если рисовать по одной (вызывая render на каждую),
  // каждый следующий вызов полностью перерисовывает контейнер и стирает
  // выбор, уже сделанный в предыдущих строках.
  if (!items.length) return;
  const ids = items.map(() => {
    if (prefix === 'other') { otherStockSeq++; otherStockRows.push(otherStockSeq); return otherStockSeq; }
    svcOtherStockSeq++; svcOtherStockRows.push(svcOtherStockSeq); return svcOtherStockSeq;
  });
  renderOtherStockRows(prefix);
  items.forEach((it, i) => {
    const id = ids[i];
    const selectEl = document.getElementById(`${prefix}_stock_product_${id}`);
    if (selectEl) selectOrPreserveBrand(selectEl, { product_id: it.product_id, brand: it.name });
    document.getElementById(`${prefix}_stock_price_${id}`).value = it.unit_price ?? '';
    document.getElementById(`${prefix}_stock_qty_${id}`).value = it.qty ?? '';
  });
}

function renderItemLists() {
  document.getElementById('fluidsList').innerHTML = FLUID_KEYS.map((key, i) => `
    <div class="item-row">
      <span class="item-name">${T[key]}</span>
      ${brandFieldHtml(key, `fluid_brand_${i}`, `onFluidProductPicked(${i})`)}
      <input id="fluid_price_${i}" type="number" placeholder="${T.price_per_liter_ph}" oninput="updateTotal()">
      <input id="fluid_liters_${i}" type="number" step="0.1" placeholder="${T.liters_ph}" oninput="updateTotal()">
    </div>
  `).join('');
  document.getElementById('filtersList').innerHTML = FILTER_KEYS.map((key, i) => {
    const prods = productsForCategory(key);
    const brandField = prods.length ? brandFieldHtml(key, `filter_brand_${i}`, `onFilterProductPicked(${i})`) : '';
    return `
    <div class="item-row">
      <span class="item-name" style="flex:${prods.length ? '1.3' : '2.3'};">${T[key]}</span>
      ${brandField}
      <input id="filter_price_${i}" type="number" placeholder="${T.price_ph}" oninput="updateTotal()">
    </div>
  `;
  }).join('');
}

function onFluidProductPicked(i) {
  const el = document.getElementById(`fluid_brand_${i}`);
  const opt = el.options[el.selectedIndex];
  if (opt && opt.dataset.price) document.getElementById(`fluid_price_${i}`).value = opt.dataset.price;
  updateTotal();
}

function onFilterProductPicked(i) {
  const el = document.getElementById(`filter_brand_${i}`);
  const opt = el.options[el.selectedIndex];
  if (opt && opt.dataset.price) document.getElementById(`filter_price_${i}`).value = opt.dataset.price;
  updateTotal();
}

function collectItems() {
  const items = [];
  FLUID_KEYS.forEach((key, i) => {
    const price = parseFloat(document.getElementById(`fluid_price_${i}`).value) || 0;
    const liters = parseFloat(document.getElementById(`fluid_liters_${i}`).value) || 0;
    if (price > 0 && liters > 0) {
      const { brand, product_id } = readBrandField(`fluid_brand_${i}`);
      items.push({
        key, name: T[key], brand, product_id,
        unit_price: price, qty: liters, total: Math.round(price * liters),
      });
    }
  });
  FILTER_KEYS.forEach((key, i) => {
    const price = parseFloat(document.getElementById(`filter_price_${i}`).value) || 0;
    if (price > 0) {
      const brandEl = document.getElementById(`filter_brand_${i}`);
      const { brand, product_id } = brandEl ? readBrandField(`filter_brand_${i}`) : { brand: null, product_id: null };
      items.push({key, name: T[key], brand, product_id, unit_price: price, qty: 1, total: Math.round(price)});
    }
  });
  const otherName = document.getElementById('other_name').value.trim();
  const otherPrice = parseFloat(document.getElementById('other_price').value) || 0;
  if (otherPrice > 0) {
    items.push({key: 'other', name: `${T.other_prefix}: ${otherName || T.other_unnamed}`, unit_price: otherPrice, qty: 1, total: Math.round(otherPrice)});
  }
  items.push(...collectOtherStockItems('other'));
  return items;
}

let paymentSplitTouched = false;

function updateTotal() {
  const total = collectItems().reduce((sum, i) => sum + i.total, 0);
  document.getElementById('totalCost').textContent = total.toLocaleString('ru-RU') + ' ' + T.currency;
  const payCash = document.getElementById('pay_cash');
  const payCard = document.getElementById('pay_card');
  if (payCash && payCard && !paymentSplitTouched) {
    payCash.value = total || '';
    payCard.value = '';
  }
  updateDebtRemaining();
}

function onPayCashInput() {
  paymentSplitTouched = true;
  const debtEnabled = document.getElementById('debt_enabled').checked;
  if (!debtEnabled) {
    const total = collectItems().reduce((sum, i) => sum + i.total, 0);
    const cash = parseFloat(document.getElementById('pay_cash').value) || 0;
    document.getElementById('pay_card').value = Math.max(0, Math.round(total - cash));
  }
  updateDebtRemaining();
}

function onPayCardInput() {
  paymentSplitTouched = true;
  const debtEnabled = document.getElementById('debt_enabled').checked;
  if (!debtEnabled) {
    const total = collectItems().reduce((sum, i) => sum + i.total, 0);
    const card = parseFloat(document.getElementById('pay_card').value) || 0;
    document.getElementById('pay_cash').value = Math.max(0, Math.round(total - card));
  }
  updateDebtRemaining();
}

function updateDebtRemaining() {
  const debtEl = document.getElementById('debtRemaining');
  if (!debtEl) return;
  const total = collectItems().reduce((sum, i) => sum + i.total, 0);
  const cash = parseFloat(document.getElementById('pay_cash').value) || 0;
  const card = parseFloat(document.getElementById('pay_card').value) || 0;
  const remaining = Math.max(0, Math.round(total - cash - card));
  debtEl.textContent = remaining.toLocaleString('ru-RU') + ' ' + T.currency;
}

function toggleDebtSection() {
  const enabled = document.getElementById('debt_enabled').checked;
  document.getElementById('debtFields').style.display = enabled ? 'block' : 'none';
  if (!enabled) {
    // выключили - возвращаем обычное поведение (наличные/карта снова = вся сумма)
    paymentSplitTouched = false;
    updateTotal();
    document.getElementById('debt_installment_amount').value = '';
    document.getElementById('debt_interval_days').value = '';
  } else {
    updateDebtRemaining();
  }
}

function resetItemInputs() {
  FLUID_KEYS.forEach((_, i) => {
    document.getElementById(`fluid_brand_${i}`).value = '';
    document.getElementById(`fluid_price_${i}`).value = '';
    document.getElementById(`fluid_liters_${i}`).value = '';
  });
  FILTER_KEYS.forEach((_, i) => document.getElementById(`filter_price_${i}`).value = '');
  document.getElementById('other_name').value = '';
  document.getElementById('other_price').value = '';
  document.getElementById('knownClientPanel').innerHTML = '';
  otherStockRows = [];
  renderOtherStockRows('other');
  updateTotal();
}

let plateSuggestLoading = false;

async function ensureCarsCacheLoaded() {
  if (carsCache.length || plateSuggestLoading) return;
  plateSuggestLoading = true;
  try {
    const res = await fetch('/api/cars');
    carsCache = await res.json();
  } catch (e) { /* тихо — просто не будет подсказок в этот раз */ }
  plateSuggestLoading = false;
}

async function onPlateInput() {
  const val = document.getElementById('plate').value.trim().toUpperCase();
  const dropdown = document.getElementById('plateSuggest');
  if (!val) { dropdown.style.display = 'none'; dropdown.innerHTML = ''; return; }
  await ensureCarsCacheLoaded();
  const matches = carsCache.filter(c => (c.plate_number || '').toUpperCase().includes(val)).slice(0, 8);
  if (!matches.length) { dropdown.style.display = 'none'; dropdown.innerHTML = ''; return; }
  dropdown.innerHTML = matches.map(c => `
    <div class="ps-item" onmousedown="selectPlateSuggestion(${escapeHtml(JSON.stringify(c.plate_number))})">
      <div class="ps-plate">${escapeHtml(c.plate_number)}</div>
      <div class="ps-owner">${escapeHtml(c.owner_name || '')}${c.car_brand ? ' · ' + escapeHtml(c.car_brand) + (c.car_model ? ' ' + escapeHtml(c.car_model) : '') : ''}</div>
    </div>
  `).join('');
  dropdown.style.display = '';
}

function selectPlateSuggestion(plate) {
  document.getElementById('plate').value = plate;
  document.getElementById('plateSuggest').style.display = 'none';
  lookupPlate();
}

function onPlateBlur() {
  setTimeout(() => {
    const dropdown = document.getElementById('plateSuggest');
    if (dropdown) dropdown.style.display = 'none';
  }, 150);
  lookupPlate();
}

async function lookupPlate() {
  const plate = document.getElementById('plate').value.trim();
  const panel = document.getElementById('knownClientPanel');
  if (!plate) { panel.innerHTML = ''; lastKnownNextMileage = null; checkMileageVsDue(); return; }
  try {
    const res = await fetch('/api/history/' + encodeURIComponent(plate));
    const data = await res.json();
    if (!data.car) { panel.innerHTML = ''; lastKnownNextMileage = null; checkMileageVsDue(); return; }

    document.getElementById('owner_name').value = data.car.owner_name || '';
    document.getElementById('owner_phone').value = data.car.owner_phone || '';
    if (data.car.car_brand) document.getElementById('car_brand').value = data.car.car_brand;
    document.getElementById('car_model').value = data.car.car_model || '';

    const name = data.car.owner_name || T.kc_no_name;
    const initials = name.trim().split(/\s+/).filter(Boolean).slice(0, 2).map(w => w[0].toUpperCase()).join('') || '?';
    const carLine = [data.car.car_brand, data.car.car_model].filter(Boolean).join(' ');
    const metaParts = [data.car.owner_phone, carLine].filter(Boolean);

    const last = data.history[0];
    const visitCount = data.history.length;
    lastKnownNextMileage = last ? last.next_mileage : null;
    let lastItemsHtml = '';
    if (last && last.items_json) {
      try {
        const items = JSON.parse(last.items_json);
        lastItemsHtml = '<div class="kc-lv-items">' + items.map(it =>
          `<div class="kc-lv-item"><span>${escapeHtml(it.name)}${it.brand ? ' (' + escapeHtml(it.brand) + ')' : ''}${it.qty && it.qty !== 1 ? ' — ' + it.qty + ' ' + T.liters_ph : ''}</span><span>${it.total.toLocaleString('ru-RU')} ${T.currency}</span></div>`
        ).join('') + '</div>';
      } catch (e) { /* старая запись без items_json — просто не показываем разбивку */ }
    }
    const mileageParts = [];
    if (last && last.mileage) mileageParts.push(`${T.kc_last_mileage} ${last.mileage.toLocaleString('ru-RU')} ${T.km_short}`);
    if (last && last.next_mileage) mileageParts.push(`${T.kc_due_mileage} ${last.next_mileage.toLocaleString('ru-RU')} ${T.km_short}`);
    const lastVisitHtml = last ? `
      <div class="kc-last-visit">
        <div class="kc-lv-top">
          <div>
            <div class="kc-lv-service">${escapeHtml(last.service_type || T.history_service_fallback)}</div>
            <div class="kc-lv-date">${last.change_date}${visitCount > 1 ? ` · ${T.kc_visits_total} ${visitCount}` : ''}</div>
          </div>
          <div class="kc-lv-cost">${last.cost ? last.cost.toLocaleString('ru-RU') + ' ' + T.currency : '—'}</div>
        </div>
        ${mileageParts.length ? `<div class="kc-lv-mileage">${mileageParts.join(' · ')}</div>` : ''}
        ${lastItemsHtml}
      </div>
    ` : `<div class="kc-last-visit"><span style="color:var(--hint); font-size:13px;">${T.kc_no_history}</span></div>`;

    panel.innerHTML = `
      <div class="known-client">
        <div class="kc-header"><i class="fa-solid fa-circle-check"></i><span>${T.kc_found_title}</span></div>
        <div class="kc-body">
          <div class="kc-person">
            <div class="kc-avatar">${escapeHtml(initials)}</div>
            <div>
              <div class="kc-name">${escapeHtml(name)}</div>
              <div class="kc-meta">${escapeHtml(metaParts.join(' · '))}</div>
            </div>
          </div>
          <div class="kc-history-label">${T.kc_history_label}</div>
          ${lastVisitHtml}
          <button type="button" class="kc-action-btn" onclick="focusNextEntry()"><i class="fa-solid fa-plus"></i>${T.kc_add_service_btn}</button>
          <div class="kc-action-hint">${T.kc_add_service_hint}</div>
        </div>
      </div>
    `;
  } catch (e) {
    panel.innerHTML = '';
    lastKnownNextMileage = null;
    checkMileageVsDue();
  }
}

function focusNextEntry() {
  const mileage = document.getElementById('mileage');
  mileage.scrollIntoView({ behavior: 'smooth', block: 'center' });
  mileage.focus();
}

async function initItemForms() {
  if (WAREHOUSE_ENABLED) {
    try {
      const res = await fetch('/api/products');
      productsCache = await res.json();
    } catch (e) { /* остаёмся с пустым каталогом — поля просто будут текстовыми */ }
  }
  renderItemLists();
  renderSvcItemLists();
}
initItemForms();

async function submitCar() {
  const items = collectItems();
  const payCashEl = document.getElementById('pay_cash');
  const payCardEl = document.getElementById('pay_card');
  const debtEnabled = document.getElementById('debt_enabled') && document.getElementById('debt_enabled').checked;
  const payload = {
    plate: document.getElementById('plate').value.trim(),
    owner_name: document.getElementById('owner_name').value.trim(),
    owner_phone: document.getElementById('owner_phone').value.trim(),
    car_brand: document.getElementById('car_brand').value,
    car_model: document.getElementById('car_model').value.trim(),
    mileage: document.getElementById('mileage').value,
    next_mileage: document.getElementById('next_mileage').value,
    items: items,
    interval_value: document.getElementById('interval_value').value,
    interval_unit: document.getElementById('interval_unit').value,
    notes: document.getElementById('notes').value.trim(),
    cash_amount: payCashEl ? payCashEl.value : null,
    card_amount: payCardEl ? payCardEl.value : null,
  };
  if (debtEnabled) {
    const total = items.reduce((sum, i) => sum + i.total, 0);
    const cash = parseFloat(payCashEl.value) || 0;
    const card = parseFloat(payCardEl.value) || 0;
    payload.debt_amount = Math.max(0, Math.round(total - cash - card));
    payload.installment_amount = document.getElementById('debt_installment_amount').value;
    payload.interval_days = document.getElementById('debt_interval_days').value;
    if (payload.debt_amount > 0 && (!payload.installment_amount || !payload.interval_days)) {
      showMsg(T.debt_fill_required, false);
      return;
    }
  }
  if (!payload.plate || !payload.owner_name || !payload.interval_value) {
    showMsg(T.msg_fill_required, false);
    return;
  }
  const res = await fetch('/api/add', {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify(payload)
  });
  const data = await res.json();
  if (data.ok) {
    showMsg(`✅ ${T.msg_saved} ${data.next_date || '—'}.`, true);
    ['plate','owner_name','owner_phone','car_model','mileage','next_mileage','notes'].forEach(id => document.getElementById(id).value = '');
    resetItemInputs();
    document.getElementById('interval_value').value = 3;
    document.getElementById('interval_unit').value = 'months';
    paymentSplitTouched = false;
    const payCash = document.getElementById('pay_cash');
    const payCard = document.getElementById('pay_card');
    if (payCash) payCash.value = '';
    if (payCard) payCard.value = '';
    const debtCheckbox = document.getElementById('debt_enabled');
    if (debtCheckbox) {
      debtCheckbox.checked = false;
      document.getElementById('debtFields').style.display = 'none';
      document.getElementById('debt_installment_amount').value = '';
      document.getElementById('debt_interval_days').value = '';
    }
    lastKnownNextMileage = null;
    document.getElementById('mileageCompare').innerHTML = '';
    document.getElementById('knownClientPanel').innerHTML = '';
    if (data.client_link) {
      openModal(payload.plate, data.client_link, payload.owner_phone);
    }
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

async function loadCars() {
  const res = await fetch('/api/cars');
  carsCache = await res.json();
  renderTable();
}

async function loadDebts() {
  const debts = await (await fetch('/api/debts')).json();
  const list = document.getElementById('debtsList');
  if (!debts.length) { list.innerHTML = `<div class="hint-text" style="text-align:center; padding:24px;">${T.debts_empty}</div>`; return; }
  list.innerHTML = debts.map(d => `
    <div class="debt-card ${d.is_overdue ? 'overdue' : ''}">
      <div class="dc-top">
        <div>
          <div class="dc-owner">${escapeHtml(d.owner_name || T.kc_no_name)}</div>
          <div class="dc-meta">${escapeHtml(d.plate_number)}${d.owner_phone ? ' · ' + escapeHtml(d.owner_phone) : ''}</div>
        </div>
        <div>
          <div class="dc-remaining">${d.remaining.toLocaleString('ru-RU')} ${T.currency}</div>
          <div class="dc-due ${d.is_overdue ? 'overdue-text' : ''}">${d.is_overdue ? T.debt_overdue : T.debt_next_due} ${d.next_due_date}</div>
        </div>
      </div>
      <div class="dc-pay-row">
        <input id="debt_pay_${d.id}" type="number" placeholder="${T.debt_pay_placeholder}" style="flex:1;">
        <button class="badge active" style="flex:none;" onclick="payDebt(${d.id})">${T.debt_pay_btn}</button>
      </div>
    </div>
  `).join('');
}

async function payDebt(planId) {
  const input = document.getElementById(`debt_pay_${planId}`);
  const amount = input.value;
  if (!amount || parseFloat(amount) <= 0) { showMsg(T.debt_pay_invalid, false); return; }
  const res = await fetch(`/api/debts/${planId}/pay`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({amount})
  });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.debt_pay_success, true);
    loadDebts();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

let expenseCategoriesCache = [];

async function loadExpenseCategoryOptions(selectId) {
  if (!expenseCategoriesCache.length) {
    expenseCategoriesCache = await (await fetch('/api/expenses/categories')).json();
  }
  const select = document.getElementById(selectId);
  if (!select) return;
  select.innerHTML = expenseCategoriesCache.map(c => `<option value="${escapeHtml(c)}">${escapeHtml(c)}</option>`).join('')
    + `<option value="__custom__">${T.expense_custom_category_option}</option>`;
}

function onExpenseCategoryChange(selectId, customInputId) {
  const select = document.getElementById(selectId);
  const customInput = document.getElementById(customInputId);
  const isCustom = select.value === '__custom__';
  customInput.style.display = isCustom ? 'block' : 'none';
  if (isCustom) customInput.focus();
}

function getSelectedCategory(selectId, customInputId) {
  const select = document.getElementById(selectId);
  if (select.value === '__custom__') {
    return document.getElementById(customInputId).value.trim();
  }
  return select.value;
}

async function loadExpensesTab() {
  await loadExpenseCategoryOptions('exp_category');
  await loadExpenseCategoryOptions('rec_category');
  loadRecurringExpenses();
  loadExpensesJournal();
}

async function submitExpense() {
  const category = getSelectedCategory('exp_category', 'exp_category_custom');
  const name = document.getElementById('exp_name').value.trim();
  const amount = document.getElementById('exp_amount').value;
  if (!category || !amount) { showMsg(T.msg_fill_required, false); return; }
  const res = await fetch('/api/expenses', {
    method: 'POST', headers: {'Content-Type': 'application/json'},
    body: JSON.stringify({category, name, amount}),
  });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.expense_saved, true);
    document.getElementById('exp_name').value = '';
    document.getElementById('exp_amount').value = '';
    document.getElementById('exp_category_custom').value = '';
    document.getElementById('exp_amount_usd').value = '';
    loadExpensesJournal();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

async function loadRecurringExpenses() {
  const list = await (await fetch('/api/recurring_expenses')).json();
  const el = document.getElementById('recurringExpensesList');
  if (!el) return;
  const today = new Date().toISOString().slice(0, 10);
  el.innerHTML = list.length ? list.map(r => `
    <div class="recur-exp-card ${r.next_due_date < today ? 'overdue' : ''}">
      <div class="rec-top">
        <div>
          <div class="rec-name">${escapeHtml(r.name || r.category)}</div>
          <div class="rec-meta">${escapeHtml(r.category)} · ${T.expense_next_due} ${r.next_due_date}</div>
        </div>
        <div class="rec-amount">${r.amount.toLocaleString('ru-RU')} ${T.currency}</div>
      </div>
      <div class="rec-actions">
        <button class="badge active" style="flex:1;" onclick="payRecurringExpense(${r.id})">${T.expense_mark_paid_btn}</button>
        <button class="badge inactive" style="flex:none;" onclick="deleteRecurringExpenseBtn(${r.id})">${T.entry_delete}</button>
      </div>
    </div>
  `).join('') : `<div class="hint-text">${T.expense_no_recurring}</div>`;
}

async function createRecurringExpense() {
  const category = getSelectedCategory('rec_category', 'rec_category_custom');
  const name = document.getElementById('rec_name').value.trim();
  const amount = document.getElementById('rec_amount').value;
  const day_of_month = document.getElementById('rec_day').value;
  if (!category || !amount || !day_of_month) { showMsg(T.msg_fill_required, false); return; }
  const res = await fetch('/api/recurring_expenses', {
    method: 'POST', headers: {'Content-Type': 'application/json'},
    body: JSON.stringify({category, name, amount, day_of_month}),
  });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.expense_recurring_created, true);
    document.getElementById('rec_name').value = '';
    document.getElementById('rec_amount').value = '';
    document.getElementById('rec_day').value = '';
    document.getElementById('rec_category_custom').value = '';
    document.getElementById('rec_amount_usd').value = '';
    loadRecurringExpenses();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

async function payRecurringExpense(id) {
  const res = await fetch(`/api/recurring_expenses/${id}/pay`, { method: 'POST' });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.expense_mark_paid_success, true);
    loadRecurringExpenses();
    loadExpensesJournal();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

async function deleteRecurringExpenseBtn(id) {
  if (!confirm(T.expense_delete_recurring_confirm)) return;
  const res = await fetch(`/api/recurring_expenses/${id}`, { method: 'DELETE' });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.entry_deleted, true);
    loadRecurringExpenses();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

let journalCache = [];

async function loadExpensesJournal() {
  journalCache = await (await fetch('/api/expenses')).json();
  const el = document.getElementById('expensesJournal');
  if (!el) return;
  el.innerHTML = journalCache.length ? journalCache.slice(0, 30).map(e => `
    <div>
      <div class="journal-row">
        <div>
          <div class="jr-cat">${escapeHtml(e.name || e.category)}</div>
          <div class="jr-meta">${escapeHtml(e.category)} · ${e.expense_date}</div>
        </div>
        <div style="display:flex; align-items:center; gap:10px;">
          <div class="jr-amount">−${e.amount.toLocaleString('ru-RU')} ${T.currency}</div>
          <button type="button" class="jr-icon-btn" onclick="toggleExpenseEditForm(${e.id})" title="${T.entry_edit}"><i class="fa-solid fa-pen"></i></button>
          <button type="button" class="jr-icon-btn" style="color:#B3241C;" onclick="deleteExpenseEntry(${e.id})" title="${T.entry_delete}"><i class="fa-solid fa-trash"></i></button>
        </div>
      </div>
      <div id="jredit-${e.id}" class="jr-edit-form" style="display:none;"></div>
    </div>
  `).join('') : `<div class="hint-text">${T.dash_no_data}</div>`;
}

async function toggleExpenseEditForm(id) {
  const panel = document.getElementById('jredit-' + id);
  if (!panel) return;
  const isOpen = panel.style.display !== 'none';
  document.querySelectorAll('.jr-edit-form').forEach(el => { if (el !== panel) el.style.display = 'none'; });
  if (isOpen) { panel.style.display = 'none'; return; }
  const entry = journalCache.find(e => e.id === id);
  if (!entry) return;
  const catSelId = `jr_cat_${id}`, catCustomId = `jr_cat_custom_${id}`;
  await loadExpenseCategoryOptions(catSelId);
  const select = document.getElementById(catSelId);
  const isKnown = expenseCategoriesCache.includes(entry.category);
  panel.innerHTML = `
    <div class="field">
      <label>${T.expense_category_label}</label>
      <select id="${catSelId}" onchange="onExpenseCategoryChange('${catSelId}', '${catCustomId}')"></select>
      <input id="${catCustomId}" placeholder="${T.expense_custom_category_ph}" style="display:none; margin-top:6px;" value="${isKnown ? '' : escapeHtml(entry.category)}">
    </div>
    <div class="field">
      <label>${T.expense_name_label}</label>
      <input id="jr_name_${id}" value="${escapeHtml(entry.name || '')}">
    </div>
    <div class="row2">
      <div class="field">
        <label>${T.expense_amount_label}</label>
        <input id="jr_amount_${id}" type="number" value="${entry.amount}">
      </div>
      <div class="field">
        <label>${T.expense_date_label}</label>
        <input id="jr_date_${id}" type="date" value="${entry.expense_date}">
      </div>
    </div>
    <button class="submit" onclick="saveExpenseEdit(${id})">${T.btn_save}</button>
  `;
  await loadExpenseCategoryOptions(catSelId);
  const selectEl = document.getElementById(catSelId);
  if (isKnown) {
    selectEl.value = entry.category;
  } else {
    selectEl.value = '__custom__';
    onExpenseCategoryChange(catSelId, catCustomId);
  }
  panel.style.display = 'block';
}

async function saveExpenseEdit(id) {
  const category = getSelectedCategory(`jr_cat_${id}`, `jr_cat_custom_${id}`);
  const name = document.getElementById(`jr_name_${id}`).value.trim();
  const amount = document.getElementById(`jr_amount_${id}`).value;
  const expense_date = document.getElementById(`jr_date_${id}`).value;
  if (!category || !amount) { showMsg(T.msg_fill_required, false); return; }
  const res = await fetch(`/api/expenses/${id}`, {
    method: 'PUT', headers: {'Content-Type': 'application/json'},
    body: JSON.stringify({category, name, amount, expense_date}),
  });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.expense_saved, true);
    loadExpensesJournal();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

async function deleteExpenseEntry(id) {
  if (!confirm(T.expense_delete_entry_confirm)) return;
  const res = await fetch(`/api/expenses/${id}`, { method: 'DELETE' });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.entry_deleted, true);
    loadExpensesJournal();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

function renderTable() {
  const q = (document.getElementById('search').value || '').toLowerCase();
  const rows = carsCache.filter(c =>
    (c.plate_number || '').toLowerCase().includes(q) || (c.owner_name || '').toLowerCase().includes(q)
  );
  document.getElementById('table-body').innerHTML = rows.length ? rows.map((c, i) => `
    <div class="car-card" onclick="toggleHistory(${escapeHtml(JSON.stringify(c.plate_number))})">
      <div class="cc-top">
        <span class="cc-plate">${escapeHtml(c.plate_number)}</span>
        ${c.telegram_id
            ? `<span class="badge linked cc-linkbtn">${T.badge_linked}</span>`
            : `<button class="badge unlinked cc-linkbtn" onclick="event.stopPropagation(); openModal(${escapeHtml(JSON.stringify(c.plate_number))}, ${escapeHtml(JSON.stringify(c.client_link || ''))}, ${escapeHtml(JSON.stringify(c.owner_phone || ''))})">${T.badge_unlinked_btn}</button>`}
      </div>
      <div class="cc-owner">${escapeHtml(c.owner_name || T.kc_no_name)}</div>
      <div class="cc-meta">${escapeHtml([c.owner_phone, [c.car_brand, c.car_model].filter(Boolean).join(' ')].filter(Boolean).join(' · '))}</div>
      <div class="cc-bottom">
        <div>
          <div class="cc-cost">${c.cost ? c.cost.toLocaleString('ru-RU') + ' ' + T.currency : '—'}</div>
          <div class="cc-cost-date">${c.change_date || '—'}</div>
        </div>
        <div class="cc-next">${T.th_next_change}<br>${c.next_change_date || '—'}</div>
        <i class="fa-solid fa-chevron-right cc-chevron"></i>
      </div>
    </div>
  `).join('') : `<div class="hint-text" style="text-align:center; padding:20px;">${T.table_empty}</div>`;
}

async function toggleHistory(plate) {
  const panel = document.getElementById('clientCardPanel');
  if (openHistoryRow === plate) {
    panel.style.display = 'none';
    panel.innerHTML = '';
    openHistoryRow = null;
    return;
  }
  openHistoryRow = plate;
  panel.style.display = '';
  panel.innerHTML = T.history_loading;
  panel.scrollIntoView({ behavior: 'smooth', block: 'start' });
  const res = await fetch('/api/history/' + encodeURIComponent(plate));
  const data = await res.json();
  const history = data.history || [];
  historyDataCache[plate] = history;
  const body = panel;

  const car = data.car || {};
  const name = car.owner_name || T.kc_no_name;
  const initials = name.trim().split(/\s+/).filter(Boolean).slice(0, 2).map(w => w[0].toUpperCase()).join('') || '?';
  const carLine = [car.car_brand, car.car_model].filter(Boolean).join(' ');
  const metaParts = [car.owner_phone, carLine].filter(Boolean);

  const entriesHtml = history.length ? history.map(h => {
    let itemsHtml = '';
    if (h.items_json) {
      try {
        const items = JSON.parse(h.items_json);
        itemsHtml = '<div class="kc-he-items">' + items.map(it =>
          `${escapeHtml(it.name)}${it.brand ? ' (' + escapeHtml(it.brand) + ')' : ''}${it.qty && it.qty !== 1 ? ' — ' + it.qty + ' ' + T.liters_ph : ''}: ${it.total.toLocaleString('ru-RU')} ${T.currency}`
        ).join('<br>') + '</div>';
      } catch (e) { /* старая запись без items_json */ }
    }
    return `
    <div class="kc-hist-entry" id="hist-entry-${h.id}">
      <div class="kc-he-top">
        <div>
          <div class="kc-he-service">${escapeHtml(h.service_type || T.history_service_fallback)}</div>
          <div class="kc-he-date">${h.change_date} · ${T.history_mileage_label} ${h.mileage || '—'} км</div>
        </div>
        <div class="kc-he-cost">${h.cost ? h.cost.toLocaleString('ru-RU') + ' ' + T.currency : '—'}</div>
      </div>
      <div class="kc-he-meta">${T.history_next_mileage_label} ${h.next_mileage || '—'} км · ${T.history_next_label} ${h.next_change_date || '—'}</div>
      ${itemsHtml}
      ${h.notes ? `<div class="kc-he-notes">${T.history_notes_label} ${escapeHtml(h.notes)}</div>` : ''}
      <div class="kc-he-actions">
        <button class="history-toggle" onclick="openEditModalById(${h.id}, ${escapeHtml(JSON.stringify(plate))})">${T.entry_edit}</button>
        &nbsp;·&nbsp;
        <button class="history-toggle" style="color:#B3241C;" onclick="deleteEntry(${h.id}, ${escapeHtml(JSON.stringify(plate))})">${T.entry_delete}</button>
      </div>
    </div>
  `;
  }).join('') : `<div class="kc-hist-entry" style="color:var(--hint); text-align:center;">${T.history_empty}</div>`;

  body.innerHTML = `
    <div class="known-client">
      <div class="kc-header"><i class="fa-solid fa-user"></i><span>${T.kc_history_card_title}</span></div>
      <div class="kc-body">
        <div class="kc-person">
          <div class="kc-avatar">${escapeHtml(initials)}</div>
          <div>
            <div class="kc-name">${escapeHtml(name)}</div>
            <div class="kc-meta">${escapeHtml(metaParts.join(' · '))}</div>
          </div>
        </div>
        <button type="button" class="kc-action-btn" onclick="openAddServiceModal(${escapeHtml(JSON.stringify(plate))})"><i class="fa-solid fa-plus"></i>${T.kc_add_service_btn}</button>
        <div style="display:flex; gap:8px; margin-top:8px;">
          <button type="button" class="history-toggle" style="flex:1;" onclick="toggleCarEditForm()">${T.kc_edit_car_btn}</button>
          <button type="button" class="history-toggle" style="flex:1; color:#B3241C;" onclick="deleteCarCompletely(${escapeHtml(JSON.stringify(plate))})">${T.kc_delete_car_btn}</button>
        </div>
        <div id="carEditForm" style="display:none; margin-top:10px; padding:12px; background:var(--field-bg); border-radius:10px;">
          <div class="field">
            <label style="font-size:11px;">${T.field_plate}</label>
            <input id="edit_car_plate" value="${escapeHtml(plate)}">
          </div>
          <div class="field">
            <label style="font-size:11px;">${T.field_owner_name}</label>
            <input id="edit_car_owner_name" value="${escapeHtml(name)}">
          </div>
          <div class="field">
            <label style="font-size:11px;">${T.field_owner_phone}</label>
            <input id="edit_car_owner_phone" value="${escapeHtml(car.owner_phone || '')}">
          </div>
          <div class="row2">
            <div class="field">
              <label style="font-size:11px;">${T.field_car_brand}</label>
              <input id="edit_car_brand" value="${escapeHtml(car.car_brand || '')}">
            </div>
            <div class="field">
              <label style="font-size:11px;">${T.field_car_model}</label>
              <input id="edit_car_model" value="${escapeHtml(car.car_model || '')}">
            </div>
          </div>
          <button type="button" class="submit" onclick="saveCarEdit(${escapeHtml(JSON.stringify(plate))})">${T.btn_save}</button>
        </div>
        <div class="kc-hist-list">
          <div class="kc-history-label">${T.kc_full_history_label} (${history.length})</div>
          ${entriesHtml}
        </div>
      </div>
    </div>
  `;
}

function escapeHtml(str) {
  return String(str)
    .replace(/&/g, '&amp;')
    .replace(/</g, '&lt;')
    .replace(/>/g, '&gt;')
    .replace(/"/g, '&quot;')
    .replace(/'/g, '&#039;');
}

function openModal(plate, link, phone) {
  if (!link) { showMsg(T.modal_no_bot_username, false); return; }
  document.getElementById('modalPlate').textContent = plate;
  document.getElementById('modalLink').textContent = link;
  document.getElementById('modalQr').src = 'https://api.qrserver.com/v1/create-qr-code/?size=200x200&data=' + encodeURIComponent(link);
  const tgBtn = document.getElementById('modalTg');
  tgBtn.href = 'https://t.me/share/url?url=' + encodeURIComponent(link) + '&text=' + encodeURIComponent(T.share_text_tg);
  const waBtn = document.getElementById('modalWa');
  if (phone) {
    let digits = phone.replace(/\\D/g, '');
    if (digits && !digits.startsWith('998') && digits.length <= 9) digits = '998' + digits;
    const text = encodeURIComponent(T.share_text_wa + ' ' + link);
    waBtn.href = 'https://wa.me/' + digits + '?text=' + text;
    waBtn.style.display = 'block';
  } else {
    waBtn.style.display = 'none';
  }
  document.getElementById('linkModal').classList.add('open');
}

function closeModal() {
  document.getElementById('linkModal').classList.remove('open');
}

function copyLink() {
  const text = document.getElementById('modalLink').textContent;
  navigator.clipboard.writeText(text).then(() => showMsg(T.link_copied, true));
}

// ---- Единое модальное окно: и "изменить запись", и "внести новую замену уже известному клиенту" ----
let svcModal = { mode: null, id: null, plate: null };  // mode: 'edit' | 'add'

function renderSvcItemLists() {
  document.getElementById('svcFluidsList').innerHTML = FLUID_KEYS.map((key, i) => `
    <div class="item-row">
      <span class="item-name">${T[key]}</span>
      ${brandFieldHtml(key, `svc_fluid_brand_${i}`, `onSvcFluidProductPicked(${i})`)}
      <input id="svc_fluid_price_${i}" type="number" placeholder="${T.price_per_liter_ph}" oninput="updateSvcTotal()">
      <input id="svc_fluid_liters_${i}" type="number" step="0.1" placeholder="${T.liters_ph}" oninput="updateSvcTotal()">
    </div>
  `).join('');
  document.getElementById('svcFiltersList').innerHTML = FILTER_KEYS.map((key, i) => {
    const prods = productsForCategory(key);
    const brandField = prods.length ? brandFieldHtml(key, `svc_filter_brand_${i}`, `onSvcFilterProductPicked(${i})`) : '';
    return `
    <div class="item-row">
      <span class="item-name" style="flex:${prods.length ? '1.3' : '2.3'};">${T[key]}</span>
      ${brandField}
      <input id="svc_filter_price_${i}" type="number" placeholder="${T.price_ph}" oninput="updateSvcTotal()">
    </div>
  `;
  }).join('');
}

function onSvcFluidProductPicked(i) {
  const el = document.getElementById(`svc_fluid_brand_${i}`);
  const opt = el.options[el.selectedIndex];
  if (opt && opt.dataset.price) document.getElementById(`svc_fluid_price_${i}`).value = opt.dataset.price;
  updateSvcTotal();
}

function onSvcFilterProductPicked(i) {
  const el = document.getElementById(`svc_filter_brand_${i}`);
  const opt = el.options[el.selectedIndex];
  if (opt && opt.dataset.price) document.getElementById(`svc_filter_price_${i}`).value = opt.dataset.price;
  updateSvcTotal();
}

function collectSvcItems() {
  const items = [];
  FLUID_KEYS.forEach((key, i) => {
    const price = parseFloat(document.getElementById(`svc_fluid_price_${i}`).value) || 0;
    const liters = parseFloat(document.getElementById(`svc_fluid_liters_${i}`).value) || 0;
    if (price > 0 && liters > 0) {
      const { brand, product_id } = readBrandField(`svc_fluid_brand_${i}`);
      items.push({
        key, name: T[key], brand, product_id,
        unit_price: price, qty: liters, total: Math.round(price * liters),
      });
    }
  });
  FILTER_KEYS.forEach((key, i) => {
    const price = parseFloat(document.getElementById(`svc_filter_price_${i}`).value) || 0;
    if (price > 0) {
      const brandEl = document.getElementById(`svc_filter_brand_${i}`);
      const { brand, product_id } = brandEl ? readBrandField(`svc_filter_brand_${i}`) : { brand: null, product_id: null };
      items.push({key, name: T[key], brand, product_id, unit_price: price, qty: 1, total: Math.round(price)});
    }
  });
  const otherName = document.getElementById('svc_other_name').value.trim();
  const otherPrice = parseFloat(document.getElementById('svc_other_price').value) || 0;
  if (otherPrice > 0) {
    items.push({key: 'other', name: `${T.other_prefix}: ${otherName || T.other_unnamed}`, unit_price: otherPrice, qty: 1, total: Math.round(otherPrice)});
  }
  items.push(...collectOtherStockItems('svcOther'));
  return items;
}

let svcPaymentSplitTouched = false;

function updateSvcTotal() {
  const total = collectSvcItems().reduce((sum, i) => sum + i.total, 0);
  document.getElementById('svcTotalCost').textContent = total.toLocaleString('ru-RU') + ' ' + T.currency;
  const payCash = document.getElementById('svc_pay_cash');
  const payCard = document.getElementById('svc_pay_card');
  if (payCash && payCard && !svcPaymentSplitTouched) {
    payCash.value = total || '';
    payCard.value = '';
  }
}

function onSvcPayCashInput() {
  svcPaymentSplitTouched = true;
  const total = collectSvcItems().reduce((sum, i) => sum + i.total, 0);
  const cash = parseFloat(document.getElementById('svc_pay_cash').value) || 0;
  document.getElementById('svc_pay_card').value = Math.max(0, Math.round(total - cash));
}

function onSvcPayCardInput() {
  svcPaymentSplitTouched = true;
  const total = collectSvcItems().reduce((sum, i) => sum + i.total, 0);
  const card = parseFloat(document.getElementById('svc_pay_card').value) || 0;
  document.getElementById('svc_pay_cash').value = Math.max(0, Math.round(total - card));
}

function resetSvcItemInputs() {
  FLUID_KEYS.forEach((_, i) => {
    document.getElementById(`svc_fluid_brand_${i}`).value = '';
    document.getElementById(`svc_fluid_price_${i}`).value = '';
    document.getElementById(`svc_fluid_liters_${i}`).value = '';
  });
  FILTER_KEYS.forEach((_, i) => {
    document.getElementById(`svc_filter_price_${i}`).value = '';
    const brandEl = document.getElementById(`svc_filter_brand_${i}`);
    if (brandEl) brandEl.value = '';
  });
  document.getElementById('svc_other_name').value = '';
  document.getElementById('svc_other_price').value = '';
  svcOtherStockRows = [];
  renderOtherStockRows('svcOther');
  updateSvcTotal();
}

function fillSvcItemsFrom(items) {
  // подставляет уже сохранённые позиции в поля модалки (режим редактирования)
  const otherStockItems = [];
  (items || []).forEach(it => {
    if (!it.key) return;
    if (it.key === 'other') {
      const label = (it.name || '').replace(T.other_prefix + ': ', '');
      document.getElementById('svc_other_name').value = label === T.other_unnamed ? '' : label;
      document.getElementById('svc_other_price').value = it.total ?? '';
      return;
    }
    if (it.key === 'other_stock') {
      otherStockItems.push(it);
      return;
    }
    const fi = FLUID_KEYS.indexOf(it.key);
    if (fi !== -1) {
      const brandEl = document.getElementById(`svc_fluid_brand_${fi}`);
      selectOrPreserveBrand(brandEl, it);
      document.getElementById(`svc_fluid_price_${fi}`).value = it.unit_price ?? '';
      document.getElementById(`svc_fluid_liters_${fi}`).value = it.qty ?? '';
      return;
    }
    const filI = FILTER_KEYS.indexOf(it.key);
    if (filI !== -1) {
      const brandEl = document.getElementById(`svc_filter_brand_${filI}`);
      if (brandEl) selectOrPreserveBrand(brandEl, it);
      document.getElementById(`svc_filter_price_${filI}`).value = it.total ?? '';
    }
  });
  fillOtherStockRowsFrom('svcOther', otherStockItems);
}

function openEditModalById(id, plate) {
  const entry = (historyDataCache[plate] || []).find(h => h.id === id);
  if (!entry) { showMsg(T.msg_error + ' запись не найдена в кэше, обновите список', false); return; }
  openEditModal(plate, entry);
}

function openEditModal(plate, entry) {
  svcModal = { mode: 'edit', id: entry.id, plate };
  document.getElementById('svcModalTitle').textContent = T.entry_edit_title;
  document.getElementById('edit_mileage').value = entry.mileage ?? '';
  document.getElementById('edit_next_mileage').value = entry.next_mileage ?? '';
  document.getElementById('edit_interval_value').value = entry.interval_months ?? '';
  document.getElementById('edit_interval_unit').value = entry.interval_unit || 'months';
  document.getElementById('edit_notes').value = entry.notes || '';
  resetSvcItemInputs();
  if (entry.items_json) {
    try { fillSvcItemsFrom(JSON.parse(entry.items_json)); } catch (e) {}
  }
  updateSvcTotal();
  document.getElementById('svc_pay_cash').value = entry.cash_amount ?? entry.cost ?? '';
  document.getElementById('svc_pay_card').value = entry.card_amount ?? '';
  svcPaymentSplitTouched = true;  // это реальная сохранённая разбивка, не пересчитывать автоматически
  document.getElementById('editModal').classList.add('open');
}

function openAddServiceModal(plate) {
  svcModal = { mode: 'add', id: null, plate };
  document.getElementById('svcModalTitle').textContent = T.add_service_title;
  document.getElementById('edit_mileage').value = '';
  document.getElementById('edit_next_mileage').value = '';
  document.getElementById('edit_interval_value').value = 3;
  document.getElementById('edit_interval_unit').value = 'months';
  document.getElementById('edit_notes').value = '';
  svcPaymentSplitTouched = false;
  document.getElementById('svc_pay_cash').value = '';
  document.getElementById('svc_pay_card').value = '';
  resetSvcItemInputs();
  document.getElementById('editModal').classList.add('open');
}

function closeEditModal() {
  document.getElementById('editModal').classList.remove('open');
}

async function saveEdit() {
  const items = collectSvcItems();
  const mileage = document.getElementById('edit_mileage').value || null;
  const next_mileage = document.getElementById('edit_next_mileage').value || null;
  const interval_value = document.getElementById('edit_interval_value').value || null;
  const interval_unit = document.getElementById('edit_interval_unit').value;
  const notes = document.getElementById('edit_notes').value;
  const cash_amount = document.getElementById('svc_pay_cash').value || null;
  const card_amount = document.getElementById('svc_pay_card').value || null;

  let res;
  if (svcModal.mode === 'edit') {
    res = await fetch('/api/oil_change/' + svcModal.id, {
      method: 'PUT', headers: {'Content-Type': 'application/json'},
      body: JSON.stringify({ mileage, next_mileage, interval_value, interval_unit, notes, items, cash_amount, card_amount }),
    });
  } else {
    const carRow = carsCache.find(c => c.plate_number === svcModal.plate);
    res = await fetch('/api/add', {
      method: 'POST', headers: {'Content-Type': 'application/json'},
      body: JSON.stringify({
        plate: svcModal.plate, owner_name: carRow ? carRow.owner_name : '', mileage, next_mileage,
        interval_value, interval_unit, notes, items, cash_amount, card_amount,
      }),
    });
  }
  const data = await res.json();
  if (data.ok) {
    closeEditModal();
    showMsg(svcModal.mode === 'edit' ? T.entry_saved : T.service_added, true);
    openHistoryRow = null;  // чтобы toggleHistory ниже заново открыл панель со свежими данными, а не закрыл её
    toggleHistory(svcModal.plate);
    loadCars();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

function toggleCarEditForm() {
  const form = document.getElementById('carEditForm');
  form.style.display = form.style.display === 'none' ? 'block' : 'none';
}

async function saveCarEdit(oldPlate) {
  const payload = {
    plate: document.getElementById('edit_car_plate').value.trim(),
    owner_name: document.getElementById('edit_car_owner_name').value.trim(),
    owner_phone: document.getElementById('edit_car_owner_phone').value.trim(),
    car_brand: document.getElementById('edit_car_brand').value.trim(),
    car_model: document.getElementById('edit_car_model').value.trim(),
  };
  if (!payload.plate || !payload.owner_name) {
    showMsg(T.msg_fill_required, false);
    return;
  }
  const res = await fetch('/api/car/' + encodeURIComponent(oldPlate), {
    method: 'PUT', headers: {'Content-Type': 'application/json'}, body: JSON.stringify(payload)
  });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.kc_car_saved, true);
    openHistoryRow = null;
    toggleHistory(data.plate);
    loadCars();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

async function deleteCarCompletely(plate) {
  if (!confirm(T.kc_delete_car_confirm.replace('{plate}', plate))) return;
  const res = await fetch('/api/car/' + encodeURIComponent(plate), { method: 'DELETE' });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.kc_car_deleted, true);
    document.getElementById('clientCardPanel').style.display = 'none';
    document.getElementById('clientCardPanel').innerHTML = '';
    openHistoryRow = null;
    loadCars();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}

async function deleteEntry(id, plate) {
  if (!confirm(T.entry_delete_confirm)) return;
  const res = await fetch('/api/oil_change/' + id, { method: 'DELETE' });
  const data = await res.json();
  if (data.ok) {
    showMsg(T.entry_deleted, true);
    openHistoryRow = null;
    toggleHistory(plate);
    loadCars();
  } else {
    showMsg(T.msg_error + ' ' + data.error, false);
  }
}
</script>
</body>
</html>
"""

PAGE = PAGE + MODAL_AND_SCRIPT


@app.route("/")
@login_required
def index():
    import json as _json
    shop = db.get_shop(g.shop_id)
    return render_template_string(
        PAGE, brands=CAR_BRANDS, service_types=SERVICE_TYPES,
        shop_name=session.get("shop_name") or "Замена масла",
        T=g.T, lang=g.lang, t_json=_json.dumps(g.T, ensure_ascii=False),
        sms_enabled=bool(shop.get("sms_enabled")) if shop else False,
        eskiz_email=(shop.get("eskiz_email") or "") if shop else "",
        warehouse_enabled=bool(shop.get("warehouse_enabled")) if shop else False,
        is_employee=g.is_employee,
        is_branch=g.is_branch,
        usd_rate=shop.get("usd_rate") if shop else None,
    )


@app.route("/api/set_language", methods=["POST"])
@login_required
def api_set_language():
    data = request.get_json(force=True)
    lang = data.get("language")
    if lang not in ("ru", "uz"):
        return jsonify({"ok": False, "error": "invalid language"}), 400
    db.set_shop_language(g.shop_id, lang)
    return jsonify({"ok": True})


@app.route("/api/usd_rate", methods=["POST"])
@login_required
@profit_blocked
def api_set_usd_rate():
    """Владелец точки/главный аккаунт сам выставляет свой курс доллара —
    сотруднику и филиалу это не нужно (они не вписывают цену закупки)."""
    data = request.get_json(force=True)
    try:
        rate = float(data.get("rate")) if data.get("rate") not in (None, "") else None
        if rate is not None and rate <= 0:
            raise ValueError()
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "укажите положительное число"}), 400
    db.set_shop_usd_rate(g.shop_id, rate)
    return jsonify({"ok": True, "rate": rate})




@app.route("/api/cars")
@login_required
def api_cars():
    cars = db.get_all_cars_overview(g.shop_id)
    for c in cars:
        c["client_link"] = _client_link(c.get("link_token"))
    return jsonify(cars)


def _strip_cost_price(history):
    """Убирает cost_price (цена закупки, снятая со склада на момент продажи) из
    items_json перед отправкой сотруднику — иначе цена закупки была бы видна
    через сырой ответ сервера, даже если интерфейс её не показывает."""
    import json as _json
    for row in history:
        if not row.get("items_json"):
            continue
        try:
            items = _json.loads(row["items_json"])
        except (ValueError, TypeError):
            continue
        for item in items:
            item.pop("cost_price", None)
        row["items_json"] = _json.dumps(items, ensure_ascii=False)
    return history


@app.route("/api/history/<plate>")
@login_required
def api_history(plate):
    car, history = db.get_car_history(g.shop_id, plate)
    if g.is_employee or g.is_branch:
        history = _strip_cost_price(history)
    return jsonify({"car": car, "history": history})


@app.route("/api/oil_change/<int:oc_id>", methods=["PUT"])
@login_required
def api_update_oil_change(oc_id):
    data = request.get_json(force=True)
    try:
        ok = db.update_oil_change(
            oc_id, g.shop_id,
            mileage=int(data["mileage"]) if data.get("mileage") not in (None, "") else None,
            next_mileage=int(data["next_mileage"]) if data.get("next_mileage") not in (None, "") else None,
            cost=int(data["cost"]) if data.get("cost") not in (None, "") else None,
            interval_value=int(data["interval_value"]) if data.get("interval_value") not in (None, "") else None,
            interval_unit=data.get("interval_unit"),
            notes=data.get("notes"),
            items=data.get("items"),
            cash_amount=int(data["cash_amount"]) if data.get("cash_amount") not in (None, "") else None,
            card_amount=int(data["card_amount"]) if data.get("card_amount") not in (None, "") else None,
        )
    except (ValueError, TypeError) as e:
        return jsonify({"ok": False, "error": str(e)}), 400
    if not ok:
        return jsonify({"ok": False, "error": "запись не найдена"}), 404
    return jsonify({"ok": True})


@app.route("/api/oil_change/<int:oc_id>", methods=["DELETE"])
@login_required
def api_delete_oil_change(oc_id):
    ok = db.delete_oil_change(oc_id, g.shop_id)
    if not ok:
        return jsonify({"ok": False, "error": "запись не найдена"}), 404
    return jsonify({"ok": True})


@app.route("/api/car/<plate>", methods=["PUT"])
@login_required
def api_update_car(plate):
    """Редактирование данных клиента и машины — имя, телефон, госномер,
    марка/модель. Доступно всем ролям (владелец, филиал, сотрудник),
    ограничено только собственной точкой."""
    data = request.get_json(force=True)
    new_plate = (data.get("plate") or "").strip()
    owner_name = (data.get("owner_name") or "").strip()
    owner_phone = (data.get("owner_phone") or "").strip()
    car_brand = data.get("car_brand") or None
    car_model = (data.get("car_model") or "").strip() or None
    if not new_plate or not owner_name:
        return jsonify({"ok": False, "error": "госномер и имя обязательны"}), 400
    ok = db.update_car_and_client(g.shop_id, plate, new_plate, owner_name, owner_phone, car_brand, car_model)
    if not ok:
        return jsonify({"ok": False, "error": "машина не найдена, или новый госномер уже занят другой машиной"}), 400
    return jsonify({"ok": True, "plate": db.normalize_plate(new_plate)})


@app.route("/api/car/<plate>", methods=["DELETE"])
@login_required
def api_delete_car(plate):
    """Полное удаление машины — сама машина, вся её история, связанные
    долги. Клиент (владелец) остаётся, если у него есть другие машины."""
    ok = db.delete_car_completely(g.shop_id, plate)
    if not ok:
        return jsonify({"ok": False, "error": "машина не найдена"}), 404
    return jsonify({"ok": True})


@app.route("/api/add", methods=["POST"])
@login_required
def api_add():
    data = request.get_json(force=True)
    try:
        plate = data["plate"]
        owner_name = data["owner_name"]
        owner_phone = data.get("owner_phone") or None
        car_brand = data.get("car_brand") or None
        car_model = data.get("car_model") or None
        mileage = int(data["mileage"]) if data.get("mileage") else None
        next_mileage = int(data["next_mileage"]) if data.get("next_mileage") else None
        items = data.get("items") or []
        interval_value = int(data["interval_value"])
        interval_unit = data.get("interval_unit") or "months"
        if interval_unit not in ("days", "months"):
            interval_unit = "months"
        notes = data.get("notes") or ""
        cash_amount = int(data["cash_amount"]) if data.get("cash_amount") not in (None, "") else None
        card_amount = int(data["card_amount"]) if data.get("card_amount") not in (None, "") else None
        debt_amount = int(data["debt_amount"]) if data.get("debt_amount") not in (None, "") else 0
        installment_amount = int(data["installment_amount"]) if data.get("installment_amount") not in (None, "") else None
        interval_days = int(data["interval_days"]) if data.get("interval_days") not in (None, "") else None

        existing_car = db.find_car(g.shop_id, plate)
        if existing_car:
            client_id = existing_car["client_id"]
            car_id = db.create_or_update_car(g.shop_id, plate, client_id, car_brand, car_model)
        else:
            client = db.get_or_create_client(g.shop_id, owner_name, owner_phone)
            client_id = client["id"]
            car_id = db.create_or_update_car(g.shop_id, plate, client_id, car_brand, car_model)

        oc_id, next_date = db.add_oil_change(
            car_id, mileage, None, None, False, None, interval_value, interval_unit, notes,
            next_mileage=next_mileage, items=items, cash_amount=cash_amount, card_amount=card_amount
        )

        if debt_amount > 0:
            if not installment_amount or not interval_days:
                return jsonify({"ok": False, "error": "укажите сумму платежа и период для рассрочки"}), 400
            db.create_installment_plan(g.shop_id, car_id, debt_amount, installment_amount, interval_days, oil_change_id=oc_id)

        car_after, _ = db.get_car_history(g.shop_id, plate)
        link = None
        if car_after and not car_after["telegram_id"]:
            link = _client_link(car_after["link_token"])

        return jsonify({"ok": True, "next_date": next_date, "client_link": link})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400


@app.route("/api/debts")
@login_required
def api_list_debts():
    return jsonify(db.get_active_debts(g.shop_id))


@app.route("/api/dashboard")
@login_required
def api_dashboard():
    """Всё для dashboard одним запросом — график выручки за 30 дней, топ
    товаров по количеству, сводка по долгам, и товары с малым остатком.
    Склад и товары не относятся к филиалу/сотруднику без склада — тогда
    просто возвращаются пустыми, без ошибки."""
    result = {
        "daily_revenue": db.get_daily_revenue(g.shop_id, days=30),
        "top_products": db.get_top_products_by_qty(g.shop_id, days=30, limit=5),
        "debt_summary": db.get_debt_summary(g.shop_id),
        "low_stock": [],
        "net_profit": None,
    }
    if not g.is_employee:
        result["net_profit"] = db.get_net_profit_30d(g.shop_id)
    shop = db.get_shop(g.shop_id)
    if shop and shop.get("warehouse_enabled") and not g.is_employee:
        low_stock = db.get_low_stock_products(g.shop_id)
        if g.is_branch:
            for p in low_stock:
                p.pop("purchase_price", None)
        result["low_stock"] = low_stock
    return jsonify(result)


@app.route("/api/dashboard/top_brands")
@login_required
def api_dashboard_top_brands():
    """Топ-10 брендов внутри одной категории — для раскрытия по клику на
    строку категории в dashboard."""
    category = request.args.get("category", "")
    if not category:
        return jsonify({"ok": False, "error": "укажите категорию"}), 400
    return jsonify(db.get_top_brands_for_category(g.shop_id, category, days=30, limit=10))


@app.route("/api/expenses/categories")
@login_required
def api_expense_categories():
    return jsonify(db.EXPENSE_PRESET_CATEGORIES)


@app.route("/api/expenses")
@login_required
@employee_blocked
def api_list_expenses():
    """Журнал разовых расходов за период — сотруднику не показываем
    (финансовая информация точки, как и цена закупки)."""
    date_from = request.args.get("from") or "2000-01-01"
    date_to = request.args.get("to") or "2100-01-01"
    return jsonify(db.get_expenses(g.shop_id, date_from, date_to))


@app.route("/api/expenses", methods=["POST"])
@login_required
@employee_blocked
def api_log_expense():
    data = request.get_json(force=True)
    category = (data.get("category") or "").strip()
    name = (data.get("name") or "").strip() or None
    try:
        amount = int(data["amount"])
        if amount <= 0:
            raise ValueError()
    except (KeyError, ValueError, TypeError):
        return jsonify({"ok": False, "error": "укажите положительную сумму"}), 400
    if not category:
        return jsonify({"ok": False, "error": "укажите категорию"}), 400
    db.log_expense(g.shop_id, category, name, amount, data.get("expense_date"))
    return jsonify({"ok": True})


@app.route("/api/expenses/<int:entry_id>", methods=["PUT"])
@login_required
@employee_blocked
def api_update_expense(entry_id):
    data = request.get_json(force=True)
    category = (data.get("category") or "").strip()
    name = (data.get("name") or "").strip() or None
    try:
        amount = int(data["amount"])
        if amount <= 0:
            raise ValueError()
    except (KeyError, ValueError, TypeError):
        return jsonify({"ok": False, "error": "укажите положительную сумму"}), 400
    if not category:
        return jsonify({"ok": False, "error": "укажите категорию"}), 400
    from datetime import datetime as _dt
    expense_date = data.get("expense_date") or _dt.now().strftime("%Y-%m-%d")
    ok = db.update_expense_entry(entry_id, g.shop_id, category, name, amount, expense_date)
    if not ok:
        return jsonify({"ok": False, "error": "запись не найдена"}), 404
    return jsonify({"ok": True})


@app.route("/api/expenses/<int:entry_id>", methods=["DELETE"])
@login_required
@employee_blocked
def api_delete_expense(entry_id):
    ok = db.delete_expense_entry(entry_id, g.shop_id)
    if not ok:
        return jsonify({"ok": False, "error": "запись не найдена"}), 404
    return jsonify({"ok": True})


@app.route("/api/recurring_expenses")
@login_required
@employee_blocked
def api_list_recurring_expenses():
    return jsonify(db.get_recurring_expenses(g.shop_id))


@app.route("/api/recurring_expenses", methods=["POST"])
@login_required
@employee_blocked
def api_create_recurring_expense():
    data = request.get_json(force=True)
    category = (data.get("category") or "").strip()
    name = (data.get("name") or "").strip() or None
    try:
        amount = int(data["amount"])
        day_of_month = int(data["day_of_month"])
        if amount <= 0 or not (1 <= day_of_month <= 28):
            raise ValueError()
    except (KeyError, ValueError, TypeError):
        return jsonify({"ok": False, "error": "укажите сумму и число месяца (1-28)"}), 400
    if not category:
        return jsonify({"ok": False, "error": "укажите категорию"}), 400
    plan = db.create_recurring_expense(g.shop_id, category, name, amount, day_of_month)
    return jsonify({"ok": True, "expense": plan})


@app.route("/api/recurring_expenses/<int:expense_id>", methods=["DELETE"])
@login_required
@employee_blocked
def api_delete_recurring_expense(expense_id):
    ok = db.delete_recurring_expense(expense_id, g.shop_id)
    if not ok:
        return jsonify({"ok": False, "error": "не найдено"}), 404
    return jsonify({"ok": True})


@app.route("/api/recurring_expenses/<int:expense_id>/pay", methods=["POST"])
@login_required
@employee_blocked
def api_pay_recurring_expense(expense_id):
    plan = db.get_recurring_expense(expense_id, g.shop_id)
    if not plan:
        return jsonify({"ok": False, "error": "не найдено"}), 404
    db.log_expense(g.shop_id, plan["category"], plan["name"], plan["amount"], recurring_expense_id=expense_id)
    return jsonify({"ok": True})


@app.route("/api/debts/<int:plan_id>/pay", methods=["POST"])
@login_required
def api_pay_debt(plan_id):
    data = request.get_json(force=True)
    try:
        amount = int(data["amount"])
        if amount <= 0:
            raise ValueError()
    except (KeyError, ValueError, TypeError):
        return jsonify({"ok": False, "error": "укажите положительную сумму платежа"}), 400
    plan = db.log_installment_payment(plan_id, g.shop_id, amount, data.get("paid_date"))
    if not plan:
        return jsonify({"ok": False, "error": "долг не найден"}), 404
    return jsonify({"ok": True, "plan": plan})


@app.route("/api/debts/<int:plan_id>/payments")
@login_required
def api_debt_payments(plan_id):
    return jsonify(db.get_installment_payments(plan_id, g.shop_id))


# ============ Рассылки ============

@app.route("/api/broadcast/recipients")
@login_required
def api_broadcast_recipients():
    return jsonify({"count": len(db.get_all_linked_clients(g.shop_id))})


@app.route("/api/broadcast/history")
@login_required
def api_broadcast_history():
    return jsonify(db.get_recent_broadcasts(g.shop_id))


@app.route("/api/broadcast", methods=["POST"])
@login_required
def api_broadcast_create():
    data = request.get_json(force=True)
    message = (data.get("message") or "").strip()
    if not message:
        return jsonify({"ok": False, "error": "empty message"}), 400
    broadcast_id = db.create_broadcast(g.shop_id, message)
    return jsonify({"ok": True, "id": broadcast_id})


# ============ Экспорт / бэкап ============

@app.route("/api/stats")
@login_required
@employee_blocked
def api_stats():
    return jsonify(db.get_revenue_stats(g.shop_id))


@app.route("/api/aggregated_stats")
@login_required
@profit_blocked
def api_aggregated_stats():
    """Для главного аккаунта — сложенные выручка и прибыль по нему самому и
    всем его филиалам вместе, плюс разбивка кто сколько продал по отдельности.
    Для точки без филиалов просто вернёт её же собственные числа (сумма по
    пустому списку филиалов — это она сама)."""
    branches = db.get_branches(g.shop_id)
    return jsonify({
        "has_branches": len(branches) > 0,
        "branch_count": len(branches),
        "revenue": db.get_aggregated_revenue_stats(g.shop_id),
        "profit": db.get_aggregated_profit_stats(g.shop_id),
        "breakdown": db.get_branch_breakdown_stats(g.shop_id),
    })


@app.route("/api/aggregated_stats/range")
@login_required
@profit_blocked
def api_aggregated_stats_range():
    """То же самое, но за произвольный период — для календаря на вкладке
    'Все филиалы', как и у собственной статистики."""
    date_from = request.args.get("from", "")
    date_to = request.args.get("to", "")
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", date_from) or not re.match(r"^\d{4}-\d{2}-\d{2}$", date_to):
        return jsonify({"ok": False, "error": "invalid date"}), 400
    return jsonify({
        "ok": True,
        "revenue": db.get_aggregated_revenue_range(g.shop_id, date_from, date_to),
        "profit": db.get_aggregated_profit_range(g.shop_id, date_from, date_to),
        "breakdown": db.get_branch_breakdown_range(g.shop_id, date_from, date_to),
    })


@app.route("/api/my_branches")
@login_required
@profit_blocked
def api_my_branches():
    """Список филиалов главного аккаунта со сводкой по складу — для панели
    управления ценами закупки. Филиалу самому это не нужно (заблокировано
    profit_blocked)."""
    branches = db.get_branches(g.shop_id)
    summary = {s["id"]: s for s in db.get_branch_warehouse_summary(g.shop_id)}
    result = []
    for b in branches:
        s = summary.get(b["id"], {"product_count": 0, "missing_price_count": 0, "stock_value": 0})
        result.append({
            "id": b["id"], "shop_name": b["shop_name"], "username": b["username"],
            "product_count": s["product_count"], "missing_price_count": s["missing_price_count"],
            "stock_value": s.get("stock_value", 0),
        })
    return jsonify(result)


@app.route("/api/branches/<int:branch_id>/products")
@login_required
@profit_blocked
def api_branch_products(branch_id):
    """Главный аккаунт смотрит склад конкретного своего филиала — с ценой
    закупки, которую сам филиал не видит. Проверяем, что это реально его
    филиал, а не чужая точка."""
    if not db.is_branch_of(branch_id, g.shop_id):
        return jsonify({"ok": False, "error": "это не ваш филиал"}), 403
    return jsonify(db.list_products(branch_id))


@app.route("/api/branches/<int:branch_id>/products/<int:product_id>/purchase_price", methods=["POST"])
@login_required
@profit_blocked
def api_set_branch_purchase_price(branch_id, product_id):
    if not db.is_branch_of(branch_id, g.shop_id):
        return jsonify({"ok": False, "error": "это не ваш филиал"}), 403
    data = request.get_json(force=True)
    try:
        purchase_price = int(data["purchase_price"]) if data.get("purchase_price") not in (None, "") else None
    except (ValueError, TypeError):
        return jsonify({"ok": False, "error": "неверная цена"}), 400
    ok = db.set_product_purchase_price(product_id, branch_id, purchase_price)
    if not ok:
        return jsonify({"ok": False, "error": "товар не найден"}), 404
    return jsonify({"ok": True})


@app.route("/api/stats/range")
@login_required
@employee_blocked
def api_stats_range():
    date_from = request.args.get("from", "")
    date_to = request.args.get("to", "")
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", date_from) or not re.match(r"^\d{4}-\d{2}-\d{2}$", date_to):
        return jsonify({"ok": False, "error": "invalid date"}), 400
    result = db.get_revenue_range(g.shop_id, date_from, date_to)
    result["ok"] = True
    return jsonify(result)


@app.route("/api/sms_settings", methods=["POST"])
@login_required
def api_sms_settings():
    shop = db.get_shop(g.shop_id)
    if not shop or not shop.get("sms_enabled"):
        return jsonify({"ok": False, "error": "SMS не включены для вашей точки"}), 403
    data = request.get_json(force=True)
    email = (data.get("email") or "").strip()
    password = data.get("password") or ""
    if not email:
        return jsonify({"ok": False, "error": "укажите email от Eskiz"}), 400
    if not password:
        # пароль не прислан -> оставляем прежний, меняем только email
        _, existing_password = db.get_shop_eskiz_credentials(g.shop_id)
        password = existing_password or ""
    if not password:
        return jsonify({"ok": False, "error": "укажите пароль от Eskiz"}), 400
    db.set_shop_eskiz_credentials(g.shop_id, email, password)
    return jsonify({"ok": True})


def _warehouse_required():
    shop = db.get_shop(g.shop_id)
    if not shop or not shop.get("warehouse_enabled"):
        return jsonify({"ok": False, "error": "Склад не включён для вашей точки"}), 403
    return None


@app.route("/api/products")
@login_required
def api_list_products():
    products = db.list_products(g.shop_id)
    if g.is_employee or g.is_branch:
        for p in products:
            p.pop("purchase_price", None)
    return jsonify(products)


@app.route("/api/products", methods=["POST"])
@login_required
@employee_blocked
def api_create_product():
    denied = _warehouse_required()
    if denied:
        return denied
    data = request.get_json(force=True)
    try:
        category = data["category"]
        name = data["name"].strip()
        if not name:
            return jsonify({"ok": False, "error": "укажите название товара"}), 400
        if category == "other":
            unit = "pc" if data.get("unit") == "pc" else "l"
        else:
            unit = "pc" if category.startswith("filter_") else "l"
        sell_price = int(data["sell_price"]) if data.get("sell_price") not in (None, "") else None
        purchase_price = int(data["purchase_price"]) if data.get("purchase_price") not in (None, "") else None
        if g.is_branch:
            purchase_price = None  # филиал не вписывает цену закупки — это делает только главный аккаунт
        initial_stock = float(data["initial_stock"]) if data.get("initial_stock") not in (None, "") else 0
    except (KeyError, ValueError, TypeError) as e:
        return jsonify({"ok": False, "error": str(e)}), 400
    product = db.create_product(g.shop_id, category, name, unit, sell_price, purchase_price, initial_stock)
    return jsonify({"ok": True, "product": product})


@app.route("/api/products/<int:product_id>", methods=["PUT"])
@login_required
@employee_blocked
def api_update_product(product_id):
    data = request.get_json(force=True)
    try:
        purchase_price = int(data["purchase_price"]) if data.get("purchase_price") not in (None, "") else None
        if g.is_branch:
            purchase_price = None  # филиал не может менять цену закупки — только главный аккаунт
        ok = db.update_product(
            product_id, g.shop_id,
            name=data.get("name"),
            sell_price=int(data["sell_price"]) if data.get("sell_price") not in (None, "") else None,
            purchase_price=purchase_price,
        )
    except (ValueError, TypeError) as e:
        return jsonify({"ok": False, "error": str(e)}), 400
    if not ok:
        return jsonify({"ok": False, "error": "товар не найден"}), 404
    return jsonify({"ok": True})


@app.route("/api/products/<int:product_id>", methods=["DELETE"])
@login_required
@employee_blocked
def api_delete_product(product_id):
    ok = db.delete_product(product_id, g.shop_id)
    if not ok:
        return jsonify({"ok": False, "error": "товар не найден"}), 404
    return jsonify({"ok": True})


@app.route("/api/products/<int:product_id>/restock", methods=["POST"])
@login_required
@employee_blocked
def api_restock_product(product_id):
    data = request.get_json(force=True)
    try:
        quantity = float(data["quantity"])
        purchase_price = int(data["purchase_price"]) if data.get("purchase_price") not in (None, "") else None
        if g.is_branch:
            purchase_price = None  # филиал не вписывает цену закупки при пополнении — только главный аккаунт
        restock_date = data.get("restock_date") or None
    except (KeyError, ValueError, TypeError) as e:
        return jsonify({"ok": False, "error": str(e)}), 400
    ok = db.restock_product(product_id, g.shop_id, quantity, purchase_price, restock_date)
    if not ok:
        return jsonify({"ok": False, "error": "товар не найден"}), 404
    return jsonify({"ok": True})


@app.route("/api/restock_history")
@login_required
@employee_blocked
def api_restock_history():
    history = db.get_restock_history(g.shop_id)
    if g.is_branch:
        for r in history:
            r.pop("purchase_price", None)
    return jsonify(history)


@app.route("/api/profit_stats")
@login_required
@profit_blocked
def api_profit_stats():
    return jsonify(db.get_profit_stats(g.shop_id))


@app.route("/api/export")
@login_required
@employee_blocked
def api_export():
    import json
    data = db.export_shop_data(g.shop_id)
    body = json.dumps(data, ensure_ascii=False, indent=2)
    filename = f"backup_shop_{g.shop_id}_{data['exported_at'][:10]}.json"
    return Response(
        body, mimetype="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@app.route("/api/export_excel")
@login_required
@employee_blocked
def api_export_excel():
    import io
    from datetime import datetime as _dt
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    rows = db.get_full_history_flat(g.shop_id)
    shop = db.get_shop(g.shop_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "История"

    headers = [
        "Дата", "Госномер", "Владелец", "Телефон", "Марка авто", "Модель",
        "Пробег, км", "Менять при, км", "Услуга", "Марка масла",
        "Итого, сум", "След. замена", "Заметки",
    ]
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="3A5FCD", end_color="3A5FCD", fill_type="solid")
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    body_font = Font(name="Arial", size=11)
    for row_idx, r in enumerate(rows, start=2):
        values = [
            r["change_date"], r["plate_number"], r["owner_name"] or "", r["owner_phone"] or "",
            r["car_brand"] or "", r["car_model"] or "", r["mileage"], r["next_mileage"],
            r["service_type"] or "", r["oil_brand"] or "", r["cost"], r["next_change_date"] or "",
            r["notes"] or "",
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = body_font
            if col_idx == 11 and value is not None:  # "Итого, сум"
                cell.number_format = "#,##0"

    widths = [12, 12, 20, 15, 14, 14, 12, 14, 22, 16, 13, 12, 24]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    shop_name_part = (shop.get("shop_name") or "shop").replace(" ", "_") if shop else "shop"
    filename = f"{shop_name_part}_{_dt.now().strftime('%Y-%m-%d')}.xlsx"
    # Content-Disposition допускает только ASCII в filename="..." — название
    # точки почти всегда на кириллице, поэтому кодируем правильно (RFC 5987):
    # ASCII-заглушка для старых клиентов + filename*= с настоящим именем
    # в UTF-8 для всех современных браузеров.
    ascii_fallback = "history_export.xlsx"
    encoded_filename = urllib.parse.quote(filename)
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{encoded_filename}"
        }
    )


# ============ Админ-панель (платформа) ============

ADMIN_PAGE = """
<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Админ-панель — точки</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:ital,wght@0,400;0,500;0,600;0,700;0,800;1,700&family=Space+Grotesk:wght@600;700&family=IBM+Plex+Mono:wght@500;600&display=swap" rel="stylesheet">
<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
<style>
  :root {
    --bg:#F1F5F9; --text:#1E293B; --hint:#94A3B8; --btn:#E63946; --btn-text:#FFFFFF; --blue:#0F52BA; --darkblue:#0A2540; --cyan:#00A8E8;
    --card:#FFFFFF; --border:#E2E8F0; --field-bg:#F4F7FA; --ok:#059669; --ok-bg:#ECFDF5; --danger:#B3241C; --danger-bg:#FDECEA;
    --font-display:'Space Grotesk', sans-serif; --font-body:'Plus Jakarta Sans', -apple-system, sans-serif; --font-mono:'IBM Plex Mono', monospace;
  }
  * { box-sizing: border-box; }
  body { margin:0; background:var(--bg); color:var(--text); font-family: var(--font-body); }
  .speedline { height:6px; width:100%; background: linear-gradient(90deg, var(--blue) 0%, var(--blue) 33%, #fff 33%, #fff 38%, var(--btn) 38%, var(--btn) 70%, #fff 70%, #fff 75%, var(--cyan) 75%, var(--cyan) 100%); }
  .container { padding: 12px; max-width: 900px; margin: 0 auto; }
  .topbar { display:flex; justify-content:space-between; align-items:center; margin: 14px 0 16px; }
  .logo-badge {
    display:inline-flex; align-items:center; justify-content:center; width:40px; height:40px;
    background:var(--darkblue); color:var(--cyan); border-radius:12px; transform:rotate(-8deg);
    box-shadow:0 4px 10px rgba(10,37,64,.25); font-size:16px; flex:none;
  }
  h1 { font-family: var(--font-display); font-weight:800; font-style:italic; letter-spacing:-0.3px; font-size: 20px; margin: 0; line-height:1.1; color:var(--darkblue); }
  .logo-sub { font-size:10px; font-weight:700; letter-spacing:1.5px; text-transform:uppercase; color:var(--hint); }
  .logout { color: var(--btn); font-size: 12px; font-weight:700; text-decoration:none; background:var(--danger-bg); padding:6px 10px; border-radius:10px; }
  .card { background: rgba(255,255,255,.94); backdrop-filter: blur(10px); border:2px solid #DBEAFE; border-radius: 22px; padding: 16px; margin-bottom: 16px; box-shadow: 0 10px 25px -5px rgba(15,82,186,.08); }
  .field { margin-bottom: 10px; }
  .row2 { display:flex; gap:10px; }
  .row2 .field { flex:1; }
  label { display:block; font-size: 12px; color: var(--hint); margin-bottom: 4px; text-transform:uppercase; letter-spacing:0.4px; }
  input { width: 100%; padding: 10px; border-radius: 10px; border: 1.5px solid var(--border); background: var(--field-bg); color: var(--text); font-size: 15px; font-family: var(--font-body); }
  input:focus { outline:none; border-color: var(--blue); box-shadow:0 0 0 3px rgba(15,82,186,.12); }
  button.submit { width: 100%; padding: 14px; border: none; border-radius: 14px; background: linear-gradient(135deg, #1D4ED8 0%, #1E40AF 55%, #312E81 100%); color:#fff; font-size: 15px; font-weight: 700; font-family: var(--font-display); letter-spacing:0.5px; text-transform:uppercase; cursor: pointer; margin-top: 6px; box-shadow: 0 8px 18px rgba(29,78,216,.30); }
  table { width: 100%; border-collapse: collapse; font-size: 13px; }
  th, td { text-align: left; padding: 8px 6px; border-bottom: 1px solid var(--border); }
  th { color:#fff; background:#1E293B; font-weight: 700; font-size:10px; text-transform:uppercase; letter-spacing:0.4px; }
  .badge { display:inline-block; padding: 3px 9px; border-radius: 20px; font-size: 11px; font-weight:600; cursor:pointer; border:none; text-transform:uppercase; letter-spacing:0.3px; }
  .badge.active { background: var(--ok-bg); color: var(--ok); }
  .badge.inactive { background: var(--danger-bg); color: var(--danger); }
  .hint-text { color: var(--hint); font-size: 12px; margin-top: 6px; }
  .msg { padding: 12px; border-radius: 10px; margin-bottom: 10px; font-size: 14px; }
  .msg.ok { background:var(--ok-bg); color:var(--ok); }
  .msg.err { background:var(--danger-bg); color:var(--danger); }
  .new-creds { background:var(--field-bg); border:1px dashed var(--blue); border-radius:10px; padding:10px; font-size:13px; margin-top:10px; font-family: var(--font-mono); }
</style>
</head>
<body>
<div class="speedline"></div>
<div class="container">
  <div class="topbar">
    <div style="display:flex; align-items:center; gap:10px;">
      <div class="logo-badge"><i class="fa-solid fa-flag-checkered"></i></div>
      <div>
        <h1>Точки замены масла</h1>
        <div class="logo-sub">MITAL PLATFORM</div>
      </div>
    </div>
    <a class="logout" href="/logout">Выйти</a>
  </div>

  <div id="msg"></div>

  <div class="card">
    <h3 style="margin-top:0;">➕ Добавить новую точку</h3>
    <div class="field">
      <label>Название точки</label>
      <input id="new_shop_name" placeholder="MITAL Namangan">
    </div>
    <div class="field">
      <label>Клиент / группа (необяз.) — для филиала укажи то же, что у других точек этого клиента</label>
      <input id="new_client_group" list="clientGroupsList" placeholder="например: Sinov01">
      <datalist id="clientGroupsList"></datalist>
    </div>
    <div class="row2">
      <div class="field">
        <label>Логин</label>
        <input id="new_username" placeholder="namangan_point">
      </div>
      <div class="field">
        <label>Пароль (пусто = сгенерировать)</label>
        <input id="new_password" placeholder="необязательно">
      </div>
    </div>
    <div class="row2">
      <div class="field">
        <label>Телефон точки</label>
        <input id="new_phone" placeholder="+998901112233">
      </div>
      <div class="field">
        <label>Telegram ID для уведомлений (необяз.)</label>
        <input id="new_notify_id" placeholder="123456789">
      </div>
    </div>
    <div class="field">
      <label>Адрес</label>
      <input id="new_address" placeholder="Наманган, ул. ...">
    </div>
    <div class="field">
      <label>Локация (необяз.) — вставь широту и долготу из Google Карт через запятую</label>
      <input id="new_location" placeholder="40.782123, 72.344567">
      <div class="hint-text">Открой точку на Google Картах, нажми и удержи на месте — внизу появятся два числа через запятую, скопируй их сюда целиком.</div>
    </div>
    <button class="submit" onclick="createShop()">Создать точку</button>
    <div id="newCreds"></div>
  </div>

  <div class="card">
    <h3 style="margin-top:0;">Все точки</h3>
    <div class="field">
      <input id="shopSearch" placeholder="Поиск по названию, логину или телефону..." oninput="filterShops()">
    </div>
    <div class="table-wrap" style="overflow-x:auto;">
    <table>
      <thead><tr><th>Название</th><th>Логин</th><th>Пароль</th><th>Телефон</th><th>Telegram</th><th>Клиентов</th><th>Статус</th><th>SMS</th><th>Склад</th><th>Сотрудники</th><th>Филиалы</th><th>Группа</th></tr></thead>
      <tbody id="shops-body"></tbody>
    </table>
    </div>
  </div>
</div>

<script>
function showMsg(text, ok) {
  const el = document.getElementById('msg');
  el.innerHTML = `<div class="msg ${ok ? 'ok' : 'err'}">${text}</div>`;
  setTimeout(() => { el.innerHTML = ''; }, 8000);
}

function showMsgSticky(text) {
  // не исчезает сама - для одноразовых паролей, которые больше нигде не
  // будут показаны повторно, чтобы админ точно успел их скопировать
  const el = document.getElementById('msg');
  el.innerHTML = `<div class="msg ok" style="display:flex; justify-content:space-between; align-items:center; gap:10px;">
    <span>${text}</span>
    <button type="button" onclick="this.closest('.msg').remove()" style="flex:none; background:none; border:none; color:inherit; font-size:18px; cursor:pointer; padding:0 4px;">×</button>
  </div>`;
}

function escapeHtml(str) {
  return String(str)
    .replace(/&/g, '&amp;')
    .replace(/</g, '&lt;')
    .replace(/>/g, '&gt;')
    .replace(/"/g, '&quot;')
    .replace(/'/g, '&#039;');
}

let allShopsCache = [];

async function loadShops() {
  const res = await fetch('/api/admin/shops');
  allShopsCache = await res.json();

  // список известных групп — для автодополнения в форме создания новой точки
  const groupNames = [...new Set(allShopsCache.map(s => s.client_group).filter(Boolean))].sort();
  document.getElementById('clientGroupsList').innerHTML = groupNames.map(g => `<option value="${escapeHtml(g)}">`).join('');

  renderShopsTable(allShopsCache);
}

function filterShops() {
  const q = document.getElementById('shopSearch').value.trim().toLowerCase();
  if (!q) { renderShopsTable(allShopsCache); return; }
  const filtered = allShopsCache.filter(s =>
    (s.shop_name || '').toLowerCase().includes(q) ||
    (s.username || '').toLowerCase().includes(q) ||
    (s.phone || '').toLowerCase().includes(q)
  );
  renderShopsTable(filtered);
}

function renderShopsTable(shops) {
  // группируем: сначала точки с группой (по алфавиту группы), потом без группы
  const grouped = {};
  const standalone = [];
  shops.forEach(s => {
    if (s.client_group) {
      (grouped[s.client_group] = grouped[s.client_group] || []).push(s);
    } else {
      standalone.push(s);
    }
  });

  function renderShopRow(s) {
    return `
    <tr>
      <td>${s.shop_name || '—'}</td>
      <td>${s.username}</td>
      <td><span class="hint-text">🔒 скрыт</span>
          <br><button class="badge" style="background:var(--border);color:var(--hint);margin-top:4px;" onclick="resetPassword(${s.id}, ${escapeHtml(JSON.stringify(s.username))})">сбросить</button></td>
      <td>${s.phone || '—'}</td>
      <td style="min-width:160px;">
        ${s.notify_telegram_id
          ? `<div style="font-size:11.5px; color:#1B8A5A; font-weight:600;">✅ привязан</div>`
          : `<div style="font-size:11.5px; color:#B3241C; font-weight:600;">не привязан</div>`}
        ${s.owner_link
          ? `<button class="badge" style="background:#EFF6FF;color:var(--blue);margin-top:4px;" onclick="copyOwnerLink(${escapeHtml(JSON.stringify(s.owner_link))}, this)">🔗 ссылка для владельца</button>`
          : ''}
        <button class="badge" style="background:var(--border);color:var(--hint);margin-top:4px;" onclick="testNotifyTelegram(${s.id})">проверить</button>
        <details style="margin-top:4px;">
          <summary style="font-size:11px; color:var(--hint); cursor:pointer;">вручную</summary>
          <input id="notify_id_${s.id}" value="${escapeHtml(s.notify_telegram_id || '')}" placeholder="123456789" style="width:100px; font-size:11px; padding:4px; margin-top:4px;">
          <button class="badge" style="background:var(--border);color:var(--hint);margin-top:4px;" onclick="saveNotifyTelegram(${s.id})">сохранить</button>
        </details>
      </td>
      <td>${s.client_count}</td>
      <td><button class="badge ${s.is_active ? 'active' : 'inactive'}" onclick="toggleShop(${s.id}, ${s.is_active ? 0 : 1})">
        ${s.is_active ? 'активна' : 'выключена'}
      </button></td>
      <td><button class="badge ${s.sms_enabled ? 'active' : 'inactive'}" onclick="toggleSms(${s.id}, ${s.sms_enabled ? 0 : 1})">
        ${s.sms_enabled ? 'включён' : 'выключен'}
      </button></td>
      <td><button class="badge ${s.warehouse_enabled ? 'active' : 'inactive'}" onclick="toggleWarehouse(${s.id}, ${s.warehouse_enabled ? 0 : 1})">
        ${s.warehouse_enabled ? 'включён' : 'выключен'}
      </button></td>
      <td><button class="badge" style="background:#EFF6FF;color:var(--blue);" onclick="toggleEmployees(${s.id})">👥 сотрудники</button></td>
      <td><button class="badge" style="background:#FFF7ED;color:#9A3412;" onclick="toggleBranches(${s.id})">🏢 филиалы</button></td>
      <td>
        <input value="${escapeHtml(s.client_group || '')}" list="clientGroupsList" placeholder="без группы"
               style="width:120px; padding:4px 6px; font-size:12px;"
               onchange="setClientGroup(${s.id}, this.value)">
      </td>
    </tr>
    <tr id="emp-row-${s.id}" style="display:none;"><td colspan="10"><div id="emp-panel-${s.id}" style="padding:10px; background:var(--field-bg); border-radius:10px;">…</div></td></tr>
    <tr id="branch-row-${s.id}" style="display:none;"><td colspan="10"><div id="branch-panel-${s.id}" style="padding:10px; background:var(--field-bg); border-radius:10px;">…</div></td></tr>
  `;
  }

  let html = '';
  Object.keys(grouped).sort().forEach(group => {
    html += `<tr><td colspan="10" style="background:#EFF6FF; font-weight:700; color:var(--blue); padding:8px 6px;">🏷️ ${escapeHtml(group)} (${grouped[group].length})</td></tr>`;
    grouped[group].forEach(s => html += renderShopRow(s));
  });
  standalone.forEach(s => html += renderShopRow(s));

  document.getElementById('shops-body').innerHTML = html || `<tr><td colspan="10" class="hint-text" style="padding:14px;">Ничего не найдено.</td></tr>`;
}

async function setClientGroup(shopId, value) {
  await fetch(`/api/admin/shops/${shopId}/client_group`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({client_group: value.trim()})
  });
  loadShops();
}

async function toggleEmployees(shopId) {
  const row = document.getElementById(`emp-row-${shopId}`);
  const opening = row.style.display === 'none';
  row.style.display = opening ? '' : 'none';
  if (opening) await loadEmployees(shopId);
}

async function loadEmployees(shopId) {
  const panel = document.getElementById(`emp-panel-${shopId}`);
  const res = await fetch(`/api/admin/shops/${shopId}/employees`);
  const employees = await res.json();
  const list = employees.length ? employees.map(e => `
    <div style="display:flex; justify-content:space-between; align-items:center; padding:6px 0; border-bottom:1px dashed var(--border); font-size:13px;">
      <span>${escapeHtml(e.full_name || e.username)} <span class="hint-text">(${e.username})</span></span>
      <span style="display:flex; align-items:center; gap:8px;">
        <span class="hint-text">🔒 скрыт</span>
        <button class="badge" style="background:var(--border);color:var(--hint);" onclick="resetEmployeePassword(${e.id}, ${shopId}, ${escapeHtml(JSON.stringify(e.username))})">сбросить</button>
        <button class="badge inactive" onclick="deleteEmployee(${e.id}, ${shopId}, ${escapeHtml(JSON.stringify(e.username))})">удалить</button>
      </span>
    </div>
  `).join('') : `<div class="hint-text">Пока нет сотрудников у этой точки.</div>`;
  panel.innerHTML = `
    <div style="font-weight:700; font-size:13px; margin-bottom:8px;">Сотрудники точки (ограниченный доступ — без статистики, прибыли, экспорта, склада)</div>
    ${list}
    <div style="display:flex; gap:6px; margin-top:10px;">
      <input id="new-emp-username-${shopId}" placeholder="логин сотрудника" style="flex:1;">
      <button class="badge active" style="padding:6px 14px;" onclick="createEmployee(${shopId})">+ добавить</button>
    </div>
  `;
}

async function createEmployee(shopId) {
  const input = document.getElementById(`new-emp-username-${shopId}`);
  const username = input.value.trim();
  if (!username) return;
  const res = await fetch(`/api/admin/shops/${shopId}/employees`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({username})
  });
  const data = await res.json();
  if (data.ok) {
    showMsgSticky(`✅ Сотрудник «${username}» создан. Пароль (больше не увидите — сохраните сейчас): <b>${data.password}</b>`);
    loadEmployees(shopId);
  } else {
    showMsg('Ошибка: ' + data.error, false);
  }
}

async function resetEmployeePassword(employeeId, shopId, username) {
  if (!confirm(`Сбросить пароль для сотрудника «${username}»?`)) return;
  const res = await fetch(`/api/admin/employees/${employeeId}/reset_password`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({shop_id: shopId})
  });
  const data = await res.json();
  if (data.ok) {
    showMsgSticky(`✅ Новый пароль для «${username}» (больше не увидите — сохраните сейчас): <b>${data.password}</b>`);
    loadEmployees(shopId);
  } else {
    showMsg('Ошибка: ' + data.error, false);
  }
}

async function deleteEmployee(employeeId, shopId, username) {
  if (!confirm(`Удалить сотрудника «${username}»? Он больше не сможет войти.`)) return;
  const res = await fetch(`/api/admin/employees/${employeeId}?shop_id=${shopId}`, { method: 'DELETE' });
  const data = await res.json();
  if (data.ok) { loadEmployees(shopId); } else { showMsg('Ошибка: ' + data.error, false); }
}

async function toggleBranches(shopId) {
  const row = document.getElementById(`branch-row-${shopId}`);
  const opening = row.style.display === 'none';
  row.style.display = opening ? '' : 'none';
  if (opening) await loadBranches(shopId);
}

async function loadBranches(shopId) {
  const panel = document.getElementById(`branch-panel-${shopId}`);
  const res = await fetch(`/api/admin/shops/${shopId}/branches`);
  const branches = await res.json();
  const list = branches.length ? branches.map(b => `
    <div style="padding:8px 0; border-bottom:1px dashed var(--border); font-size:13px;">
      <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:6px;">
        <span>${escapeHtml(b.shop_name || b.username)} <span class="hint-text">(${b.username}, клиентов: ${b.client_count})</span></span>
        <span style="display:flex; align-items:center; gap:8px;">
          <span class="hint-text">🔒 скрыт</span>
          <button class="badge" style="background:var(--border);color:var(--hint);" onclick="resetPassword(${b.id}, ${escapeHtml(JSON.stringify(b.username))})">сбросить</button>
        </span>
      </div>
      <div style="display:flex; gap:6px; flex-wrap:wrap;">
        <button class="badge ${b.is_active ? 'active' : 'inactive'}" onclick="toggleBranchField(${shopId}, ${b.id}, 'toggle', ${b.is_active ? 0 : 1})">${b.is_active ? 'активен' : 'выключен'}</button>
        <button class="badge ${b.sms_enabled ? 'active' : 'inactive'}" onclick="toggleBranchField(${shopId}, ${b.id}, 'toggle_sms', ${b.sms_enabled ? 0 : 1})">SMS: ${b.sms_enabled ? 'включён' : 'выключен'}</button>
        <button class="badge ${b.warehouse_enabled ? 'active' : 'inactive'}" onclick="toggleBranchField(${shopId}, ${b.id}, 'toggle_warehouse', ${b.warehouse_enabled ? 0 : 1})">Склад: ${b.warehouse_enabled ? 'включён' : 'выключен'}</button>
      </div>
    </div>
  `).join('') : `<div class="hint-text">У этой точки пока нет филиалов.</div>`;
  panel.innerHTML = `
    <div style="font-weight:700; font-size:13px; margin-bottom:8px;">Филиалы (полноценные точки — свой склад, своя база, без цены закупки и без прибыли по отдельности)</div>
    ${list}
    <div style="display:flex; gap:6px; margin-top:10px; flex-wrap:wrap;">
      <input id="new-branch-name-${shopId}" placeholder="название филиала" style="flex:1; min-width:140px;">
      <input id="new-branch-username-${shopId}" placeholder="логин" style="flex:1; min-width:120px;">
      <button class="badge active" style="padding:6px 14px;" onclick="createBranch(${shopId})">+ добавить филиал</button>
    </div>
  `;
}

async function createBranch(shopId) {
  const shopName = document.getElementById(`new-branch-name-${shopId}`).value.trim();
  const username = document.getElementById(`new-branch-username-${shopId}`).value.trim();
  if (!shopName || !username) { showMsg('Укажите название филиала и логин.', false); return; }
  const res = await fetch(`/api/admin/shops/${shopId}/branches`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({shop_name: shopName, username})
  });
  const data = await res.json();
  if (data.ok) {
    showMsgSticky(`✅ Филиал «${shopName}» создан. Логин: <b>${username}</b>, пароль (больше не увидите — сохраните сейчас): <b>${data.password}</b>`);
    loadBranches(shopId);
  } else {
    showMsg('Ошибка: ' + data.error, false);
  }
}

async function toggleWarehouse(id, makeEnabled) {
  await fetch(`/api/admin/shops/${id}/toggle_warehouse`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({enabled: !!makeEnabled})
  });
  loadShops();
}

async function toggleShop(id, makeActive) {
  await fetch(`/api/admin/shops/${id}/toggle`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({active: !!makeActive})
  });
  loadShops();
}

async function toggleBranchField(headShopId, branchId, endpoint, value) {
  const bodyKey = endpoint === 'toggle' ? 'active' : 'enabled';
  await fetch(`/api/admin/shops/${branchId}/${endpoint}`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({[bodyKey]: !!value})
  });
  loadBranches(headShopId);
}

async function toggleSms(id, makeEnabled) {
  await fetch(`/api/admin/shops/${id}/toggle_sms`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({enabled: !!makeEnabled})
  });
  loadShops();
}

async function resetPassword(id, username) {
  if (!confirm(`Сбросить пароль для «${username}»? Старый пароль перестанет работать.`)) return;
  const res = await fetch(`/api/admin/shops/${id}/reset_password`, { method: 'POST' });
  const data = await res.json();
  if (data.ok) {
    showMsgSticky(`✅ Новый пароль для «${username}» (больше не увидите — сохраните сейчас): <b>${data.password}</b>`);
    loadShops();
  } else {
    showMsg('Ошибка: ' + data.error, false);
  }
}

function copyOwnerLink(link, btn) {
  const done = () => {
    const original = btn.textContent;
    btn.textContent = '✅ скопировано';
    setTimeout(() => { btn.textContent = original; }, 1500);
  };
  if (navigator.clipboard && navigator.clipboard.writeText) {
    navigator.clipboard.writeText(link).then(done).catch(() => {
      prompt('Скопируйте ссылку вручную:', link);
    });
  } else {
    prompt('Скопируйте ссылку вручную:', link);
  }
}

async function saveNotifyTelegram(id) {
  const notify_telegram_id = document.getElementById(`notify_id_${id}`).value.trim();
  const res = await fetch(`/api/admin/shops/${id}/notify_telegram`, {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({notify_telegram_id}),
  });
  const data = await res.json();
  if (data.ok) {
    showMsg('✅ Telegram ID сохранён', true);
  } else {
    showMsg('Ошибка: ' + data.error, false);
  }
}

async function testNotifyTelegram(id) {
  const res = await fetch(`/api/admin/shops/${id}/test_telegram`, { method: 'POST' });
  const data = await res.json();
  if (data.ok) {
    showMsg('✅ Тестовое сообщение доставлено — связь работает', true);
  } else {
    showMsgSticky(`❌ Не удалось отправить: ${data.error}`);
  }
}

async function createShop() {
  const payload = {
    shop_name: document.getElementById('new_shop_name').value.trim(),
    client_group: document.getElementById('new_client_group').value.trim(),
    username: document.getElementById('new_username').value.trim(),
    password: document.getElementById('new_password').value.trim(),
    phone: document.getElementById('new_phone').value.trim(),
    notify_telegram_id: document.getElementById('new_notify_id').value.trim(),
    address: document.getElementById('new_address').value.trim(),
  };
  const loc = document.getElementById('new_location').value.trim();
  if (loc) {
    const parts = loc.split(',').map(p => p.trim()).filter(Boolean);
    if (parts.length === 2 && !isNaN(parseFloat(parts[0])) && !isNaN(parseFloat(parts[1]))) {
      payload.lat = parts[0];
      payload.lon = parts[1];
    } else {
      showMsg('Локация должна быть в формате: широта, долгота (два числа через запятую).', false);
      return;
    }
  }
  if (!payload.shop_name || !payload.username) {
    showMsg('Заполните хотя бы название точки и логин.', false);
    return;
  }
  const res = await fetch('/api/admin/shops', {
    method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify(payload)
  });
  const data = await res.json();
  if (data.ok) {
    document.getElementById('newCreds').innerHTML =
      `<div class="new-creds">✅ Точка создана. Логин: <b>${data.username}</b>, пароль: <b>${data.password}</b><br>
       Сохраните пароль сейчас — второй раз он нигде не показывается.</div>`;
    ['new_shop_name','new_client_group','new_username','new_password','new_phone','new_notify_id','new_address','new_location'].forEach(id => document.getElementById(id).value = '');
    loadShops();
  } else {
    showMsg('Ошибка: ' + data.error, false);
  }
}

loadShops();
</script>
</body>
</html>
"""


@app.route("/admin")
@admin_required
def admin_page():
    return render_template_string(ADMIN_PAGE)


@app.route("/api/admin/shops")
@admin_required
def api_admin_shops():
    shops = db.list_shops()
    for s in shops:
        s["owner_link"] = _client_link(f"owner_{s['owner_link_token']}") if s.get("owner_link_token") else None
    return jsonify(shops)


@app.route("/api/admin/shops", methods=["POST"])
@admin_required
def api_admin_create_shop():
    data = request.get_json(force=True)
    username = (data.get("username") or "").strip()
    shop_name = (data.get("shop_name") or "").strip()
    if not username or not shop_name:
        return jsonify({"ok": False, "error": "укажите логин и название точки"}), 400
    if db.username_taken(username):
        return jsonify({"ok": False, "error": "такой логин уже занят"}), 400

    password = (data.get("password") or "").strip() or secrets.token_urlsafe(6)
    shop = db.create_shop(
        username, password, shop_name=shop_name,
        phone=data.get("phone") or None, address=data.get("address") or None,
        hours=data.get("hours") or None,
        lat=float(data["lat"]) if data.get("lat") else None,
        lon=float(data["lon"]) if data.get("lon") else None,
        notify_telegram_id=data.get("notify_telegram_id") or None,
        client_group=(data.get("client_group") or "").strip() or None,
    )
    return jsonify({"ok": True, "id": shop["id"], "username": username, "password": password})


@app.route("/api/admin/shops/<int:shop_id>/client_group", methods=["POST"])
@admin_required
def api_admin_set_client_group(shop_id):
    data = request.get_json(force=True)
    db.set_shop_client_group(shop_id, (data.get("client_group") or "").strip())
    return jsonify({"ok": True})


@app.route("/api/admin/shops/<int:shop_id>/toggle", methods=["POST"])
@admin_required
def api_admin_toggle_shop(shop_id):
    data = request.get_json(force=True)
    db.set_shop_active(shop_id, bool(data.get("active")))
    return jsonify({"ok": True})


@app.route("/api/admin/shops/<int:shop_id>/toggle_sms", methods=["POST"])
@admin_required
def api_admin_toggle_sms(shop_id):
    data = request.get_json(force=True)
    db.set_shop_sms_enabled(shop_id, bool(data.get("enabled")))
    return jsonify({"ok": True})


@app.route("/api/admin/shops/<int:shop_id>/toggle_warehouse", methods=["POST"])
@admin_required
def api_admin_toggle_warehouse(shop_id):
    data = request.get_json(force=True)
    db.set_shop_warehouse_enabled(shop_id, bool(data.get("enabled")))
    return jsonify({"ok": True})


@app.route("/api/admin/shops/<int:shop_id>/reset_password", methods=["POST"])
@admin_required
def api_admin_reset_password(shop_id):
    shop = db.get_shop(shop_id)
    if not shop:
        return jsonify({"ok": False, "error": "shop not found"}), 404
    new_password = secrets.token_urlsafe(6)
    db.reset_shop_password(shop_id, new_password)
    return jsonify({"ok": True, "password": new_password})


@app.route("/api/admin/shops/<int:shop_id>/notify_telegram", methods=["POST"])
@admin_required
def api_admin_set_notify_telegram(shop_id):
    shop = db.get_shop(shop_id)
    if not shop:
        return jsonify({"ok": False, "error": "shop not found"}), 404
    data = request.get_json(force=True)
    notify_id = (data.get("notify_telegram_id") or "").strip()
    db.set_shop_notify_telegram_id(shop_id, notify_id or None)
    return jsonify({"ok": True, "notify_telegram_id": notify_id or None})


@app.route("/api/admin/shops/<int:shop_id>/test_telegram", methods=["POST"])
@admin_required
def api_admin_test_telegram(shop_id):
    """Пробная отправка — сразу видно, реально ли бот может писать этому
    получателю (частая причина 'не приходит' — получатель ни разу не писал
    боту первым, Telegram такое запрещает)."""
    shop = db.get_shop(shop_id)
    if not shop:
        return jsonify({"ok": False, "error": "shop not found"}), 404
    notify_id = shop.get("notify_telegram_id")
    if not notify_id:
        return jsonify({"ok": False, "error": "у точки не указан Telegram ID"}), 400
    text = f"✅ Тестовое сообщение от платформы — если вы это видите, связь с точкой «{shop.get('shop_name') or shop['username']}» настроена верно."
    sent = _send_telegram_message(notify_id, text)
    if not sent:
        return jsonify({"ok": False, "error": "не удалось отправить — скорее всего, получатель ни разу не писал этому боту. Попросите его открыть бота в Telegram и нажать «Старт»"}), 400
    return jsonify({"ok": True})


@app.route("/api/admin/shops/<int:shop_id>/employees")
@admin_required
def api_admin_list_employees(shop_id):
    return jsonify(db.list_shop_employees(shop_id))


@app.route("/api/admin/shops/<int:shop_id>/employees", methods=["POST"])
@admin_required
def api_admin_create_employee(shop_id):
    shop = db.get_shop(shop_id)
    if not shop:
        return jsonify({"ok": False, "error": "точка не найдена"}), 404
    data = request.get_json(force=True)
    username = (data.get("username") or "").strip()
    full_name = (data.get("full_name") or "").strip() or None
    if not username:
        return jsonify({"ok": False, "error": "укажите логин"}), 400
    result = db.create_shop_employee(shop_id, username, full_name=full_name)
    if not result:
        return jsonify({"ok": False, "error": "такой логин уже занят"}), 400
    return jsonify({"ok": True, **result})


@app.route("/api/admin/shops/<int:shop_id>/branches")
@admin_required
def api_admin_list_branches(shop_id):
    return jsonify(db.get_branches(shop_id))


@app.route("/api/admin/shops/<int:shop_id>/branches", methods=["POST"])
@admin_required
def api_admin_create_branch(shop_id):
    parent = db.get_shop(shop_id)
    if not parent:
        return jsonify({"ok": False, "error": "главная точка не найдена"}), 404
    data = request.get_json(force=True)
    username = (data.get("username") or "").strip()
    shop_name = (data.get("shop_name") or "").strip()
    if not username or not shop_name:
        return jsonify({"ok": False, "error": "укажите логин и название филиала"}), 400
    if db.username_taken(username):
        return jsonify({"ok": False, "error": "такой логин уже занят"}), 400
    password = (data.get("password") or "").strip() or secrets.token_urlsafe(6)
    branch = db.create_branch_shop(
        shop_id, username, password, shop_name=shop_name,
        phone=data.get("phone") or None, address=data.get("address") or None,
    )
    db.set_shop_warehouse_enabled(branch["id"], True)  # филиалу склад нужен сразу, это весь смысл филиала
    return jsonify({"ok": True, "id": branch["id"], "username": username, "password": password})


@app.route("/api/admin/employees/<int:employee_id>/reset_password", methods=["POST"])
@admin_required
def api_admin_reset_employee_password(employee_id):
    shop_id = request.get_json(force=True).get("shop_id")
    new_password = db.reset_shop_employee_password(employee_id, shop_id)
    if not new_password:
        return jsonify({"ok": False, "error": "сотрудник не найден"}), 404
    return jsonify({"ok": True, "password": new_password})


@app.route("/api/admin/employees/<int:employee_id>", methods=["DELETE"])
@admin_required
def api_admin_delete_employee(employee_id):
    shop_id = request.get_json(force=True).get("shop_id") if request.data else None
    if not shop_id:
        shop_id = request.args.get("shop_id", type=int)
    ok = db.delete_shop_employee(employee_id, shop_id)
    if not ok:
        return jsonify({"ok": False, "error": "сотрудник не найден"}), 404
    return jsonify({"ok": True})


# ============ ТАБЛО (для ANPR-камеры + телевизора у входа, своё на каждую точку) ============

def _extract_plate_from_request():
    plate = request.args.get("plate")
    if plate:
        return plate
    if request.is_json:
        data = request.get_json(silent=True) or {}
        for key in ("plate", "plateNumber", "licensePlate", "carNumber", "car_number"):
            if data.get(key):
                return data[key]
    if request.form.get("plate"):
        return request.form.get("plate")
    for f in request.files.values():
        if (f.mimetype or "").endswith("xml") or (f.filename or "").endswith(".xml"):
            xml_text = f.read().decode("utf-8", errors="ignore")
            m = re.search(r"<(?:licensePlate|plateNumber)>([^<]+)</(?:licensePlate|plateNumber)>", xml_text)
            if m:
                return m.group(1)
    return None


@app.route("/api/anpr/<anpr_token>", methods=["GET", "POST"])
def api_anpr(anpr_token):
    """Сюда камера ОДНОЙ КОНКРЕТНОЙ точки присылает распознанный номер —
    токен в самом адресе однозначно определяет точку, так что табло разных
    точек никогда не пересекаются."""
    shop = db.get_shop_by_anpr_token(anpr_token)
    if not shop or not shop["is_active"]:
        return jsonify({"ok": False, "error": "invalid shop token"}), 403

    plate = _extract_plate_from_request()
    if not plate:
        return jsonify({"ok": False, "error": "no plate found in request"}), 400

    plate = db.normalize_plate(plate)
    with _display_lock:
        _display_states[shop["id"]] = {"plate": plate, "shown_at": time.time()}

    return jsonify({"ok": True, "plate": plate})


@app.route("/api/display_state/<anpr_token>")
def api_display_state(anpr_token):
    shop = db.get_shop_by_anpr_token(anpr_token)
    if not shop or not shop["is_active"]:
        return jsonify({"active": False})

    with _display_lock:
        state = _display_states.get(shop["id"], {"plate": None, "shown_at": 0})

    if not state["plate"] or (time.time() - state["shown_at"]) > DISPLAY_SHOW_SECONDS:
        return jsonify({"active": False})

    car, history = db.get_car_history(shop["id"], state["plate"])
    if not car:
        return jsonify({"active": True, "found": False, "plate": state["plate"]})

    last = history[0] if history else None
    return jsonify({
        "active": True, "found": True, "plate": state["plate"],
        "owner_name": car["owner_name"],
        "last_service_date": last["change_date"] if last else None,
        "oil_brand": last["oil_brand"] if last else None,
        "service_type": last["service_type"] if last else None,
    })


DISPLAY_PAGE = """
<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0, user-scalable=no">
<title>{{ T.app_title }}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Teko:wght@500;600;700&family=Exo+2:wght@400;500;600;700&family=IBM+Plex+Mono:wght@500;600&display=swap" rel="stylesheet">
<style>
  * { box-sizing: border-box; margin:0; padding:0; }
  body {
    background: radial-gradient(circle at center, #1A0B0A 0%, #0A0A0B 100%);
    color: #F5F5F2; font-family: 'Exo 2', -apple-system, sans-serif;
    height: 100vh; display:flex; align-items:center; justify-content:center;
    overflow: hidden; text-align:center;
  }
  .idle .shop { font-family:'Teko', sans-serif; font-weight:600; font-size: 4.5vw; letter-spacing:1px; opacity:.9; text-transform:uppercase; }
  .idle .clock { font-family:'IBM Plex Mono', monospace; font-weight:600; font-size: 10vw; margin-top: 2vh; font-variant-numeric: tabular-nums; color:#E8352E; }
  .idle .date { font-size: 2.2vw; opacity:.6; margin-top:1vh; }
  .active { animation: fadein .4s ease; }
  .active .greet { font-family:'Teko', sans-serif; font-weight:600; font-size: 5vw; color:#3FBE7E; text-transform:uppercase; }
  .active .plate { font-family:'IBM Plex Mono', monospace; font-size: 6vw; font-weight:600; letter-spacing:4px; margin: 3vh 0; padding: 1vh 3vw;
    border: 4px solid #F5F5F2; border-radius: 6px; display:inline-block; }
  .active .info { font-size: 2.4vw; opacity:.85; line-height:1.7; margin-top:2vh; }
  .active .notfound { font-size: 3vw; opacity:.8; margin-top:3vh; }
  @keyframes fadein { from{opacity:0; transform:scale(.97);} to{opacity:1; transform:scale(1);} }
</style>
</head>
<body>
<div id="screen"></div>
<script>
const T = {{ t_json|safe }};
const SHOP_NAME = {{ shop_name|tojson }};
const ANPR_TOKEN = {{ anpr_token|tojson }};

function pad(n) { return n.toString().padStart(2, '0'); }
function renderIdle() {
  const now = new Date();
  const days = T.days_of_week;
  document.getElementById('screen').innerHTML = `
    <div class="idle">
      <div class="shop">${SHOP_NAME || '🔧 ' + T.app_title}</div>
      <div class="clock">${pad(now.getHours())}:${pad(now.getMinutes())}</div>
      <div class="date">${days[now.getDay()]}, ${pad(now.getDate())}.${pad(now.getMonth()+1)}.${now.getFullYear()}</div>
    </div>`;
}

function renderActive(d) {
  if (!d.found) {
    document.getElementById('screen').innerHTML = `
      <div class="active">
        <div class="greet">${T.display_welcome} 👋</div>
        <div class="plate">${d.plate}</div>
        <div class="notfound">${T.display_not_client_yet}</div>
      </div>`;
    return;
  }
  const last = d.last_service_date
    ? `${T.display_last_service} ${d.last_service_date}${d.service_type ? ' — ' + d.service_type : ''}${d.oil_brand ? ' (' + d.oil_brand + ')' : ''}`
    : T.display_no_history;
  document.getElementById('screen').innerHTML = `
    <div class="active">
      <div class="greet">${T.display_greeting} ${d.owner_name || ''}!</div>
      <div class="plate">${d.plate}</div>
      <div class="info">${last}</div>
    </div>`;
}

async function tick() {
  try {
    const res = await fetch('/api/display_state/' + ANPR_TOKEN);
    const d = await res.json();
    if (d.active) renderActive(d); else renderIdle();
  } catch (e) { renderIdle(); }
}

tick();
setInterval(tick, 2000);
</script>
</body>
</html>
"""


@app.route("/display/<anpr_token>")
def display_page(anpr_token):
    import json as _json
    shop = db.get_shop_by_anpr_token(anpr_token)
    if not shop or not shop["is_active"]:
        return "Табло не найдено — проверьте ссылку.", 404
    T = i18n.get_texts(shop.get("language") or "ru")
    return render_template_string(
        DISPLAY_PAGE, shop_name=shop.get("shop_name") or "", anpr_token=anpr_token,
        T=T, t_json=_json.dumps(T, ensure_ascii=False),
    )


def run_webapp():
    port = int(os.environ.get("PORT", 8000))
    db.init_db()
    app.run(host="0.0.0.0", port=port, use_reloader=False, threaded=True)


def run_webapp_in_thread():
    t = threading.Thread(target=run_webapp, daemon=True)
    t.start()
    return t


if __name__ == "__main__":
    run_webapp()
